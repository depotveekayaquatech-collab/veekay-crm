"""Capping history: the effective-date rule, changing a capping (backdated, duplicate dates, who may), the capping sheet and its downloads, and the
average-vs-limit alert (average above the limit, one day above it, exactly the limit is fine)."""
import csv
import io
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.config import settings
from app.db.session import engine
from app.models.capping_history import CappingHistory
from app.models.store_cap import StoreCap
from app.services import capping_history_service as svc

from .conftest import API, plant_entry


@pytest.fixture()
def zstore(make_store):
    return make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")


@pytest.fixture(autouse=True)
def _forget_history(db):
    """Rows of the stores a test made go with the stores (cascade); the seeded Zepto stores keep none."""
    yield
    db.rollback()


def _today():
    return svc.today_ist()


def _change(client, headers, store, capping, day):
    return client.post(f"{API}/capping/change", json={"storeId": str(store.id), "capping": capping, "effectiveFrom": day.isoformat()}, headers=headers)


def _seed(db, store, *rows):
    for eff, cap in rows:
        db.add(CappingHistory(store_id=store.id, capping=cap, effective_from=eff))
    db.commit()


def _sheet_row(body, store):
    return next(r for r in body["rows"] if r["store_id"] == str(store.id))


# ---------------------------------------------------------------- the effective-date rule (pure)

def test_the_capping_on_a_day_is_the_one_with_the_greatest_effective_from_on_or_before_it():
    history = [(date(2026, 10, 1), 10), (date(2026, 10, 15), 11)]
    assert svc.effective_capping(history, date(2026, 9, 30)) is None             # before the first entry
    assert [svc.effective_capping(history, date(2026, 10, d)) for d in (1, 2, 14)] == [10, 10, 10]
    assert [svc.effective_capping(history, date(2026, 10, d)) for d in (15, 16, 31)] == [11, 11, 11]
    assert svc.effective_capping([], date(2026, 10, 1)) is None


def test_a_backdated_entry_slots_into_the_timeline_without_moving_the_later_ones():
    history = sorted([(date(2026, 10, 1), 10), (date(2026, 10, 15), 11), (date(2026, 10, 8), 12)])      # added last, dated in the middle
    assert [svc.effective_capping(history, date(2026, 10, d)) for d in (7, 8, 14, 15)] == [10, 12, 12, 11]


# ---------------------------------------------------------------- the alert rules (pure)

def test_average_above_the_limit_is_an_alert():
    # average 12.5 over two days: the average and both days are above 11
    a = svc.assess(total=25, days=2, breaches=[(date(2026, 10, 1), 12), (date(2026, 10, 2), 13)], limit=11)
    assert a["alert"] and a["average_breach"] and a["day_breach"] and a["reasons"] == ["average", "day"] and a["status"] == "over"
    # the average rule on its own, without being given the days: 11.1 is above 11, 11.0 is not
    assert svc.assess(total=111, days=10, breaches=[], limit=11)["average_breach"] is True
    assert svc.assess(total=110, days=10, breaches=[], limit=11)["average_breach"] is False


def test_one_day_above_the_limit_is_an_alert_even_when_the_average_is_fine():
    a = svc.assess(total=27, days=4, breaches=[(date(2026, 10, 2), 12)], limit=11)       # 5, 12, 6, 4: average 6.75
    assert a["alert"] and a["day_breach"] and not a["average_breach"] and a["reasons"] == ["day"] and a["status"] == "over"


def test_exactly_the_limit_is_not_a_breach():
    a = svc.assess(total=11 * 5, days=5, breaches=[], limit=11)                  # every day 11: average exactly 11 and no day above
    assert not a["alert"] and a["reasons"] == [] and a["status"] == "warn"       # within 1 of the limit: amber, not red
    assert svc.assess(total=100, days=10, breaches=[], limit=11)["status"] == "ok"           # 10.0 is not within the margin
    assert svc.assess(total=101, days=10, breaches=[], limit=11)["status"] == "warn"          # 10.1 is
    assert svc.assess(total=0, days=0, breaches=[], limit=11)["status"] == "none"            # no entries: nothing to compare


def test_the_capping_average_counts_only_days_with_a_capping_in_force():
    hist = [(date(2026, 10, 1), 10), (date(2026, 10, 15), 11)]
    assert svc.capping_totals(hist, date(2026, 10, 1), date(2026, 10, 31)) == (10 * 14 + 11 * 17, 31)
    assert svc.capping_totals(hist, date(2026, 10, 10), date(2026, 10, 16)) == (10 * 5 + 11 * 2, 7)
    assert svc.capping_totals(hist, date(2026, 9, 28), date(2026, 10, 2)) == (20, 2)          # nothing in force before the 1st: those days are left out
    assert svc.capping_totals(hist, date(2026, 9, 1), date(2026, 9, 30)) == (0, 0)
    assert svc.capping_totals([], date(2026, 10, 1), date(2026, 10, 31)) == (0, 0)
    assert svc.capping_totals(hist, date(2026, 10, 5), date(2026, 10, 4)) == (0, 0)


def test_a_capping_average_above_the_limit_is_an_alert_on_its_own():
    a = svc.assess(total=50, days=5, breaches=[], limit=11, capping_total=12 * 10, capping_days=10)       # entries average 10, capping averages 12
    assert a["alert"] and a["capping_breach"] and not a["average_breach"] and a["reasons"] == ["capping"] and a["status"] == "over"
    exact = svc.assess(total=50, days=5, breaches=[], limit=11, capping_total=11 * 10, capping_days=10)   # exactly 11 is fine
    assert not exact["alert"] and exact["reasons"] == [] and exact["status"] == "warn"
    assert svc.assess(total=0, days=0, breaches=[], limit=11, capping_total=90, capping_days=10)["status"] == "ok"      # capping alone, 9.0


def test_the_limit_is_one_setting():
    assert settings.MAX_DAILY_LIMIT == 11 and svc.max_daily_limit() == settings.MAX_DAILY_LIMIT


# ---------------------------------------------------------------- changing a capping

def test_a_change_today_is_saved_and_becomes_the_current_capping(client, admin, zstore, db):
    r = _change(client, admin, zstore, 10, _today())
    assert r.status_code == 201, r.text
    body = r.json()
    assert (body["capping"], body["backdated"], body["current_capping"], body["effective_from"]) == (10, False, 10, _today().isoformat())
    db.expire_all()
    assert db.query(StoreCap).filter_by(store_id=zstore.id).one().daily_cap == 10           # the old field is kept in step
    h = client.get(f"{API}/capping/history", params={"storeId": str(zstore.id)}, headers=admin).json()
    assert [(x["capping"], x["effective_from"]) for x in h["rows"]] == [(10, _today().isoformat())]
    assert h["rows"][0]["changed_by"] and h["rows"][0]["created_at"]


def test_the_example_of_10_from_the_first_and_11_from_the_fifteenth(client, admin, zstore, db):
    first, mid = date(2026, 10, 1), date(2026, 10, 15)
    assert _change(client, admin, zstore, 10, first).status_code == 201
    assert _change(client, admin, zstore, 11, mid).status_code == 201
    sheet = client.get(f"{API}/capping/sheet", params={"from": "2026-10-01", "to": "2026-10-31", "storeId": str(zstore.id)}, headers=admin).json()
    cells = _sheet_row(sheet, zstore)["cells"]
    assert cells[:14] == [10] * 14 and cells[14:] == [11] * 17 and len(sheet["dates"]) == 31


def test_a_backdated_change_is_allowed_and_flagged_and_does_not_move_the_current_capping(client, admin, zstore, db):
    today = _today()
    assert _change(client, admin, zstore, 11, today - timedelta(days=2)).json()["backdated"] is True
    r = _change(client, admin, zstore, 9, today - timedelta(days=20))               # older than the existing one
    assert r.status_code == 201 and r.json()["backdated"] is True
    assert r.json()["current_capping"] == 11                                         # the later change still rules today
    db.expire_all()
    assert db.query(StoreCap).filter_by(store_id=zstore.id).one().daily_cap == 11
    dates = [x["effective_from"] for x in client.get(f"{API}/capping/history", params={"storeId": str(zstore.id)}, headers=admin).json()["rows"]]
    assert dates == sorted(dates, reverse=True) and len(dates) == 2                  # newest first


def test_a_second_change_for_the_same_date_is_refused_with_a_clear_message(client, admin, zstore, db):
    day = _today() - timedelta(days=3)
    assert _change(client, admin, zstore, 10, day).status_code == 201
    r = _change(client, admin, zstore, 12, day)
    assert r.status_code == 409 and "already exists" in r.json()["detail"] and f"{day:%d-%b-%Y}" in r.json()["detail"]
    db.expire_all()
    assert db.query(CappingHistory).filter_by(store_id=zstore.id).count() == 1       # the first stays as it was


@pytest.mark.parametrize("capping", [0, -3, 1.5, "7", True, None])
def test_the_capping_must_be_a_whole_number_of_at_least_one(client, admin, zstore, capping):
    r = client.post(f"{API}/capping/change", json={"storeId": str(zstore.id), "capping": capping, "effectiveFrom": _today().isoformat()}, headers=admin)
    assert r.status_code == 422, r.text


def test_the_date_and_the_store_are_checked(client, admin, zstore):
    assert _change(client, admin, zstore, 10, _today() + timedelta(days=1)).status_code == 422          # no future dates
    bad = client.post(f"{API}/capping/change", json={"storeId": str(zstore.id), "capping": 10, "effectiveFrom": "31-02-2026"}, headers=admin)
    assert bad.status_code == 422
    ghost = client.post(f"{API}/capping/change", json={"storeId": str(uuid.uuid4()), "capping": 10, "effectiveFrom": _today().isoformat()}, headers=admin)
    assert ghost.status_code == 404
    assert client.post(f"{API}/capping/change", json={"storeId": str(zstore.id), "capping": 10, "effectiveFrom": _today().isoformat(), "extra": 1}, headers=admin).status_code == 422


def test_a_store_of_another_platform_is_not_found(client, admin, make_store):
    blinkit = make_store(platform="blinkit")
    assert _change(client, admin, blinkit, 10, _today()).status_code == 404


def test_only_someone_holding_capping_change_may_change_and_everyone_who_sees_capping_may_view(client, zepto_emp, emp, admin, zstore, make_user):
    assert _change(client, zepto_emp, zstore, 10, _today()).status_code == 403                    # a Zepto employee can view, not change
    assert _change(client, emp, zstore, 10, _today()).status_code == 403
    assert client.post(f"{API}/capping/change", json={}).status_code in (401, 403)
    for url in ("/capping/sheet", "/capping/summary", "/capping/sheet/download"):
        assert client.get(f"{API}{url}", headers=zepto_emp).status_code == 200, url
        assert client.get(f"{API}{url}", headers=emp).status_code == 403, url                     # a Blinkit employee sees no Zepto capping
    assert client.get(f"{API}/capping/history", params={"storeId": str(zstore.id)}, headers=zepto_emp).status_code == 200

    from .conftest import auth, login
    code, pw, _ = make_user("employee", platform="zepto", grants=("orders.view", "capping.change"))   # ticked for one person, like a manager
    ticked = auth(login(client, code, pw).json()["access_token"])
    assert _change(client, ticked, zstore, 10, _today()).status_code == 201
    assert _change(client, admin, zstore, 12, _today() - timedelta(days=1)).status_code == 201


def test_the_admin_role_holds_the_permission(db):
    from app.models.permission import Permission, RolePermission
    from app.models.role import Role

    n = (db.query(RolePermission).join(Role, Role.id == RolePermission.role_id).join(Permission, Permission.id == RolePermission.permission_id)
         .filter(Role.code == "admin", Permission.code == "capping.change").count())
    assert n == 1


def test_there_is_no_way_to_edit_or_delete_a_history_row(client, admin, zstore):
    row = _change(client, admin, zstore, 10, _today()).json()
    for method in ("put", "patch", "delete"):
        for url in (f"{API}/capping/change", f"{API}/capping/history", f"{API}/capping/history/{row['id']}", f"{API}/capping/change/{row['id']}"):
            if method == "patch" and url.endswith(f"/history/{row['id']}"):
                continue                                                  # the one allowed edit: the START DATE (see the effective-date tests)
            assert getattr(client, method)(url, headers=admin).status_code in (404, 405), (method, url)


def test_the_older_single_store_edit_keeps_the_history_in_step(client, zepto_emp, zstore, db):
    r = client.put(f"{API}/capping/stores/{zstore.id}", json={"daily_cap": 9}, headers=zepto_emp)
    assert r.status_code == 200, r.text
    assert client.put(f"{API}/capping/stores/{zstore.id}", json={"daily_cap": 10}, headers=zepto_emp).status_code == 200      # the same day again: corrected, not a clash
    db.expire_all()
    rows = db.query(CappingHistory).filter_by(store_id=zstore.id).all()
    assert [(x.capping, x.effective_from) for x in rows] == [(10, _today())]
    assert client.put(f"{API}/capping/stores/{zstore.id}", json={"daily_cap": None}, headers=zepto_emp).status_code == 200      # "no cap" cannot be history; nothing breaks
    db.expire_all()
    assert db.query(CappingHistory).filter_by(store_id=zstore.id).count() == 1


# ---------------------------------------------------------------- the sheet and its downloads

def test_the_sheet_has_one_row_per_store_and_one_column_per_date_and_empty_before_the_first_change(client, admin, zstore, db):
    _seed(db, zstore, (date(2026, 10, 3), 10))
    r = client.get(f"{API}/capping/sheet", params={"from": "2026-10-01", "to": "2026-10-05"}, headers=admin)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dates"] == [f"2026-10-0{d}" for d in range(1, 6)]
    assert _sheet_row(body, zstore)["cells"] == [None, None, 10, 10, 10]
    assert len({x["store_id"] for x in body["rows"]}) == len(body["rows"])


def test_the_sheet_defaults_to_this_month_and_refuses_a_bad_period(client, admin):
    body = client.get(f"{API}/capping/sheet", headers=admin).json()
    today = _today()
    assert body["from"] == today.replace(day=1).isoformat() and body["to"] >= today.isoformat()
    assert client.get(f"{API}/capping/sheet", params={"from": "2026-10-05", "to": "2026-10-01"}, headers=admin).status_code == 422
    assert client.get(f"{API}/capping/sheet", params={"from": "2020-01-01", "to": "2020-01-31"}, headers=admin).status_code == 422
    assert client.get(f"{API}/capping/sheet", params={"from": "2025-01-01", "to": "2026-10-01"}, headers=admin).status_code == 422     # over a year
    assert client.get(f"{API}/capping/sheet", params={"storeId": str(uuid.uuid4())}, headers=admin).status_code == 404


def test_the_sheet_costs_the_same_number_of_queries_however_many_stores_there_are(client, admin, make_store, db):
    def queries() -> int:
        n = []
        listener = lambda *a, **k: n.append(1)  # noqa: E731
        event.listen(engine, "before_cursor_execute", listener)
        try:
            for url in ("/capping/sheet", "/capping/summary"):
                assert client.get(f"{API}{url}", params={"from": "2026-10-01", "to": "2026-10-31"}, headers=admin).status_code == 200
        finally:
            event.remove(engine, "before_cursor_execute", listener)
        return len(n)

    def add_stores(k):
        for _ in range(k):
            s = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")
            _seed(db, s, (date(2026, 10, 1), 10), (date(2026, 10, 15), 11))
            plant_entry(db, s, date(2026, 10, 2), 5)

    add_stores(2)
    few = queries()
    add_stores(8)
    assert queries() == few and few <= 20            # ten times the stores, the same queries


def test_the_csv_download_has_store_then_dd_mmm_dates(client, admin, zstore, db):
    _seed(db, zstore, (date(2026, 10, 1), 10), (date(2026, 10, 3), 11))
    r = client.get(f"{API}/capping/sheet/download", params={"from": "2026-10-01", "to": "2026-10-04", "storeId": str(zstore.id), "format": "csv"}, headers=admin)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('_2026-10-01_to_2026-10-04.csv"')
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["Store", "01-Oct", "02-Oct", "03-Oct", "04-Oct"]
    assert rows[1] == [zstore.name, "10", "10", "11", "11"] and len(rows) == 2


def test_the_excel_download_matches_and_the_format_is_checked(client, admin, zstore, db):
    _seed(db, zstore, (date(2026, 10, 2), 10))
    r = client.get(f"{API}/capping/sheet/download", params={"from": "2026-10-01", "to": "2026-10-03", "storeId": str(zstore.id), "format": "xlsx"}, headers=admin)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Store", "01-Oct", "02-Oct", "03-Oct"]
    assert [c.value for c in ws[2]] == [zstore.name, None, 10, 10] and ws.freeze_panes == "B2"
    assert client.get(f"{API}/capping/sheet/download", params={"format": "pdf"}, headers=admin).status_code == 422


def test_a_store_name_that_looks_like_a_formula_is_defused_in_the_file(client, admin, make_store, db):
    s = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru", name="=SUM(A1)")
    r = client.get(f"{API}/capping/sheet/download", params={"from": "2026-10-01", "to": "2026-10-01", "storeId": str(s.id)}, headers=admin)
    assert list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))[1][0] == "'=SUM(A1)"


# ---------------------------------------------------------------- the average-vs-limit summary

def _summary(client, headers, store, start="2026-10-01", end="2026-10-31"):
    r = client.get(f"{API}/capping/summary", params={"from": start, "to": end}, headers=headers)
    assert r.status_code == 200, r.text
    return r.json(), _sheet_row(r.json(), store)


def test_summary_average_is_over_the_days_that_have_an_entry_and_a_zero_counts(client, admin, zstore, db):
    for i, n in enumerate((10, 0, 8)):
        plant_entry(db, zstore, date(2026, 10, 1) + timedelta(days=i), n)
    plant_entry(db, zstore, date(2026, 9, 30), 200)             # outside the period
    body, row = _summary(client, admin, zstore)
    assert (row["average"], row["entry_days"], row["max_day"], row["total"]) == (6.0, 3, 10, 18)
    assert row["alert"] is False and row["breach_days"] == [] and row["status"] == "ok"
    assert body["limit"] == 11 and body["from"] == "2026-10-01" and body["totals"]["stores"] == len(body["rows"])


def test_summary_flags_the_average_above_the_limit_and_names_the_breach_days(client, admin, zstore, db):
    for i, n in enumerate((12, 13)):
        plant_entry(db, zstore, date(2026, 10, 1) + timedelta(days=i), n)
    _, row = _summary(client, admin, zstore)
    assert row["average"] == 12.5 and row["alert"] and row["average_breach"] and row["reasons"] == ["average", "day"]
    assert row["breach_days"] == [{"date": "2026-10-01", "value": 12}, {"date": "2026-10-02", "value": 13}]


def test_summary_flags_a_single_day_above_the_limit_even_when_the_average_is_fine(client, admin, zstore, db):
    for i, n in enumerate((5, 12, 6, 4)):
        plant_entry(db, zstore, date(2026, 10, 1) + timedelta(days=i), n)
    _, row = _summary(client, admin, zstore)
    assert row["average"] == 6.75 and row["alert"] and row["day_breach"] and not row["average_breach"] and row["reasons"] == ["day"]
    assert row["breach_days"] == [{"date": "2026-10-02", "value": 12}] and row["status"] == "over"


def test_summary_gives_both_averages_to_two_decimals_and_alerts_on_the_capping_average(client, admin, zstore, db):
    today = _today()
    start = today.replace(day=1)
    _seed(db, zstore, (start - timedelta(days=40), 12))                  # a capping of 12 in force the whole period
    for i, n in enumerate((10, 9, 9)):
        plant_entry(db, zstore, start + timedelta(days=i), n)         # entries average 9.33
    body = client.get(f"{API}/capping/summary", params={"from": start.isoformat(), "to": today.isoformat()}, headers=admin).json()
    row = _sheet_row(body, zstore)
    assert row["average"] == 9.33 and row["capping_average"] == 12.0 and row["capping_days"] == (today - start).days + 1
    assert row["alert"] and row["capping_breach"] and not row["average_breach"] and row["reasons"] == ["capping"] and row["status"] == "over"
    assert row["breach_days"] == []


def test_summary_capping_average_stops_at_today_for_a_period_that_runs_on(client, admin, zstore, db):
    today = _today()
    _seed(db, zstore, (today.replace(day=1) - timedelta(days=5), 10))
    end = today + timedelta(days=10)
    row = _sheet_row(client.get(f"{API}/capping/summary", params={"from": today.replace(day=1).isoformat(), "to": end.isoformat()}, headers=admin).json(), zstore)
    assert row["capping_days"] == today.day and row["capping_average"] == 10.0 and not row["alert"]


def test_summary_exactly_the_limit_is_not_a_breach(client, admin, zstore, db):
    for i in range(3):
        plant_entry(db, zstore, date(2026, 10, 1) + timedelta(days=i), 11)
    _, row = _summary(client, admin, zstore)
    assert row["average"] == 11.0 and row["max_day"] == 11
    assert row["alert"] is False and row["reasons"] == [] and row["breach_days"] == []
    assert row["status"] == "warn"                                        # amber: within 1 of the limit, not over


def test_summary_a_store_with_no_entries_has_no_average_and_no_alert(client, admin, zstore):
    _, row = _summary(client, admin, zstore)
    assert row["average"] is None and row["max_day"] is None and row["alert"] is False and row["status"] == "none"


def test_summary_gives_the_capping_in_force_today_and_lists_alerts_first(client, admin, zstore, make_store, db):
    calm = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")
    plant_entry(db, calm, date(2026, 10, 1), 3)
    plant_entry(db, zstore, date(2026, 10, 1), 15)
    _seed(db, zstore, (date(2026, 9, 1), 10), (_today() - timedelta(days=1), 11), (_today() - timedelta(days=30), 9))
    body, row = _summary(client, admin, zstore)
    assert row["current_capping"] == 11
    rows = [r["store_id"] for r in body["rows"]]
    assert rows.index(str(zstore.id)) < rows.index(str(calm.id))          # in alert before not
    assert body["totals"]["in_alert"] == sum(1 for r in body["rows"] if r["alert"])


def test_summary_period_is_checked_and_defaults_to_this_month(client, admin):
    body = client.get(f"{API}/capping/summary", headers=admin).json()
    assert body["from"] == _today().replace(day=1).isoformat()
    assert client.get(f"{API}/capping/summary", params={"from": "2026-10-09", "to": "2026-10-01"}, headers=admin).status_code == 422


def test_summary_alerts_on_the_collective_averages_not_on_a_single_store(client, admin, zstore, make_store, db, monkeypatch):
    other = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")
    start = _today().replace(day=1)
    for i, n in enumerate((20, 20)):
        plant_entry(db, zstore, start + timedelta(days=i), n)           # this store alone averages 20 ...
    for i, n in enumerate((2, 2, 2, 2)):
        plant_entry(db, other, start + timedelta(days=i), n)            # ... and this one 2
    params = {"from": start.isoformat(), "to": (start + timedelta(days=5)).isoformat()}

    def overall():
        body = client.get(f"{API}/capping/summary", params=params, headers=admin).json()
        rows = body["rows"]
        o = body["overall"]
        assert o["total"] == sum(r["total"] for r in rows) and o["entry_days"] == sum(r["entry_days"] for r in rows)
        assert o["entry_average"] == round(o["total"] / o["entry_days"], 2)                           # all bottles / all entry days, not an average of averages
        return o, _sheet_row(body, zstore)

    monkeypatch.setattr(settings, "MAX_DAILY_LIMIT", 11)
    o, store_row = overall()
    assert o["entry_average"] == round((40 + 8) / 6, 2) == 8.0                                        # 8.00 overall: under 11
    assert o["alert"] is False and o["reasons"] == []
    assert store_row["average"] == 20.0                                                               # the store keeps its own figure for reference
    monkeypatch.setattr(settings, "MAX_DAILY_LIMIT", 7)
    o, _ = overall()
    assert o["alert"] and o["entry_breach"] and o["reasons"] == ["average"] and o["status"] == "over"
    monkeypatch.setattr(settings, "MAX_DAILY_LIMIT", 8)
    assert overall()[0]["alert"] is False                                                             # exactly the limit is fine


# ---------------------------------------------------------------- zero capping on specific days

def _zero(client, headers, store, *days):
    return client.post(f"{API}/capping/zero-days", json={"storeId": str(store.id), "days": [d.isoformat() for d in days]}, headers=headers)


def test_with_zero_days_a_marked_day_counts_as_zero_in_the_capping_average():
    hist = [(date(2026, 10, 1), 10)]
    base = svc.capping_totals(hist, date(2026, 10, 1), date(2026, 10, 10))
    assert base == (100, 10)
    assert svc.with_zero_days(*base, hist, [date(2026, 10, 4), date(2026, 10, 5)], date(2026, 10, 1), date(2026, 10, 10)) == (80, 10)    # two days at 0: 8.0 on average
    assert svc.with_zero_days(*base, hist, [date(2026, 11, 1)], date(2026, 10, 1), date(2026, 10, 10)) == (100, 10)                    # outside the period: ignored
    assert svc.with_zero_days(0, 0, [], [date(2026, 10, 4)], date(2026, 10, 1), date(2026, 10, 10)) == (0, 1)                         # no capping that day: it still counts, as 0
    assert svc.with_zero_days(*base, hist, [date(2026, 10, 4), date(2026, 10, 4)], date(2026, 10, 1), date(2026, 10, 10)) == (90, 10)  # the same day twice is one day


def test_marking_zero_days_shows_on_the_sheet_and_lowers_the_capping_average(client, admin, zstore, db):
    today = _today()
    start = today.replace(day=1)
    _seed(db, zstore, (start - timedelta(days=3), 12))
    d1, d2 = start + timedelta(days=0), start + timedelta(days=1)
    r = _zero(client, admin, zstore, d2, d1)
    assert r.status_code == 201, r.text
    assert r.json()["marked"] == [d1.isoformat(), d2.isoformat()] and r.json()["already"] == []
    sheet = client.get(f"{API}/capping/sheet", params={"from": start.isoformat(), "to": (start + timedelta(days=3)).isoformat(), "storeId": str(zstore.id)}, headers=admin).json()
    assert _sheet_row(sheet, zstore)["cells"] == [0, 0, 12, 12]                         # the days either side keep their capping
    body = client.get(f"{API}/capping/summary", params={"from": start.isoformat(), "to": today.isoformat()}, headers=admin).json()
    row = _sheet_row(body, zstore)
    n = (today - start).days + 1
    expected = round((12 * (n - 2)) / n, 2) if n >= 2 else 0.0
    assert row["capping_zero_days"] == (2 if n >= 2 else 1) and row["capping_average"] == expected


def test_marking_the_same_day_again_is_reported_not_an_error_and_the_store_cap_is_untouched(client, admin, zstore, db):
    day = _today() - timedelta(days=2)
    assert _change(client, admin, zstore, 11, day - timedelta(days=5)).status_code == 201
    assert _zero(client, admin, zstore, day).json()["marked"] == [day.isoformat()]
    again = _zero(client, admin, zstore, day).json()
    assert again["marked"] == [] and again["already"] == [day.isoformat()]
    db.expire_all()
    assert db.query(StoreCap).filter_by(store_id=zstore.id).one().daily_cap == 11        # the store's own capping is not changed by a zero day
    assert db.query(CappingHistory).filter_by(store_id=zstore.id).count() == 1


def test_a_marked_day_is_listed_in_the_history_and_can_be_taken_back(client, admin, zstore):
    day = _today() - timedelta(days=1)
    _zero(client, admin, zstore, day)
    h = client.get(f"{API}/capping/history", params={"storeId": str(zstore.id)}, headers=admin).json()
    assert [z["day"] for z in h["zero_days"]] == [day.isoformat()] and h["zero_days"][0]["changed_by"]
    assert client.delete(f"{API}/capping/zero-days", params={"storeId": str(zstore.id), "day": day.isoformat()}, headers=admin).status_code == 200
    assert client.get(f"{API}/capping/history", params={"storeId": str(zstore.id)}, headers=admin).json()["zero_days"] == []
    assert client.delete(f"{API}/capping/zero-days", params={"storeId": str(zstore.id), "day": day.isoformat()}, headers=admin).status_code == 404


def test_zero_days_need_capping_change_and_sensible_dates(client, admin, zepto_emp, emp, zstore, make_store):
    day = _today()
    assert _zero(client, zepto_emp, zstore, day).status_code == 403
    assert _zero(client, emp, zstore, day).status_code == 403
    assert client.delete(f"{API}/capping/zero-days", params={"storeId": str(zstore.id), "day": day.isoformat()}, headers=zepto_emp).status_code == 403
    assert _zero(client, admin, zstore, _today() + timedelta(days=svc.ZERO_DAYS_AHEAD + 1)).status_code == 422
    assert _zero(client, admin, zstore, date(2020, 1, 1)).status_code == 422
    assert _zero(client, admin, zstore, _today() + timedelta(days=5)).status_code == 201                     # a planned day ahead is fine
    assert client.post(f"{API}/capping/zero-days", json={"storeId": str(zstore.id), "days": []}, headers=admin).status_code == 422
    too_many = [(_today() - timedelta(days=i)).isoformat() for i in range(svc.MAX_ZERO_DAYS + 1)]
    assert client.post(f"{API}/capping/zero-days", json={"storeId": str(zstore.id), "days": too_many}, headers=admin).status_code == 422
    assert client.post(f"{API}/capping/zero-days", json={"storeId": str(uuid.uuid4()), "days": [day.isoformat()]}, headers=admin).status_code == 404
    assert _zero(client, admin, make_store(platform="blinkit"), day).status_code == 404


def test_stores_with_no_capping_all_get_the_default_from_today_and_nobody_else_is_touched(client, admin, zepto_emp, zstore, make_store, db):
    from app.models.user import User

    have = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")
    zero = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")
    _seed(db, have, (_today() - timedelta(days=3), 12))
    db.add(StoreCap(store_id=zero.id, daily_cap=0, source="manual"))                       # a 0 is a capping that was set
    db.commit()
    before_caps = {r.store_id for r in db.query(StoreCap).all()}
    admin_id = db.query(User).filter_by(employee_code="ADMIN001").one().id
    try:
        assert client.post(f"{API}/capping/defaults", json={"dry_run": False}, headers=zepto_emp).status_code == 403     # only capping.change
        count = client.post(f"{API}/capping/defaults", json={}, headers=admin).json()                                    # a dry run by default
        assert count["dry_run"] is True and count["unset"] >= 1 and count["applied"] == 0 and count["capping"] == 10
        assert db.query(CappingHistory).filter_by(store_id=zstore.id).count() == 0                                       # counting changed nothing

        done = client.post(f"{API}/capping/defaults", json={"dry_run": False}, headers=admin)
        assert done.status_code == 200, done.text
        assert done.json()["applied"] == count["unset"]
        db.expire_all()
        row = db.query(CappingHistory).filter_by(store_id=zstore.id).one()
        assert row.capping == 10 and row.effective_from == _today().replace(day=1) and row.created_by == admin_id        # from day 1 of the month by default
        assert db.query(StoreCap).filter_by(store_id=zstore.id).one().daily_cap == 10
        assert [r.capping for r in db.query(CappingHistory).filter_by(store_id=have.id)] == [12]                         # a store that has one is left alone
        assert db.query(CappingHistory).filter_by(store_id=zero.id).count() == 0 and db.query(StoreCap).filter_by(store_id=zero.id).one().daily_cap == 0
        again = client.post(f"{API}/capping/defaults", json={"dry_run": False}, headers=admin).json()
        assert again["unset"] == 0 and again["applied"] == 0                                                             # running it twice does nothing
    finally:
        db.rollback()
        db.query(CappingHistory).filter(CappingHistory.created_by == admin_id, CappingHistory.effective_from == _today().replace(day=1)).delete()
        db.query(StoreCap).filter(StoreCap.store_id.notin_(before_caps)).delete(synchronize_session=False)
        db.commit()


def test_the_capping_print_is_one_page_with_every_store_and_its_capping(client, admin, zepto_emp, zstore, db):
    import fitz

    _seed(db, zstore, (_today() - timedelta(days=2), 9))
    r = client.get(f"{API}/capping/print", headers=zepto_emp)                                  # anyone who may see Zepto capping can print it
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and ".pdf" in r.headers["content-disposition"]
    pdf = fitz.open(stream=r.content, filetype="pdf")
    assert pdf.page_count == 1
    text = pdf[0].get_text()
    assert zstore.name[:20] in text and "store capping per day" in text and "9" in text
    assert client.get(f"{API}/capping/print").status_code == 401


def test_the_default_capping_can_start_on_any_past_date_but_not_in_the_future(client, admin, zstore, db):
    first = _today().replace(day=1)
    assert client.post(f"{API}/capping/defaults", json={"dry_run": True, "effective_from": (_today() + timedelta(days=1)).isoformat()}, headers=admin).status_code == 422
    assert client.post(f"{API}/capping/defaults", json={"dry_run": True}, headers=admin).json()["effective_from"] == first.isoformat()       # day 1 unless chosen


def test_a_start_date_can_be_moved_and_the_current_capping_follows(client, admin, zepto_emp, zstore, db):
    today = _today()
    _seed(db, zstore, (today - timedelta(days=9), 10), (today - timedelta(days=2), 12))
    ids = {r.capping: r.id for r in db.query(CappingHistory).filter_by(store_id=zstore.id)}

    moved = client.patch(f"{API}/capping/history/{ids[12]}", json={"effectiveFrom": (today - timedelta(days=5)).isoformat()}, headers=admin)
    assert moved.status_code == 200, moved.text
    assert moved.json()["effective_from"] == (today - timedelta(days=5)).isoformat() and moved.json()["capping"] == 12 and moved.json()["current_capping"] == 12
    db.expire_all()
    assert {r.capping: r.effective_from for r in db.query(CappingHistory).filter_by(store_id=zstore.id)} == {10: today - timedelta(days=9), 12: today - timedelta(days=5)}

    assert client.patch(f"{API}/capping/history/{ids[12]}", json={"effectiveFrom": (today - timedelta(days=9)).isoformat()}, headers=admin).status_code == 409   # that date is used
    assert client.patch(f"{API}/capping/history/{ids[12]}", json={"effectiveFrom": (today + timedelta(days=1)).isoformat()}, headers=admin).status_code == 422
    assert client.patch(f"{API}/capping/history/{ids[12]}", json={"effectiveFrom": today.isoformat()}, headers=zepto_emp).status_code == 403          # capping.change only
    assert client.patch(f"{API}/capping/history/{uuid.uuid4()}", json={"effectiveFrom": today.isoformat()}, headers=admin).status_code == 404
    assert client.patch(f"{API}/capping/history/{ids[12]}", json={"effectiveFrom": today.isoformat(), "capping": 99}, headers=admin).status_code == 422     # only the date can change


def _upload(client, headers, text, **form):
    return client.post(f"{API}/capping/import-dated", headers=headers, files={"file": ("capping.csv", text.encode(), "text/csv")}, data={k: str(v).lower() if isinstance(v, bool) else str(v) for k, v in form.items()})


def test_a_dated_sheet_finds_where_each_stores_capping_starts_and_nothing_is_saved_until_applied(client, admin, zepto_emp, make_store, db):
    from app.models.capping_history import CappingZeroDay

    today = _today()
    if today.day < 4:
        return                                                    # the sheet needs a few days of the month
    first = today.replace(day=1)
    days = [first + timedelta(days=i) for i in range(4)]
    a, b, c = (make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru") for _ in range(3))
    head = "Store," + ",".join(d.isoformat() for d in days)
    sheet = "\n".join([
        head,
        f"{a.name},10,10,10,10",                  # 10 from day 1
        f"{b.name},,,12,12",                      # starts on day 3
        f"{c.name},10,10,12,0",                   # 10 from day 1, 12 from day 3, and day 4 marked 0
        "Not A Store,5,5,5,5",
    ])
    assert _upload(client, zepto_emp, sheet).status_code == 403                                   # capping.change only

    look = _upload(client, admin, sheet)                                                          # a preview by default
    assert look.status_code == 200, look.text
    body = look.json()
    assert body["dry_run"] is True and body["applied"] is False and body["layout"] == "dates" and body["stores"] == 3 and body["unknown_stores"] == 1
    assert body["day_one"] == 2 and body["mid_month"] == 1 and body["new_rows"] == 4 and body["zero_days"] == 1
    assert {x["date"]: x["stores"] for x in body["starts"]} == {first.isoformat(): 2, days[2].isoformat(): 1}
    assert db.query(CappingHistory).filter(CappingHistory.store_id.in_([a.id, b.id, c.id])).count() == 0

    done = _upload(client, admin, sheet, dry_run=False)
    assert done.status_code == 200 and done.json()["applied"] is True, done.text
    db.expire_all()
    rows = lambda s: sorted((r.effective_from, r.capping) for r in db.query(CappingHistory).filter_by(store_id=s.id))  # noqa: E731
    assert rows(a) == [(first, 10)] and rows(b) == [(days[2], 12)] and rows(c) == [(first, 10), (days[2], 12)]
    assert [z.day for z in db.query(CappingZeroDay).filter_by(store_id=c.id)] == [days[3]]
    assert db.query(StoreCap).filter_by(store_id=a.id).one().daily_cap == 10                       # the current capping is read from the sheet
    assert db.query(StoreCap).filter_by(store_id=c.id).one().daily_cap == 12

    again = _upload(client, admin, sheet, dry_run=False).json()
    assert again["new_rows"] == 0 and again["zero_days"] == 0                                     # applying twice adds nothing


def test_a_dated_sheet_leaves_a_different_existing_capping_alone_unless_replace_is_on(client, admin, make_store, db):
    today = _today()
    if today.day < 4:
        return
    first = today.replace(day=1)
    days = [first + timedelta(days=i) for i in range(4)]
    s = make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru")
    _seed(db, s, (first, 9), (days[3], 15))                       # the CRM disagrees on day 1 and has a change on day 4 the sheet does not
    sheet = "Store," + ",".join(d.isoformat() for d in days) + f"\n{s.name},10,10,10,10"

    keep = _upload(client, admin, sheet, dry_run=False).json()
    assert keep["conflicts"] == 1 and keep["updated_rows"] == 0 and keep["removed_rows"] == 0
    db.expire_all()
    assert sorted((r.effective_from, r.capping) for r in db.query(CappingHistory).filter_by(store_id=s.id)) == [(first, 9), (days[3], 15)]

    replaced = _upload(client, admin, sheet, dry_run=False, replace=True).json()
    assert replaced["updated_rows"] == 1 and replaced["removed_rows"] == 1 and replaced["conflicts"] == 0
    db.expire_all()
    assert sorted((r.effective_from, r.capping) for r in db.query(CappingHistory).filter_by(store_id=s.id)) == [(first, 10)]
    assert db.query(StoreCap).filter_by(store_id=s.id).one().daily_cap == 10


def test_a_one_capping_sheet_starts_each_store_on_its_date_or_on_the_chosen_default(client, admin, make_store, db):
    today = _today()
    first = today.replace(day=1)
    mid = first + timedelta(days=min(2, today.day - 1))
    a, b = (make_store(platform="zepto", region=None, state="Karnataka", city="Bengaluru") for _ in range(2))
    sheet = f"Outlet Name,Daily Capping,Effective From\n{a.name},11,\n{b.name},13,{mid.isoformat()}\n"
    body = _upload(client, admin, sheet, dry_run=False).json()
    assert body["layout"] == "simple" and body["stores"] == 2 and body["applied"] is True
    db.expire_all()
    assert [(r.effective_from, r.capping) for r in db.query(CappingHistory).filter_by(store_id=a.id)] == [(first, 11)]      # no date: day 1 of the month
    assert [(r.effective_from, r.capping) for r in db.query(CappingHistory).filter_by(store_id=b.id)] == [(mid, 13)]
    assert _upload(client, admin, "Outlet Name,Capping\nx,1\n", start_date=(today + timedelta(days=1)).isoformat()).status_code == 422
    assert _upload(client, admin, "just,some,words\n1,2,3\n").status_code == 422
