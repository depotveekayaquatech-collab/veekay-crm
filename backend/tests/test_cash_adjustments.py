"""Developer cash-purchase adjustments + Sync Cash Purchase: adds to (never replaces) the real entry, ceil maths,
preview changes nothing, duplicate-sync protection, changed-record review, developer-only access."""
import io
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest

from app.models.cash_adjustment import CashAdjustment, CashSyncRecord, CashSyncRun
from app.models.cash_purchase import CashPurchase
from app.models.order_entry import OrderEntry
from app.models.user import User
from app.services.cash_adjustment_service import calc_quantity

from .conftest import API

BASE = f"{API}/developer/cash-adjustments"
TODAY = date.today()


@pytest.fixture(autouse=True)
def _wipe(db):
    yield
    for model in (CashAdjustment, CashSyncRecord, CashSyncRun, CashPurchase):
        db.query(model).delete()
    db.commit()


def _entry(db, store, count, day=TODAY, adjustment=0):
    db.add(OrderEntry(organization_id=store.organization_id, store_id=store.id, order_date=day, bottle_count=count, cash_adjustment=adjustment))
    db.commit()


def _total(db, store, day=TODAY):
    db.expire_all()
    e = db.query(OrderEntry).filter_by(store_id=store.id, order_date=day).first()
    return (e.bottle_count, e.cash_adjustment) if e else None


def _body(stores, **kw):
    return {"store_ids": [str(s.id) for s in stores], "purchase_date": TODAY.isoformat(), "amount": "100", "price_per_item": "30", **kw}


def _purchase(db, store, amount, day=TODAY, user=None):
    u = user or db.query(User).filter_by(employee_code="EMP001").one()
    p = CashPurchase(organization_id=store.organization_id, kind="STORE", store_id=store.id, purchase_date=day, category="SUPPLY_DELAYED",
                     amount=Decimal(str(amount)), created_by_user_id=u.id)
    db.add(p)
    db.commit()
    return p


def test_quantity_is_always_rounded_up():
    for amount, price, qty in ((100, 30, 4), (100, 25, 4), (100, 40, 3), (500, 30, 17), (1000, 30, 34)):
        assert calc_quantity(Decimal(amount), Decimal(price)) == qty


def test_only_the_developer_can_reach_any_of_it(client, admin, emp, dev, make_store):
    s = make_store(code="CA-X1")
    calls = [("get", f"{BASE}/stores", None), ("post", f"{BASE}/preview", _body([s])), ("post", f"{BASE}/apply", _body([s])),
             ("get", f"{BASE}/history", None), ("get", f"{BASE}/history/export", None), ("get", f"{BASE}/purchase-sheet", None),
             ("post", f"{BASE}/sync", None), ("post", f"{BASE}/sync/standard-rate", {"price_per_item": "30"}), ("get", f"{BASE}/sync/records", None), ("get", f"{BASE}/sync/runs", None),
             ("post", f"{BASE}/sync/preview", {"record_ids": [str(s.id)]}), ("post", f"{BASE}/sync/apply", {"record_ids": [str(s.id)]})]
    for method, url, body in calls:
        for who in (admin, emp):
            r = client.request(method, url, headers=who, **({"json": body} if body else {}))
            assert r.status_code == 403, (method, url)
        assert client.request(method, url, headers=dev, **({"json": body} if body else {})).status_code != 403, url


def test_cash_is_added_to_the_employees_entry_for_several_stores(client, dev, db, make_store):
    a, b, c = (make_store(code=f"CA-M{i}") for i in range(3))
    for store, n in ((a, 10), (b, 8), (c, 15)):
        _entry(db, store, n)
    r = client.post(f"{BASE}/preview", headers=dev, json=_body([a, b, c]))
    assert r.status_code == 200, r.text
    pv = r.json()
    assert pv["auto_qty"] == 4 and pv["final_qty"] == 4 and pv["can_apply"] is True
    assert [(x["previous"], x["new_total"]) for x in pv["rows"]] == [(10, 14), (8, 12), (15, 19)]
    assert [_total(db, s) for s in (a, b, c)] == [(10, 0), (8, 0), (15, 0)]          # a preview changes nothing

    done = client.post(f"{BASE}/apply", headers=dev, json=_body([a, b, c]))
    assert done.status_code == 200 and not done.json()["failed"]
    assert [_total(db, s) for s in (a, b, c)] == [(14, 4), (12, 4), (19, 4)]          # added, not replaced
    hist = client.get(f"{BASE}/history", headers=dev).json()
    assert hist["total"] == 3 and {(h["previous_qty"], h["final_order_qty"]) for h in hist["items"]} == {(10, 14), (8, 12), (15, 19)}


def test_the_developer_can_override_the_quantity_and_it_is_recorded(client, dev, db, make_store):
    s = make_store(code="CA-O1")
    _entry(db, s, 10)
    assert client.post(f"{BASE}/apply", headers=dev, json=_body([s], quantity=5)).status_code == 200
    assert _total(db, s) == (15, 5)
    row = client.get(f"{BASE}/history", headers=dev).json()["items"][0]
    assert row["auto_qty"] == 4 and row["final_qty"] == 5 and row["source"] == "Manual"


def test_a_day_with_no_employee_entry_gets_one_and_the_employee_can_still_mark(client, dev, emp, db, make_store):
    s = make_store(code="CA-N1")
    assert client.post(f"{BASE}/apply", headers=dev, json=_body([s])).status_code == 200
    assert _total(db, s) == (4, 4)
    r = client.post(f"{API}/orders/mark", headers=emp, json={"store_id": str(s.id), "order_date": TODAY.isoformat(), "bottle_count": 10})
    assert r.status_code == 200, r.text
    assert _total(db, s) == (14, 4)                                                    # their 10 + the developer's 4
    again = client.post(f"{API}/orders/mark", headers=emp, json={"store_id": str(s.id), "order_date": TODAY.isoformat(), "bottle_count": 3})
    assert again.status_code == 409                                                    # still one-shot


def test_bad_input_changes_nothing(client, dev, db, make_store):
    s = make_store(code="CA-V1")
    _entry(db, s, 190)
    for kw in ({"amount": "0"}, {"amount": "-5"}, {"price_per_item": "0"}, {"amount": "abc"}, {"quantity": -1},
               {"purchase_date": date(TODAY.year + 1, 1, 1).isoformat()}, {"store_ids": []}):
        assert client.post(f"{BASE}/apply", headers=dev, json={**_body([s]), **kw}).status_code == 422, kw
    assert client.post(f"{BASE}/apply", headers=dev, json=_body([s], quantity=0)).status_code == 422
    over = client.post(f"{BASE}/apply", headers=dev, json=_body([s], quantity=20))      # 190 + 20 > 200
    assert over.status_code == 422 and _total(db, s) == (190, 0)
    assert client.get(f"{BASE}/history", headers=dev).json()["total"] == 0


def test_history_sheet_download(client, dev, db, make_store):
    from openpyxl import load_workbook

    s = make_store(code="CA-H1", name="History Store")
    _entry(db, s, 10)
    client.post(f"{BASE}/apply", headers=dev, json=_body([s]))
    rows = list(load_workbook(io.BytesIO(client.get(f"{BASE}/history/export", headers=dev).content)).active.iter_rows(values_only=True))
    assert rows[0][:2] == ("Sync ID", "Purchase ID") and rows[0][-2:] == ("Source", "Action")
    assert rows[1][4:6] == ("CA-H1", "History Store") and rows[1][9:13] == (4, 4, 10, 14) and rows[1][13:] == ("Manual", "Applied")


# ----------------------------------------------------------------------- sync

def test_sync_brings_purchases_in_without_touching_the_order_sheet(client, dev, db, make_store):
    s = make_store(code="CA-S1")
    _entry(db, s, 10)
    _purchase(db, s, 100)
    r = client.post(f"{BASE}/sync", headers=dev)
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["records_found"], body["new_records"], body["changed_records"], body["already_processed"]) == (1, 1, 0, 0)
    assert _total(db, s) == (10, 0)

    recs = client.get(f"{BASE}/sync/records", headers=dev).json()
    assert recs["counts"]["PENDING"] == 1 and recs["items"][0]["price_per_item"] is None and recs["items"][0]["final_qty"] is None

    # syncing again never creates a second working row
    again = client.post(f"{BASE}/sync", headers=dev).json()
    assert again["new_records"] == 0 and db.query(CashSyncRecord).count() == 1


def test_review_edit_preview_apply_and_never_twice(client, dev, db, make_store):
    s = make_store(code="CA-S2")
    _entry(db, s, 10)
    _purchase(db, s, 100)
    client.post(f"{BASE}/sync", headers=dev)
    rec = client.get(f"{BASE}/sync/records", headers=dev).json()["items"][0]
    rid = rec["id"]

    assert client.post(f"{BASE}/sync/preview", headers=dev, json={"record_ids": [rid]}).json()["can_apply"] is False   # no price yet
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]}).status_code == 422

    edited = client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"price_per_item": "30"}).json()
    assert edited["auto_qty"] == 4 and edited["final_qty"] == 4
    assert client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"amount": "120"}).json()["auto_qty"] == 4      # 120 / 30
    assert client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"amount": "110"}).json()["auto_qty"] == 4      # 3.67 -> 4
    assert client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"price_per_item": "25"}).json()["auto_qty"] == 5  # 110 / 25 = 4.4 -> 5
    assert client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"final_qty": 6}).json()["final_qty"] == 6

    pv = client.post(f"{BASE}/sync/preview", headers=dev, json={"record_ids": [rid]}).json()
    assert pv["can_apply"] and pv["rows"][0]["previous"] == 10 and pv["rows"][0]["new_total"] == 16 and pv["total_qty"] == 6
    assert _total(db, s) == (10, 0)

    ok = client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]})
    assert ok.status_code == 200 and ok.json()["applied_records"] == 1
    assert _total(db, s) == (16, 6)
    h = client.get(f"{BASE}/history", headers=dev).json()["items"][0]
    assert h["source"] == "Store Purchase Sheet" and h["auto_qty"] == 5 and h["final_qty"] == 6 and h["purchase_code"].startswith("CP-")

    # applying again, or syncing again, must not add it a second time
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]}).status_code == 422
    third = client.post(f"{BASE}/sync", headers=dev).json()
    assert third["new_records"] == 0 and third["already_processed"] == 1
    assert _total(db, s) == (16, 6)
    assert client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"price_per_item": "10"}).status_code == 409


def test_a_standard_rate_fixes_every_pending_purchase(client, dev, db, make_store):
    s = make_store(code="CA-R1")
    for amount in (100, 250, 75):
        _purchase(db, s, amount)
    client.post(f"{BASE}/sync", headers=dev)
    items = client.get(f"{BASE}/sync/records", headers=dev).json()["items"]
    client.patch(f"{BASE}/sync/records/{items[0]['id']}", headers=dev, json={"price_per_item": "20"})   # one already priced

    kept = client.post(f"{BASE}/sync/standard-rate", headers=dev, json={"price_per_item": "30"}).json()
    assert kept == {"updated": 2, "kept": 1}
    prices = sorted(r["price_per_item"] for r in client.get(f"{BASE}/sync/records", headers=dev).json()["items"])
    assert prices == ["20.00", "30.00", "30.00"]

    every = client.post(f"{BASE}/sync/standard-rate", headers=dev, json={"price_per_item": "25", "overwrite": True}).json()
    assert every == {"updated": 3, "kept": 0}
    recs = client.get(f"{BASE}/sync/records", headers=dev).json()["items"]
    assert {r["price_per_item"] for r in recs} == {"25.00"} and sorted(r["auto_qty"] for r in recs) == [3, 4, 10]
    assert client.post(f"{BASE}/sync/standard-rate", headers=dev, json={"price_per_item": "0"}).status_code == 422


def test_applied_purchases_keep_their_price_when_a_standard_rate_is_set(client, dev, db, make_store):
    s = make_store(code="CA-R2")
    _entry(db, s, 10)
    _purchase(db, s, 100)
    client.post(f"{BASE}/sync", headers=dev)
    client.post(f"{BASE}/sync/standard-rate", headers=dev, json={"price_per_item": "30"})
    rid = client.get(f"{BASE}/sync/records", headers=dev).json()["items"][0]["id"]
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]}).status_code == 200
    again = client.post(f"{BASE}/sync/standard-rate", headers=dev, json={"price_per_item": "50", "overwrite": True}).json()
    assert again["updated"] == 0
    assert client.get(f"{BASE}/sync/records?status=APPLIED", headers=dev).json()["items"][0]["price_per_item"] == "30.00"


def test_several_purchases_for_one_store_and_day_are_summed(client, dev, db, make_store):
    s = make_store(code="CA-S3")
    _entry(db, s, 10)
    for amount in (100, 90, 60):                       # at 25 each: 4 + 4 (3.6 up) + 3 (2.4 up) = 11
        _purchase(db, s, amount)
    client.post(f"{BASE}/sync", headers=dev)
    ids = [r["id"] for r in client.get(f"{BASE}/sync/records", headers=dev).json()["items"]]
    for i in ids:
        client.patch(f"{BASE}/sync/records/{i}", headers=dev, json={"price_per_item": "25"})
    pv = client.post(f"{BASE}/sync/preview", headers=dev, json={"record_ids": ids}).json()
    assert len(pv["rows"]) == 1 and pv["rows"][0]["added"] == 11 and pv["rows"][0]["new_total"] == 21 and pv["total_records"] == 3
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": ids}).status_code == 200
    assert _total(db, s) == (21, 11)
    hist = client.get(f"{BASE}/history", headers=dev).json()["items"]
    assert len(hist) == 3 and max(h["final_order_qty"] for h in hist) == 21 and min(h["previous_qty"] for h in hist) == 10
    by_prev = {h["previous_qty"]: h["final_order_qty"] for h in hist}                         # one unbroken chain 10 -> ... -> 21
    steps, cur = 0, 10
    while cur in by_prev:
        cur, steps = by_prev[cur], steps + 1
    assert cur == 21 and steps == 3


def test_the_developer_can_move_a_purchase_to_another_store_and_day(client, dev, db, make_store):
    a, b = make_store(code="CA-S4"), make_store(code="CA-S5")
    yesterday = date.fromordinal(TODAY.toordinal() - 1)
    _entry(db, a, 10)
    _entry(db, b, 7, day=yesterday)
    _purchase(db, a, 100)
    client.post(f"{BASE}/sync", headers=dev)
    rid = client.get(f"{BASE}/sync/records", headers=dev).json()["items"][0]["id"]
    client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"store_id": str(b.id), "purchase_date": yesterday.isoformat(), "price_per_item": "30"})
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]}).status_code == 200
    assert _total(db, a) == (10, 0) and _total(db, b, yesterday) == (11, 4)


def test_a_changed_source_purchase_is_flagged_not_applied(client, dev, db, make_store):
    s = make_store(code="CA-S6")
    _entry(db, s, 10)
    p = _purchase(db, s, 100)
    client.post(f"{BASE}/sync", headers=dev)
    rid = client.get(f"{BASE}/sync/records", headers=dev).json()["items"][0]["id"]
    client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"price_per_item": "30"})
    client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]})
    assert _total(db, s) == (14, 4)

    p.amount = Decimal("120")
    db.commit()
    summary = client.post(f"{BASE}/sync", headers=dev).json()
    assert summary["changed_records"] == 1 and summary["new_records"] == 0
    rec = client.get(f"{BASE}/sync/records?status=CHANGED", headers=dev).json()["items"][0]
    assert rec["previous_source"]["amount"] == "100.00" and rec["current_source"]["amount"] == "120.00"
    assert _total(db, s) == (14, 4)                                                   # nothing added automatically
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]}).status_code == 422

    assert client.post(f"{BASE}/sync/records/{rid}/resolve", headers=dev, json={"action": "acknowledge"}).json()["status"] == "APPLIED"
    assert _total(db, s) == (14, 4)
    assert client.post(f"{BASE}/sync", headers=dev).json()["changed_records"] == 0

    # a deliberate extra adjustment instead: reopen, give it a quantity, apply -> adds on top
    p.amount = Decimal("150")
    db.commit()
    client.post(f"{BASE}/sync", headers=dev)
    reopened = client.post(f"{BASE}/sync/records/{rid}/resolve", headers=dev, json={"action": "reopen"}).json()
    assert reopened["status"] == "PENDING" and reopened["amount"] == "150.00" and reopened["final_qty"] is not None
    client.patch(f"{BASE}/sync/records/{rid}", headers=dev, json={"final_qty": 2})
    assert client.post(f"{BASE}/sync/apply", headers=dev, json={"record_ids": [rid]}).status_code == 200
    assert _total(db, s) == (16, 6)


def test_sync_history_lists_every_run(client, dev, db, make_store):
    s = make_store(code="CA-S7")
    _purchase(db, s, 100)
    client.post(f"{BASE}/sync", headers=dev)
    client.post(f"{BASE}/sync", headers=dev)
    runs = client.get(f"{BASE}/sync/runs", headers=dev).json()
    assert len(runs) == 2 and runs[0]["records_found"] == 1 and {r["new_records"] for r in runs} == {0, 1}


def test_the_purchase_sheet_download(client, dev, db, make_store):
    from openpyxl import load_workbook

    s = make_store(code="CA-D1")
    _purchase(db, s, 100)
    r = client.get(f"{BASE}/purchase-sheet", headers=dev)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))[0][1] == "Type"


# -------------------------------------------------------------------- analysis

def _vendor_store(make_store, code, vendor, name):
    s = make_store(code=code, name=name)
    s.vendor_name = vendor
    return s


def test_vendor_analysis_ranks_flags_and_lists_store_names(client, admin, dev, db, make_store):
    from openpyxl import load_workbook

    a1 = _vendor_store(make_store, "AN-A1", "Alpha Aqua", "Alpha Store One")
    a2 = _vendor_store(make_store, "AN-A2", "alpha  aqua", "Alpha Store Two")      # same vendor, spelled differently
    b1 = _vendor_store(make_store, "AN-B1", "Beta Water", "Beta Store One")
    c1 = _vendor_store(make_store, "AN-C1", "Gamma", "Gamma Store One")
    d1 = _vendor_store(make_store, "AN-D1", "Delta", "Delta Store One")
    db.commit()
    for s, n in ((a1, 4), (a2, 2), (b1, 2), (c1, 1), (d1, 1)):
        for _ in range(n):
            _purchase(db, s, 100)
    _purchase(db, a1, 999).kind = "OFFICE"        # an office purchase must not count
    db.commit()

    r = client.get(f"{BASE}/analysis", headers=dev)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_purchases"] == 10 and body["vendors_count"] == 4 and body["average_purchases"] == 2.5
    top, second, *rest = body["vendors"]
    assert (top["rank"], top["vendor"].lower(), top["purchases"], top["flag"]) == (1, "alpha aqua", 6, "HIGH")
    assert {s["outlet_name"] for s in top["stores"]} == {"Alpha Store One", "Alpha Store Two"} and top["stores"][0]["purchases"] == 4
    assert (second["vendor"], second["purchases"], second["flag"]) == ("Beta Water", 2, "")      # not above the 2.5 average
    assert [v["flag"] for v in rest] == ["", ""]

    xl = client.get(f"{BASE}/analysis/export", headers=dev)
    assert xl.status_code == 200 and "attachment" in xl.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames == ["Vendor summary", "Stores by vendor", "Notes"]
    summary = list(wb["Vendor summary"].iter_rows(values_only=True))
    assert summary[0][:3] == ("Rank", "Flag", "Vendor") and summary[1][1] == "HIGH" and "Alpha Store One (4)" in summary[1][11]
    stores = list(wb["Stores by vendor"].iter_rows(values_only=True))
    assert {"Alpha Store One", "Beta Store One"} <= {row[3] for row in stores[1:]}
    assert client.get(f"{BASE}/analysis?start={TODAY.isoformat()}&end={TODAY.isoformat()}", headers=dev).json()["total_purchases"] == 10
    assert client.get(f"{BASE}/analysis?end=2000-01-01", headers=dev).json()["vendors_count"] == 0
    for who in (admin,):
        assert client.get(f"{BASE}/analysis", headers=who).status_code == 403
        assert client.get(f"{BASE}/analysis/export", headers=who).status_code == 403
