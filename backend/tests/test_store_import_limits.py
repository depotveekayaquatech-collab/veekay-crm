"""Store import: bounded memory / CPU whatever the file contains, batched writes, and an import that never blocks the event loop."""
import inspect
import io
import time
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.api.v1.endpoints import stores as stores_endpoint
from app.models.store import Store
from app.services import store_sync_service as svc
from tests.conftest import API

HEADER = ["Outlet ID", "Outlet Name", "STATUS", "State", "City"]


def _xlsx(rows, *, fill_cells=()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    fill = PatternFill("solid", fgColor="FFFF00")
    for row, col in fill_cells:
        ws.cell(row=row, column=col).fill = fill
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(n, prefix="LIM"):
    return [[f"{prefix}-{i}", f"Blinkit - {prefix} {i}", "Live", "Delhi", "Delhi"] for i in range(n)]


def test_normal_sheet_parses_unchanged():
    rows = svc._rows_from_xlsx(_xlsx(_rows(50)))
    assert len(rows) == 50 and rows[0]["Outlet ID"] == "LIM-0" and rows[49]["City"] == "Delhi"


def test_a_sheet_formatted_to_the_last_row_and_column_is_read_quickly_and_small():
    # a formatted cell at XFD / row 1,048,576 makes the declared sheet size enormous
    data = _xlsx(_rows(200), fill_cells=[(1048576, 16384), (1, 16384)])
    t = time.time()
    rows = svc._rows_from_xlsx(data)
    assert len(rows) == 200
    assert all(len(r) <= len(HEADER) for r in rows)      # only the header's width is ever read
    assert time.time() - t < 20                           # stops after the blank run instead of walking a million rows


def test_too_many_rows_is_refused_with_a_clear_message(monkeypatch):
    monkeypatch.setattr(svc, "MAX_ROWS", 100)
    with pytest.raises(svc.SyncError, match="more than 100 rows"):
        svc._rows_from_xlsx(_xlsx(_rows(150)))


def test_too_many_columns_is_refused(monkeypatch):
    monkeypatch.setattr(svc, "MAX_COLUMNS", 3)
    with pytest.raises(svc.SyncError, match="more than 3 columns"):
        svc._rows_from_xlsx(_xlsx(_rows(3)))


def test_a_workbook_that_unpacks_hugely_is_refused_before_it_is_read(monkeypatch):
    monkeypatch.setattr(svc, "MAX_UNPACKED_BYTES", 10_000)
    with pytest.raises(svc.SyncError, match="too large to import safely"):
        svc._rows_from_xlsx(_xlsx(_rows(2000)))


def test_not_a_zip_is_a_clean_error():
    with pytest.raises(svc.SyncError, match="could not be read"):
        svc._rows_from_xlsx(b"this is not an excel file")


def test_csv_is_bounded_too(monkeypatch):
    monkeypatch.setattr(svc, "MAX_ROWS", 10)
    body = "Outlet ID,Outlet Name\n" + "\n".join(f"C{i},Name {i}" for i in range(30))
    with pytest.raises(svc.SyncError, match="more than 10 rows"):
        svc._parse_csv_text(body)
    assert len(svc._parse_csv_text("Outlet ID,Outlet Name\nA1,One\n,\nA2,Two\n")) == 2


def test_import_writes_in_batches_and_is_idempotent(client, dev, db, monkeypatch):
    monkeypatch.setattr(svc, "STORE_FLUSH_EVERY", 7)      # force several flushes for a small file
    data = _xlsx(_rows(30, "BATCH"))
    r = client.post(f"{API}/stores/import", data={"platform": "blinkit"}, files={"file": ("a.xlsx", data, "application/octet-stream")}, headers=dev)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 30
    again = client.post(f"{API}/stores/import", data={"platform": "blinkit"}, files={"file": ("a.xlsx", data, "application/octet-stream")}, headers=dev)
    assert again.json()["created"] == 0 and again.json()["unchanged"] == 30     # a retry never duplicates or partially repeats
    assert db.query(Store).filter(Store.external_code.like("BATCH-%")).count() == 30
    db.query(Store).filter(Store.external_code.like("BATCH-%")).delete()
    db.commit()


def test_oversized_upload_is_refused_cleanly(client, dev):
    big = b"x" * (svc.MAX_UPLOAD_BYTES + 10)
    r = client.post(f"{API}/stores/import", data={"platform": "blinkit"}, files={"file": ("a.xlsx", big, "application/octet-stream")}, headers=dev)
    assert r.status_code == 422 and "too large" in r.json()["detail"]


def test_the_import_endpoint_does_not_block_the_event_loop():
    assert not inspect.iscoroutinefunction(stores_endpoint.import_stores)
