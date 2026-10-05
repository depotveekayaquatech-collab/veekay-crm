"""Tickets (routing, workflow, insights), partner accounts (hard limits) and admin management."""
import uuid

import pytest

from app.models.refresh_session import RefreshSession
from app.models.ticket import Ticket
from app.models.user import User, UserRole
from app.models.user_permission import UserPermission

from .conftest import API, DEMO_PASSWORD, auth, login

NEW_PW = "Brand#New1234"


@pytest.fixture()
def made(db):
    """Users created through the API in a test; removed afterwards."""
    codes: list[str] = []
    yield codes
    for code in codes:
        u = db.query(User).filter(User.employee_code == code).one_or_none()
        if u is None:
            continue
        for model in (RefreshSession, UserPermission, UserRole):
            db.query(model).filter(model.user_id == u.id).delete()
        db.query(Ticket).filter(Ticket.created_by_user_id == u.id).delete()
        db.query(User).filter(User.id == u.id).delete()
    db.commit()


def _platform(client, admin, slug):
    return next(p["id"] for p in client.get(f"{API}/partners", headers=admin).json() if p["slug"] == slug)


def _create(client, admin, made, **body):
    code = f"T{uuid.uuid4().hex[:8].upper()}"
    made.append(code)
    payload = {"employee_code": code, "full_name": f"Test {code}", "password": "Temp#Pass1234", **body}
    r = client.post(f"{API}/employees", headers=admin, json=payload)
    return code, r


def _sign_in(client, code):
    """Log in with the temporary password, replace it (required), and return headers for the fresh session."""
    t = login(client, code, "Temp#Pass1234").json()
    r = client.post(f"{API}/auth/change-password", headers=auth(t["access_token"]), json={"current_password": "Temp#Pass1234", "new_password": NEW_PW})
    assert r.status_code == 200, r.text
    return auth(r.json()["access_token"])


@pytest.fixture()
def partner(client, admin, made):
    code, r = _create(client, admin, made, account_type="partner", platform_id=_platform(client, admin, "blinkit"),
                      permission_codes=["orders.view", "tickets.view", "tickets.create"])
    assert r.status_code == 201, r.text
    return _sign_in(client, code)


def _store_of(client, headers):
    return client.get(f"{API}/tickets/stores", headers=headers).json()[0]


# ---------------------------------------------------------------- partner accounts
def test_partner_can_only_hold_partner_safe_permissions(client, admin, made):
    bid = _platform(client, admin, "blinkit")
    _, r = _create(client, admin, made, account_type="partner", platform_id=bid, permission_codes=["orders.overview"])
    assert r.status_code == 422 and "orders.overview" in r.json()["detail"]
    _, r = _create(client, admin, made, account_type="partner", permission_codes=["orders.view"])      # a platform is required
    assert r.status_code == 422


def test_partner_is_confined_to_its_own_platform_and_to_what_it_was_given(client, admin, partner):
    for path in ("/orders/insights", "/orders/daily-overview?partner=zepto", "/stores", "/employees", "/regions", "/tickets/analytics"):
        assert client.get(f"{API}{path}", headers=partner).status_code == 403, path
    inv = client.get(f"{API}/orders/inventory", headers=partner).json()
    rows = next(v for v in inv.values() if isinstance(v, list))
    assert rows and {r["platform_slug"] for r in rows} == {"blinkit"}
    stores = client.get(f"{API}/orders/my-stores", headers=partner).json()
    assert stores and {s["partner_slug"] for s in stores} == {"blinkit"}


def test_partner_cannot_mark_orders_or_check_in(client, partner):
    sid = client.get(f"{API}/orders/my-stores", headers=partner).json()[0]["id"]
    assert client.post(f"{API}/orders/mark", headers=partner, json={"store_id": sid, "order_date": "2026-01-01", "bottle_count": 5}).status_code == 403
    assert client.post(f"{API}/attendance/check-in", headers=partner, json={}).status_code == 403


def test_partner_home_overview(client, partner, admin):
    o = client.get(f"{API}/partner/overview", headers=partner)
    assert o.status_code == 200
    body = o.json()
    assert body["platform"]["slug"] == "blinkit" and body["entries"]["stores"]["live"] >= 1 and "tickets" in body
    assert client.get(f"{API}/partner/overview", headers=admin).status_code == 403                  # only partner logins


# ---------------------------------------------------------------- tickets
def test_ticket_visibility_follows_region_and_platform(client, admin, emp, partner):
    store = client.get(f"{API}/orders/my-stores", headers=emp).json()[0]          # a North / Blinkit store EMP001 handles
    t = client.post(f"{API}/tickets", headers=partner, json={"store_id": store["id"], "category": "LATE_DELIVERY", "priority": "HIGH", "title": "Delivery came 3 hours late"})
    assert t.status_code == 201, t.text
    tid = t.json()["id"]
    assert tid in {x["id"] for x in client.get(f"{API}/tickets", headers=emp).json()["items"]}            # the region's employee receives it
    assert tid in {x["id"] for x in client.get(f"{API}/tickets", headers=admin).json()["items"]}          # admin sees every region
    # a Zepto employee (a different platform and state scope) never sees it
    from .conftest import login as _login
    z = auth(_login(client, "EMP002").json()["access_token"])
    assert tid not in {x["id"] for x in client.get(f"{API}/tickets", headers=z).json()["items"]}
    assert client.get(f"{API}/tickets/{tid}", headers=z).status_code == 404


def test_partner_cannot_raise_a_ticket_for_another_platform(client, admin, partner):
    zep = next(s for s in client.get(f"{API}/tickets/stores?q=ZEP", headers=admin).json() if s["platform_slug"] == "zepto")
    r = client.post(f"{API}/tickets", headers=partner, json={"store_id": zep["id"], "category": "OTHER", "title": "Not my store"})
    assert r.status_code == 403


def test_ticket_workflow_roles(client, admin, emp, partner):
    store = client.get(f"{API}/orders/my-stores", headers=emp).json()[0]
    tid = client.post(f"{API}/tickets", headers=partner, json={"store_id": store["id"], "category": "NO_DELIVERY", "priority": "URGENT", "title": "Nothing arrived today"}).json()["id"]
    patch = lambda h, **b: client.patch(f"{API}/tickets/{tid}", headers=h, json=b)
    assert patch(emp, priority="LOW").status_code == 403                       # only admins set priority / assign
    assert patch(emp, status="CLOSED").status_code == 403                      # employees can't close
    assert patch(partner, status="IN_PROGRESS").status_code == 403             # partners can't do the field work
    assert patch(emp, status="IN_PROGRESS").status_code == 200
    assert client.post(f"{API}/tickets/{tid}/comments", headers=emp, json={"body": "Driver is on the way"}).status_code == 200
    assert client.post(f"{API}/tickets/{tid}/comments", headers=partner, json={"body": "Thanks"}).status_code == 200
    assert patch(emp, status="RESOLVED").status_code == 200
    closed = patch(partner, status="CLOSED")
    assert closed.status_code == 200 and closed.json()["status"] == "CLOSED"
    thread = [c["event"] or c["body"] for c in closed.json()["comments"]]
    assert any("In Progress" in x for x in thread if x) and "Thanks" in thread
    reopened = patch(admin, status="OPEN", priority="HIGH")
    assert reopened.json()["status"] == "OPEN" and reopened.json()["priority"] == "HIGH" and reopened.json()["resolved_at"] is None


def test_insights_are_admin_only_and_flag_delivery_problems(client, admin, emp, partner):
    store = client.get(f"{API}/orders/my-stores", headers=emp).json()[0]
    client.post(f"{API}/tickets", headers=partner, json={"store_id": store["id"], "category": "LATE_DELIVERY", "title": "Late again"})
    assert client.get(f"{API}/tickets/analytics", headers=emp).status_code == 403
    a = client.get(f"{API}/tickets/analytics", headers=admin).json()
    assert a["summary"]["delivery_active"] >= 1 and a["by_region"] and a["by_vendor"] is not None
    assert any(h["store_id"] == store["id"] for h in a["hotspots"])
    assert {"label", "total", "active", "delivery", "overdue", "severity"} <= set(a["by_region"][0])


# ---------------------------------------------------------------- admin management
def test_admin_can_create_a_full_admin_and_a_custom_admin(client, admin, made):
    code, r = _create(client, admin, made, account_type="admin", admin_access="full")
    assert r.status_code == 201 and r.json()["category"] == "admin" and r.json()["admin_level"] == "full" and r.json()["roles"] == ["admin"]
    code2, r2 = _create(client, admin, made, account_type="admin", admin_access="custom", permission_codes=["orders.view", "tickets.manage", "employees.view"])
    assert r2.status_code == 201 and r2.json()["admin_level"] == "custom" and set(r2.json()["direct_permissions"]) == {"orders.view", "tickets.manage", "employees.view"}
    me = client.get(f"{API}/auth/me", headers=_sign_in(client, code2)).json()
    assert me["is_admin"] is False and "tickets.manage" in me["permissions"] and "stores.manage" not in me["permissions"]


def test_custom_admin_cannot_create_admins_or_hand_out_what_they_lack(client, admin, made):
    code, r = _create(client, admin, made, account_type="admin", admin_access="custom", permission_codes=["orders.view", "employees.view", "employees.create"])
    assert r.status_code == 201
    custom = _sign_in(client, code)
    mk = lambda **b: client.post(f"{API}/employees", headers=custom, json={"employee_code": f"T{uuid.uuid4().hex[:8].upper()}", "full_name": "x", "password": "Temp#Pass1234", **b})
    assert mk(account_type="admin", admin_access="full").status_code == 403
    over = mk(permission_codes=["stores.manage"])
    assert over.status_code == 403 and "stores.manage" in over.json()["detail"]
    ok = mk(permission_codes=["orders.view"])
    assert ok.status_code == 201
    made.append(ok.json()["employee_code"])
    # and they can't touch a full admin's account at all
    full_id = next(e["id"] for e in client.get(f"{API}/employees?category=admin", headers=admin).json()["items"] if e["employee_code"] == "ADMIN001")
    assert client.patch(f"{API}/employees/{full_id}", headers=custom, json={"full_name": "hacked"}).status_code in (403, 404)


def test_switching_admin_access_and_keeping_one_full_admin(client, admin, made):
    code, r = _create(client, admin, made, account_type="admin", admin_access="custom", permission_codes=["orders.view"])
    uid = r.json()["id"]
    up = client.put(f"{API}/employees/{uid}/admin-access", headers=admin, json={"full": True})
    assert up.status_code == 200 and up.json()["admin_level"] == "full" and up.json()["direct_permissions"] == []
    down = client.put(f"{API}/employees/{uid}/admin-access", headers=admin, json={"full": False})
    assert down.status_code == 200 and down.json()["admin_level"] == "custom"
    me_id = client.get(f"{API}/auth/me", headers=admin).json()["id"]
    refuse = client.put(f"{API}/employees/{me_id}/admin-access", headers=admin, json={"full": False})
    assert refuse.status_code == 422                                           # you can't reduce your own access
    assert client.post(f"{API}/employees/{me_id}/deactivate", headers=admin).status_code == 422


def test_new_admins_must_replace_the_temporary_password(client, admin, made):
    code, r = _create(client, admin, made, account_type="admin", admin_access="full")
    t = login(client, code, "Temp#Pass1234").json()
    assert client.get(f"{API}/auth/me", headers=auth(t["access_token"])).json()["must_change_password"] is True


# ---------------------------------------------------------------- partner delivery report (download only, no totals)
def _report_partner(client, admin, made, perms):
    code, r = _create(client, admin, made, account_type="partner", platform_id=_platform(client, admin, "blinkit"), permission_codes=perms)
    assert r.status_code == 201, r.text
    return _sign_in(client, code)


def test_partner_delivery_report_is_an_excel_without_totals_for_its_own_platform(client, admin, made):
    import io
    from datetime import date, timedelta

    from openpyxl import load_workbook

    h = _report_partner(client, admin, made, ["orders.view", "reports.delivery"])
    end = date.today()
    r = client.get(f"{API}/partner/delivery-report?start={end - timedelta(days=2)}&end={end}", headers=h)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert "blinkit-delivery-report" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header[:6] == ["#", "Store", "Store code", "Channel", "State", "City"] and len(header) == 6 + 3
    assert "Total" not in header                                              # no row-total column
    labels = {str(ws.cell(row=i, column=2).value).upper() for i in range(2, ws.max_row + 1)}
    assert "TOTAL" not in labels                                              # no total row
    channels = {ws.cell(row=i, column=4).value for i in range(2, ws.max_row + 1)}
    assert channels == {"Blinkit"}                                            # never another platform


def test_delivery_report_needs_its_permission_and_a_sane_range(client, admin, partner, made):
    assert client.get(f"{API}/partner/delivery-report", headers=partner).status_code == 403            # partner without reports.delivery
    h = _report_partner(client, admin, made, ["reports.delivery"])
    assert client.get(f"{API}/partner/delivery-report?start=2026-01-01&end=2026-12-31", headers=h).status_code == 422   # > 62 days
    assert client.get(f"{API}/partner/delivery-report?start=2026-02-10&end=2026-02-01", headers=h).status_code == 422
    assert client.get(f"{API}/partner/delivery-report", headers=admin).status_code == 403             # staff use the main Reports page


def test_admin_report_still_has_totals(client, admin):
    import io
    from datetime import date, timedelta

    from openpyxl import load_workbook

    end = date.today()
    r = client.get(f"{API}/orders/matrix/export?start={end - timedelta(days=1)}&end={end}&partner=blinkit", headers=admin)
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]][-1] == "Total"
    assert any(str(ws.cell(row=i, column=2).value) == "TOTAL" for i in range(2, ws.max_row + 1))
