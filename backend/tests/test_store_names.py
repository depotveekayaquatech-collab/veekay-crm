"""Store names are the store's identity: unique per platform, never changed; Outlet IDs can be corrected; duplicates are listed."""
from datetime import date, timedelta

from app.models.order_entry import OrderEntry
from app.models.store import Store

from .conftest import API


def test_a_store_name_cannot_be_changed_by_anyone(client, admin, dev, make_store):
    s = make_store(code="NM-1", name="Fixed Name Mart")
    for who in (admin, dev):
        r = client.patch(f"{API}/stores/{s.id}", json={"name": "Another Name"}, headers=who)
        assert r.status_code == 422 and "can't be changed" in r.json()["detail"]
    same = client.patch(f"{API}/stores/{s.id}", json={"name": "  fixed name   mart ", "city": "Noida"}, headers=admin)
    assert same.status_code == 200 and same.json()["name"] == "Fixed Name Mart" and same.json()["city"] == "Noida"


def test_the_outlet_id_can_be_corrected_and_the_store_keeps_its_orders(client, admin, db, make_store):
    s = make_store(code="NM-TEMP", name="Id Fix Mart")
    db.add(OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=date.today() - timedelta(days=1), bottle_count=7, source="manual"))
    db.commit()
    r = client.patch(f"{API}/stores/{s.id}", json={"external_code": "NM-REAL"}, headers=admin)
    assert r.status_code == 200 and r.json()["external_code"] == "NM-REAL" and r.json()["name"] == "Id Fix Mart"
    assert db.query(OrderEntry).filter_by(store_id=s.id).count() == 1


def test_an_outlet_id_that_belongs_to_another_store_is_refused_clearly(client, admin, make_store):
    a = make_store(code="NM-A", name="Mart A")
    make_store(code="NM-B", name="Mart B")
    r = client.patch(f"{API}/stores/{a.id}", json={"external_code": "NM-B"}, headers=admin)
    assert r.status_code == 409 and "already belongs to another store" in r.json()["detail"]


def test_a_store_with_an_existing_name_cannot_be_added_by_hand(client, admin, make_store):
    s = make_store(code="NM-C", name="Taken Name Mart")
    body = {"name": "taken name  mart", "external_code": "NM-C2", "start_date": "2026-03-01", "partner_organization_id": str(s.partner_organization_id)}
    r = client.post(f"{API}/stores", json=body, headers=admin)
    assert r.status_code == 409 and "already exists" in r.json()["detail"]


def test_order_import_finds_the_store_by_name_even_with_a_different_outlet_id(client, dev, db, make_store):
    temp = make_store(code="NM-TEMP-9", name="Blinkit - Order Name Mart")
    day = date.today() - timedelta(days=1)
    csv = f"Outlet ID,Outlet Name,{day}\nREAL-ID-9,blinkit - order  name mart,12\nGHOST-1,Blinkit - Nobody,5\n".encode()
    r = client.post(f"{API}/orders/import", headers=dev, data={"platform": "blinkit"}, files={"file": ("o.csv", csv, "text/csv")})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["created"] == 1 and out["unknown_stores"] == 1
    assert "by Outlet Name" in out["warnings"][0] and "Blinkit - Nobody" in out["warnings"][0]
    assert db.query(OrderEntry).filter_by(store_id=temp.id, order_date=day).one().bottle_count == 12


def test_order_import_uses_the_id_only_when_the_sheet_has_no_name_column(client, dev, db, make_store):
    s = make_store(code="NM-ONLYID", name="Blinkit - Only Id Mart")
    day = date.today() - timedelta(days=1)
    csv = f"Outlet ID,{day}\nNM-ONLYID,4\n".encode()
    r = client.post(f"{API}/orders/import", headers=dev, data={"platform": "blinkit"}, files={"file": ("o.csv", csv, "text/csv")})
    assert r.status_code == 200 and r.json()["created"] == 1
    assert db.query(OrderEntry).filter_by(store_id=s.id).count() == 1


def test_duplicates_are_listed_with_their_order_days(client, admin, db, make_store):
    a = make_store(code="NM-D1", name="Dup List Mart")
    make_store(code="NM-D2", name="Dup List Mart")
    make_store(code="nm d3", name="Id Lookalike A")
    make_store(code="NM-D3", name="Id Lookalike B")
    db.add(OrderEntry(organization_id=a.organization_id, store_id=a.id, order_date=date.today() - timedelta(days=1), bottle_count=3, source="manual"))
    db.commit()
    got = client.get(f"{API}/stores/duplicates", headers=admin).json()
    names = [g for g in got["groups"] if g["kind"] == "name" and g["label"] == "Dup List Mart"]
    assert len(names) == 1 and sorted(s["external_code"] for s in names[0]["stores"]) == ["NM-D1", "NM-D2"]
    assert {s["external_code"]: s["order_days"] for s in names[0]["stores"]} == {"NM-D1": 1, "NM-D2": 0}
    assert any(g["kind"] == "id" and len(g["stores"]) == 2 for g in got["groups"])
    assert got["total_groups"] >= 2 and got["total_stores"] >= 4


def test_merging_a_duplicate_keeps_every_order(client, dev, db, make_store):
    old = make_store(code="NM-M1", name="Merge Keep Mart")
    new = make_store(code="NM-M2", name="Merge Keep Mart")
    d1, d2 = date.today() - timedelta(days=2), date.today() - timedelta(days=1)
    for d, n in ((d1, 2), (d2, 5)):
        db.add(OrderEntry(organization_id=old.organization_id, store_id=old.id, order_date=d, bottle_count=n, source="manual"))
    db.commit()
    nid = new.id
    r = client.post(f"{API}/stores/merge", json={"platform": "blinkit", "dry_run": False, "pairs": [{"old_code": "NM-M1", "new_code": "NM-M2"}]}, headers=dev)
    assert r.status_code == 200 and r.json()["merged"] == 1
    assert sorted((e.order_date, e.bottle_count) for e in db.query(OrderEntry).filter_by(store_id=nid)) == [(d1, 2), (d2, 5)]
    assert db.query(Store).filter(Store.name == "Merge Keep Mart").count() == 1


def _closed(make_store, code, name):
    return make_store(code=code, name=name, status="CLOSE")


def test_only_closed_stores_without_history_are_deleted(client, admin, db, make_store):
    empty = _closed(make_store, "DC-1", "Del Empty Mart")
    used = _closed(make_store, "DC-2", "Del Used Mart")
    live = make_store(code="DC-3", name="Del Live Mart", status="LIVE")
    db.add(OrderEntry(organization_id=used.organization_id, store_id=used.id, order_date=date.today() - timedelta(days=3), bottle_count=4, source="manual"))
    db.commit()
    ids = [str(empty.id), str(used.id), str(live.id)]
    eid, uid, lid = empty.id, used.id, live.id

    prev = client.post(f"{API}/stores/delete-closed", json={"store_ids": ids}, headers=admin).json()
    assert prev["dry_run"] is True and prev["deleted"] == 0
    assert prev["closed"] == 2 and prev["deletable_count"] == 1 and prev["kept_with_orders"] == 1
    assert [d["external_code"] for d in prev["deletable"]] == ["DC-1"]
    db.expire_all()
    assert db.get(Store, eid) is not None                                  # a preview deletes nothing

    done = client.post(f"{API}/stores/delete-closed", json={"store_ids": ids, "dry_run": False}, headers=admin).json()
    assert done["deleted"] == 1
    db.expire_all()
    assert db.get(Store, eid) is None
    assert db.get(Store, uid) is not None and db.query(OrderEntry).filter_by(store_id=uid).count() == 1   # history untouched
    assert db.get(Store, lid) is not None                                  # live stores are never touched


def test_a_closed_store_with_other_records_is_kept(client, admin, db, make_store):
    from app.models.cash_purchase import CashPurchase
    s = _closed(make_store, "DC-4", "Del Cash Mart")
    sid = s.id
    row = CashPurchase(organization_id=s.organization_id, kind="STORE", store_id=s.id, purchase_date=date.today(), category="EMERGENCY_REQUIREMENT",
                       amount=10, created_by_user_id=None)
    db.add(row)
    db.commit()
    r = client.post(f"{API}/stores/delete-closed", json={"store_ids": [str(sid)], "dry_run": False}, headers=admin).json()
    assert r["deleted"] == 0 and r["kept_with_other_records"] == 1
    db.expire_all()
    assert db.get(Store, sid) is not None
    db.delete(db.get(CashPurchase, row.id))
    db.commit()


def test_deleting_closed_stores_needs_the_manage_permission(client, emp):
    assert client.post(f"{API}/stores/delete-closed", json={}, headers=emp).status_code == 403


def test_stores_sharing_an_outlet_id_are_merged_by_store_id(client, dev, db, make_store):
    a = make_store(code="DUP-ID-1", name="Share Id Mart One")
    b = make_store(code="DUP-ID-1", name="Share Id Mart Two")
    day = date.today() - timedelta(days=1)
    db.add(OrderEntry(organization_id=a.organization_id, store_id=a.id, order_date=day, bottle_count=6, source="manual"))
    db.commit()
    by_code = client.post(f"{API}/stores/merge", json={"platform": "blinkit", "pairs": [{"old_code": "DUP-ID-1", "new_code": "DUP-ID-1"}]}, headers=dev).json()
    assert by_code["merged"] == 0 and "share Outlet ID" in by_code["results"][0]["reason"]      # an ambiguous ID is refused, never guessed
    exact = client.post(f"{API}/stores/merge", headers=dev, json={"platform": "blinkit", "dry_run": False,
                                                                   "pairs": [{"old_id": str(a.id), "new_id": str(b.id)}]}).json()
    assert exact["merged"] == 1
    assert db.query(OrderEntry).filter_by(store_id=b.id, order_date=day).one().bottle_count == 6


def test_the_upload_keeps_a_list_of_rows_it_could_not_use(client, dev, admin, db, make_store):
    from .test_store_import_limits import _xlsx
    make_store(code="IS-1", name="Blinkit - Issue Twin")
    make_store(code="IS-2", name="Blinkit - Issue Twin")
    rows = [
        ["IS-3", "Blinkit - Issue Twin", "Live", "Delhi", "Delhi", ""],          # two stores already carry this name, neither has IS-3
        ["IS-4", "Blinkit - Issue Fresh", "Live", "Delhi", "Delhi", ""],
        ["IS-5", "Blinkit - Issue Fresh", "Live", "Delhi", "Delhi", ""],         # the same name again in the sheet
        ["IS-6", "Blinkit - Issue Odd Status", "somewhat ok", "Delhi", "Delhi", ""],
        ["IS-7", "", "Live", "Delhi", "Delhi", ""],                              # an ID with no name
    ]
    r = client.post(f"{API}/stores/import", data={"platform": "blinkit"}, files={"file": ("a.xlsx", _xlsx(rows), "application/octet-stream")}, headers=dev)
    assert r.status_code == 200, r.text
    got = client.get(f"{API}/stores/import-issues", headers=admin).json()
    run = next(x for x in got["runs"] if x["platform_slug"] == "blinkit")
    reasons = {i["reason"] for i in run["issues"]}
    assert {"name_matches_several_stores", "name_repeated_in_sheet", "status_not_understood", "missing_name_or_id"} <= reasons
    twin = next(i for i in run["issues"] if i["reason"] == "name_matches_several_stores")
    assert twin["row"] == 2 and sorted(c["external_code"] for c in twin["crm_stores"]) == ["IS-1", "IS-2"]
    db.query(Store).filter(Store.external_code.in_(["IS-4", "IS-6"])).delete(synchronize_session=False)
    db.commit()


def test_a_store_manager_can_merge_but_an_employee_cannot(client, emp, admin, make_store):
    a, b = make_store(name="Merge Perm Mart"), make_store(name="Merge Perm Mart")
    body = {"platform": "blinkit", "pairs": [{"old_id": str(a.id), "new_id": str(b.id)}]}
    assert client.post(f"{API}/stores/merge", json=body, headers=emp).status_code == 403
    assert client.post(f"{API}/stores/merge", json=body, headers=admin).status_code == 200      # a preview


# ---------------------------------------------------------------- wrong-platform check / move, and employees correcting their own entries
def test_a_store_on_the_wrong_platform_is_flagged_and_moves_with_its_history_and_loses_its_zepto_cap(client, admin, db, make_store):
    from app.models.store_cap import StoreCap
    z = make_store(code="55123", name="SS Wrongly On Zepto Mart", platform="zepto", region=None)
    day = date.today() - timedelta(days=1)
    db.add(OrderEntry(organization_id=z.organization_id, store_id=z.id, order_date=day, bottle_count=4, source="manual"))
    db.add(StoreCap(store_id=z.id, daily_cap=10, source="manual"))
    db.commit()
    zid, zpid = z.id, z.partner_organization_id
    found = client.get(f"{API}/stores/platform-check", headers=admin).json()
    row = next(r for r in found["stores"] if r["store"]["id"] == str(zid))
    assert row["suggested_platform"] == "blinkit" and row["has_capping"] is True and row["order_days"] == 1

    moved = client.post(f"{API}/stores/{zid}/move-platform", json={"to_platform": "blinkit"}, headers=admin)
    assert moved.status_code == 200 and moved.json()["capping_removed"] == 1
    db.expire_all()
    assert db.get(Store, zid).partner_organization_id != zpid                                           # it is on the other platform now
    assert db.query(StoreCap).filter_by(store_id=zid).count() == 0                                   # the Zepto cap did not follow it
    assert db.query(OrderEntry).filter_by(store_id=zid).count() == 1                                  # its orders did
    again = client.post(f"{API}/stores/{zid}/move-platform", json={"to_platform": "blinkit"}, headers=admin)
    assert again.status_code == 422                                                                     # already there


def test_moving_to_a_platform_that_already_has_that_name_is_refused(client, admin, make_store):
    a = make_store(code="MP-1", name="Same Name Both Platforms", platform="zepto", region=None)
    make_store(code="MP-2", name="Same Name Both Platforms", platform="blinkit")
    r = client.post(f"{API}/stores/{a.id}/move-platform", json={"to_platform": "blinkit"}, headers=admin)
    assert r.status_code == 409 and "already has a store named" in r.json()["detail"]
    flagged = client.get(f"{API}/stores/platform-check", headers=admin).json()["stores"]
    assert any(x["store"]["id"] == str(a.id) and any("also exists" in why for why in x["reasons"]) for x in flagged)


def test_a_level_one_employee_can_correct_or_clear_an_entry_of_their_store_in_the_open_dates(client, emp, admin, db, make_store):
    s = make_store(code="OC-1", region="NORTH")
    today = date.today().isoformat()
    mark = client.post(f"{API}/orders/mark", headers=emp, json={"store_id": str(s.id), "order_date": today, "bottle_count": 5})
    assert mark.status_code == 200, mark.text
    again = client.post(f"{API}/orders/mark", headers=emp, json={"store_id": str(s.id), "order_date": today, "bottle_count": 9})
    assert again.status_code == 409                                                                      # marking twice is still refused

    fixed = client.patch(f"{API}/orders/my-entry", headers=emp, json={"store_id": str(s.id), "order_date": today, "bottle_count": 8})
    assert fixed.status_code == 200, fixed.text
    assert db.query(OrderEntry).filter_by(store_id=s.id).one().bottle_count == 8
    from app.models.audit_log import AuditLog
    log = db.query(AuditLog).filter_by(action="order.corrected_by_employee").order_by(AuditLog.created_at.desc()).first()
    assert log.metadata_json["previous_count"] == 5 and log.metadata_json["bottle_count"] == 8        # what was there is on record

    cleared = client.patch(f"{API}/orders/my-entry", headers=emp, json={"store_id": str(s.id), "order_date": today})
    assert cleared.status_code == 200
    db.expire_all()
    assert db.query(OrderEntry).filter_by(store_id=s.id).count() == 0
    assert client.patch(f"{API}/orders/my-entry", headers=emp, json={"store_id": str(s.id), "order_date": today, "bottle_count": 3}).status_code == 404   # nothing to correct

    db.add(OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=date.today(), bottle_count=6, source="admin"))
    db.commit()
    theirs = client.patch(f"{API}/orders/my-entry", headers=emp, json={"store_id": str(s.id), "order_date": today, "bottle_count": 1})
    assert theirs.status_code == 200, theirs.text                                                       # an admin's entry of their store: allowed at Level 1, and audited
    assert db.query(OrderEntry).filter_by(store_id=s.id).one().bottle_count == 1
    old = (date.today().replace(day=1) - timedelta(days=70)).isoformat()
    assert client.patch(f"{API}/orders/my-entry", headers=emp, json={"store_id": str(s.id), "order_date": old, "bottle_count": 1}).status_code == 422


def test_dashboard_numbers_are_per_platform_and_never_added_together(client, admin, db, make_store):
    b = make_store(code="DB-B", name="Dash Blinkit Mart", platform="blinkit")
    z = make_store(code="DB-Z", name="Dash Zepto Mart", platform="zepto", region=None)
    today = date.today()
    for s, n in ((b, 7), (z, 11)):
        db.add(OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=today, bottle_count=n, source="manual"))
    db.commit()
    both = client.get(f"{API}/orders/insights", headers=admin).json()
    blk = client.get(f"{API}/orders/insights", params={"partner": "blinkit"}, headers=admin).json()
    zpt = client.get(f"{API}/orders/insights", params={"partner": "zepto"}, headers=admin).json()
    key = "bottles_today" if "bottles_today" in blk else "bottlesToday"
    assert blk[key] + zpt[key] == both[key]                                    # the two channels add up to the old mixed number, never inside one view
    assert {p["slug"] for p in blk["platforms"]} == {"blinkit"} and {p["slug"] for p in zpt["platforms"]} == {"zepto"}
    assert {"blinkit", "zepto"} <= {p["slug"] for p in blk["status_by_platform"]}          # the Store status card lists every platform, each on its own
    assert all(sc["partner_slug"] == "blinkit" for sc in blk["scopes"])
    assert client.get(f"{API}/orders/insights", params={"partner": "nope"}, headers=admin).status_code == 404


def test_every_store_gets_a_permanent_system_code_unique_across_platforms(client, admin, db, make_store):
    import re
    b = make_store(code="9001", name="East Patel Nagar", platform="blinkit")
    z = make_store(code="9001", name="East Patel Nagar", platform="zepto", region=None)      # same name AND same Outlet ID on both platforms
    assert re.fullmatch(r"BKT\d{6}", b.crm_code) and re.fullmatch(r"ZPT\d{6}", z.crm_code) and b.crm_code != z.crm_code
    got = client.get(f"{API}/stores", params={"search": z.crm_code}, headers=admin).json()
    assert [s["crm_code"] for s in got["items"]] == [z.crm_code]                            # one code finds exactly one store
    again = client.patch(f"{API}/stores/{b.id}", json={"crm_code": "ZPT000001", "city": "Delhi"}, headers=admin)
    assert again.status_code == 200 and again.json()["crm_code"] == b.crm_code             # the code can never be edited


def test_moving_a_store_to_the_other_platform_gives_it_that_platforms_prefix(client, admin, db, make_store):
    z = make_store(code="MV-1", name="Move Code Mart", platform="zepto", region=None)
    before = z.crm_code
    r = client.post(f"{API}/stores/{z.id}/move-platform", json={"to_platform": "blinkit"}, headers=admin)
    assert r.status_code == 200
    db.expire_all()
    assert before.startswith("ZPT") and db.get(Store, z.id).crm_code.startswith("BKT")


def test_a_sheet_with_the_crm_code_matches_exactly_and_refuses_the_other_platforms_code(client, dev, db, make_store):
    from .test_store_import_limits import _xlsx, HEADER
    from openpyxl import Workbook
    import io as _io
    s = make_store(code="CC-1", name="Blinkit - Crm Code Mart")
    other = make_store(code="CC-2", name="Zepto Crm Code Mart", platform="zepto", region=None)
    wb = Workbook(); ws = wb.active
    ws.append(["CRM Code", "Outlet ID", "Outlet Name", "STATUS", "State", "City"])
    ws.append([s.crm_code, "CC-NEW", "Blinkit - Crm Code Mart", "Live", "Delhi", "Delhi"])       # the code settles it; the ID is corrected
    ws.append([other.crm_code, "CC-X", "Blinkit - Wrong Platform Row", "Live", "Delhi", "Delhi"])  # a Zepto code in a Blinkit upload
    buf = _io.BytesIO(); wb.save(buf)
    r = client.post(f"{API}/stores/import", data={"platform": "blinkit"}, files={"file": ("a.xlsx", buf.getvalue(), "application/octet-stream")}, headers=dev)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["created"] == 0 and any("not a store of" in w for w in out["warnings"])
    db.expire_all()
    assert db.get(Store, s.id).external_code == "CC-NEW"


def test_the_calendar_says_who_entered_each_day(client, emp, admin, db, make_store):
    s = make_store(code="PV-1", region="NORTH")
    today = date.today().isoformat()
    client.post(f"{API}/orders/mark", headers=emp, json={"store_id": str(s.id), "order_date": today, "bottle_count": 4})
    cal = client.get(f"{API}/orders/calendar", params={"store_id": str(s.id)}, headers=emp).json()
    day = next(d for d in cal["days"] if d["date"] == today)
    assert day["source"] == "employee" and day["entered_by"] and day["entered_at"]


def test_an_order_import_can_be_undone_for_one_month_and_nothing_else(client, dev, admin, db, make_store):
    from datetime import datetime
    from app.models.audit_log import AuditLog
    from app.models.deleted_row import DeletedRow
    s = make_store(code="RV-1", name="Revert Mart One")
    t = make_store(code="RV-2", name="Revert Mart Two")
    today = date.today()
    first = today.replace(day=1)
    last_month = first - timedelta(days=5)
    csv = f"Outlet ID,Outlet Name,{last_month.isoformat()},{first.isoformat()},{(first + timedelta(days=1)).isoformat() if first + timedelta(days=1) <= today else first.isoformat()}\n".encode()
    # an employee marks day 1 on store t first; the import will overwrite it
    emp_day = first
    db.add(OrderEntry(organization_id=t.organization_id, store_id=t.id, order_date=emp_day, bottle_count=3, source="employee"))
    db.commit()
    body = f"Outlet ID,Outlet Name,{last_month},{emp_day}\nRV-1,{s.name},9,9\nRV-2,{t.name},,20\n".encode()
    r = client.post(f"{API}/orders/import", headers=dev, data={"platform": "blinkit", "overwrite": "true"}, files={"file": ("o.csv", body, "text/csv")})
    assert r.status_code == 200, r.text
    assert db.query(OrderEntry).filter_by(store_id=t.id, order_date=emp_day).one().bottle_count == 20     # the import overwrote the employee's 3

    month = emp_day.strftime("%Y-%m")
    runs = client.get(f"{API}/orders/import-runs", params={"platform": "blinkit", "month": month}, headers=dev).json()["runs"]
    mine = [x for x in runs if x["created_in_month"] >= 1]
    assert mine and mine[0]["overwrote_in_month"] >= 1
    rid = mine[0]["run_id"]
    pre = client.post(f"{API}/orders/import-revert", headers=dev, json={"platform": "blinkit", "month": month, "run_ids": [rid]}).json()
    assert pre["dry_run"] is True and pre["will_delete"] >= 1 and pre["will_restore"] >= 1
    assert db.query(OrderEntry).filter_by(store_id=s.id, order_date=emp_day).count() == 1               # a preview deletes nothing

    assert client.post(f"{API}/orders/import-revert", headers=dev, json={"platform": "blinkit", "month": month, "run_ids": [rid], "dry_run": False}).status_code == 422   # needs the confirm text
    done = client.post(f"{API}/orders/import-revert", headers=dev, json={"platform": "blinkit", "month": month, "run_ids": [rid], "dry_run": False, "confirm": f"REVERT BLINKIT {month}"})
    assert done.status_code == 200, done.text
    db.expire_all()
    assert db.query(OrderEntry).filter_by(store_id=s.id, order_date=emp_day).count() == 0                  # what the import created this month is gone
    back = db.query(OrderEntry).filter_by(store_id=t.id, order_date=emp_day).one()
    assert back.bottle_count == 3 and back.source == "employee"                                           # the employee's own number is back
    if last_month.strftime("%Y-%m") != month:
        assert db.query(OrderEntry).filter_by(store_id=s.id, order_date=last_month).count() == 1           # other months are untouched
    assert db.query(DeletedRow).filter(DeletedRow.table_name == "order_entries", DeletedRow.row_id.isnot(None)).count() >= 1   # and the deleted rows are kept
    assert client.get(f"{API}/orders/import-runs", params={"platform": "blinkit", "month": month}, headers=admin).status_code == 403   # developer only


def test_a_store_that_exists_twice_is_flagged_to_the_employee_and_the_split_is_shown(client, emp, admin, db, make_store):
    a = make_store(code="TW-1", name="Twin Entries Mart", region="NORTH")
    b = make_store(code="TW-2", name="Twin Entries Mart", region="NORTH")
    solo = make_store(code="TW-3", name="Single Entries Mart", region="NORTH")
    today = date.today()
    for s in (a, b):
        db.add(OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=today, bottle_count=5, source="employee"))
    db.commit()
    mine = {s["id"]: s for s in client.get(f"{API}/orders/my-stores", headers=emp).json()}
    assert mine[str(a.id)]["duplicate_name"] is True and mine[str(b.id)]["duplicate_name"] is True and mine[str(solo.id)]["duplicate_name"] is False
    assert mine[str(a.id)]["crm_code"] == a.crm_code
    cal = client.get(f"{API}/orders/calendar", params={"store_id": str(a.id)}, headers=emp).json()
    assert cal["store_crm_code"] == a.crm_code
    groups = client.get(f"{API}/stores/duplicates", headers=admin).json()["groups"]
    g = next(x for x in groups if x["kind"] == "name" and x["label"] == "Twin Entries Mart")
    assert g["entries_split"] is True and all(s["order_days_this_month"] == 1 for s in g["stores"])


def test_deleted_entries_can_be_restored_for_employees_only(client, dev, db, make_store):
    s = make_store(code="RE-1", name="Restore Emp Mart")
    d1, d2 = date.today() - timedelta(days=3), date.today() - timedelta(days=2)
    emp_row = OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=d1, bottle_count=12, source="employee")
    imp_row = OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=d2, bottle_count=20, source="admin")
    db.add_all([emp_row, imp_row])
    db.commit()
    for e in db.query(OrderEntry).filter_by(store_id=s.id).all():
        db.delete(e)                                                              # the database copies each into deleted_rows
    db.commit()
    assert db.query(OrderEntry).filter_by(store_id=s.id).count() == 0
    rng = {"start": (d1 - timedelta(days=1)).isoformat(), "end": date.today().isoformat()}
    listed = client.get(f"{API}/developer/order-audit/deleted", params={**rng, "employee_only": "true"}, headers=dev).json()
    mine = [i for i in listed["items"] if i["store_code"] == "RE-1"]
    assert len(mine) == 1 and mine[0]["source"] == "employee"                      # the import's entry is not even offered
    pre = client.post(f"{API}/developer/order-audit/deleted/restore", headers=dev, json={**rng, "employee_only": True}).json()
    assert pre["dry_run"] is True and pre["restorable"] >= 1
    done = client.post(f"{API}/developer/order-audit/deleted/restore", headers=dev, json={**rng, "employee_only": True, "dry_run": False}).json()
    assert done["restored"] >= 1
    db.expire_all()
    back = db.query(OrderEntry).filter_by(store_id=s.id).all()
    assert [(e.order_date, e.bottle_count, e.source) for e in back] == [(d1, 12, "employee")]      # only the employee's entry came back


def test_the_delivery_matrix_never_drops_a_store_that_has_entries_in_the_period(client, admin, db, make_store):
    s = make_store(code="MX-1", name="Matrix Late Start Mart")
    day = date.today() - timedelta(days=2)
    db.add(OrderEntry(organization_id=s.organization_id, store_id=s.id, order_date=day, bottle_count=37, source="employee"))
    s.start_date = date.today() + timedelta(days=30)           # a start date later than the whole period (set by a sync, a typo...)
    db.commit()
    params = {"start": (date.today() - timedelta(days=6)).isoformat(), "end": date.today().isoformat(), "partner": "blinkit", "page_size": 200}
    m = client.get(f"{API}/orders/matrix", params=params, headers=admin).json()
    assert any(r["code"] == "MX-1" if "code" in r else r.get("external_code") == "MX-1" or "MX-1" in str(r) for r in m["rows"])
    rep = client.get(f"{API}/orders/report", params={"start": params["start"], "end": params["end"], "group_by": "partner", "partner": "blinkit"}, headers=admin).json()
    assert m["grand_total_bottles"] == rep["total_bottles"]                      # the matrix and the report add up to the same number


def test_a_store_that_is_not_live_never_gets_an_entry_from_a_sheet_or_an_admin(client, dev, admin, db, make_store):
    from app.models.order_entry import OrderEntry

    live = make_store(code="NL-LIVE-1", name="Blinkit - Live Mart NL")
    closed = make_store(code="NL-CLOSED-1", name="Blinkit - Closed Mart NL", status="CLOSE")
    day = date.today() - timedelta(days=1)
    csv = f"Outlet ID,Outlet Name,{day}\nNL-LIVE-1,Blinkit - Live Mart NL,12\nNL-CLOSED-1,Blinkit - Closed Mart NL,7\n".encode()
    r = client.post(f"{API}/orders/import", headers=dev, data={"platform": "blinkit"}, files={"file": ("o.csv", csv, "text/csv")})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["created"] == 1 and out["not_live_stores"] == 1 and any("not live" in w for w in out["warnings"])
    db.expire_all()
    assert db.query(OrderEntry).filter_by(store_id=live.id).count() == 1
    assert db.query(OrderEntry).filter_by(store_id=closed.id).count() == 0

    # an administrator cannot add one either...
    add = client.patch(f"{API}/orders/entry", headers=admin, json={"store_id": str(closed.id), "order_date": day.isoformat(), "bottle_count": 5})
    assert add.status_code == 409 and "not live" in add.text
    assert db.query(OrderEntry).filter_by(store_id=closed.id).count() == 0

    # ...but an entry already sitting on a store that is not live can still be cleared
    from .conftest import plant_entry
    plant_entry(db, closed, day, 9)
    cleared = client.patch(f"{API}/orders/entry", headers=admin, json={"store_id": str(closed.id), "order_date": day.isoformat(), "bottle_count": None})
    assert cleared.status_code == 200, cleared.text
    db.expire_all()
    assert db.query(OrderEntry).filter_by(store_id=closed.id).count() == 0


def test_order_sheet_headings_without_a_year_are_read_with_the_tabs_year():
    from app.services.order_import_service import _header_date as head, year_of

    today = date(2026, 10, 28)
    assert year_of("NORTH OCT-2026") == 2026 and year_of("Sheet1") is None
    assert head("1-Oct", 2026, today) == date(2026, 10, 1) and head("28-Oct", None, today) == date(2026, 10, 28)
    assert head("01 October 2026", None, today) == date(2026, 10, 1) and head("2-Sept", None, today) == date(2026, 9, 2) and head("1-Oct-26", None, today) == date(2026, 10, 1)
    assert head("15 Dec", None, date(2027, 1, 5)) == date(2026, 12, 15)            # a December sheet read in January is last year's
    assert head("2026-10-01", None, today) == date(2026, 10, 1)                      # the old formats still work
    assert head("Total Count", None, today) is None and head("Outlet ID", None, today) is None and head("Region", None, today) is None


def test_an_order_workbook_with_day_month_headings_loads_and_a_zero_only_counts_where_the_crm_has_a_zero(client, dev, db, make_store):
    import io

    from openpyxl import Workbook

    from app.models.order_entry import OrderEntry

    from .conftest import plant_entry

    today = date.today()
    if today.day < 4:
        return
    days = [today.replace(day=i) for i in range(1, 5)]
    a = make_store(code="OS-A", name="Blinkit - Sheet Mart A")
    b = make_store(code="OS-B", name="Blinkit - Sheet Mart B")
    c = make_store(code="OS-C", name="Blinkit - Sheet Mart C")
    plant_entry(db, a, days[1], 0)                                                         # the CRM already has a real 0 on day 2 for store A
    wb = Workbook()
    ws = wb.active
    ws.title = f"NORTH {today:%b}-{today.year}".upper()                                    # the tab carries the year, the headings do not
    ws.append(["Region", "Outlet ID", "Outlet Name", "Total Count"] + [f"{d.day}-{d:%b}" for d in days])
    ws.append(["NORTH", "OS-A", "Blinkit - Sheet Mart A", 12, 5, 0, 7, 0])                  # 0 on day 2 matches the CRM's 0 (nothing to change); 0 on day 4 is not marked
    ws.append(["NORTH", "OS-B", "Some Other Name In The Sheet", 12, 3, 3, 3, 3])            # the name matches no store, the Outlet ID does
    ws.append(["NORTH", "OS-C", "Blinkit - Sheet Mart C", 0, 0, 0, 0, 0])                   # zeros the CRM does not have: never marked
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post(f"{API}/orders/import", headers=dev, data={"platform": "blinkit"},
                    files={"file": ("Arti october.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["created"] == 6 and out["unchanged"] == 1 and out["matched_by_id"] == 1 and out["zeros_skipped"] == 5 and out["unknown_stores"] == 0
    assert any("only counts on a day the CRM" in w for w in out["warnings"]) and any("Outlet ID" in w for w in out["warnings"])
    db.expire_all()
    got = lambda s: {e.order_date: e.bottle_count for e in db.query(OrderEntry).filter_by(store_id=s.id)}  # noqa: E731
    assert got(a) == {days[0]: 5, days[1]: 0, days[2]: 7}                                      # day 2 is the CRM's own 0; day 4's 0 was not written
    assert got(b) == {d: 3 for d in days}
    assert got(c) == {}


def test_a_blank_cell_and_a_zero_never_mark_a_day_and_a_zero_never_replaces_a_number(client, dev, db, make_store):
    from app.models.order_entry import OrderEntry

    from .conftest import plant_entry

    day = date.today() - timedelta(days=1)
    s = make_store(code="ZR-1", name="Blinkit - Zero Rule Mart")
    t = make_store(code="ZR-2", name="Blinkit - Zero Rule Mart Two")
    plant_entry(db, s, day, 9)                                                    # the CRM has a real 9; the sheet says 0 and overwrite is on
    csv = f"Outlet ID,Outlet Name,{day}\nZR-1,Blinkit - Zero Rule Mart,0\nZR-2,Blinkit - Zero Rule Mart Two,\n".encode()   # a 0 for one store, a blank for the other
    r = client.post(f"{API}/orders/import", headers=dev, data={"platform": "blinkit", "overwrite": "true"}, files={"file": ("o.csv", csv, "text/csv")})
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 0 and r.json()["updated"] == 0 and r.json()["zeros_skipped"] == 1
    db.expire_all()
    assert db.query(OrderEntry).filter_by(store_id=s.id).one().bottle_count == 9           # never replaced by the sheet's 0
    assert db.query(OrderEntry).filter_by(store_id=t.id).count() == 0                      # a blank cell marks nothing: left for the employee


def test_the_matrix_export_can_show_days_nobody_marked_as_0_or_leave_them_blank(client, admin, make_store, db):
    import io

    from openpyxl import load_workbook

    from .conftest import plant_entry

    d1 = date.today() - timedelta(days=3)
    s = make_store(code="MX-1", name="Blinkit - Matrix Mart")
    plant_entry(db, s, d1, 5)                                                        # marked on day 1 only
    plant_entry(db, s, d1 + timedelta(days=1), 0)                                    # a real 0 on day 2; nothing on day 3

    def cells(**extra):
        r = client.get(f"{API}/orders/matrix/export", params={"start": d1.isoformat(), "end": (d1 + timedelta(days=2)).isoformat(), "partner": "blinkit", **extra}, headers=admin)
        assert r.status_code == 200, r.text
        ws = load_workbook(io.BytesIO(r.content)).active
        head = [c.value for c in ws[1]]
        row = next(r_ for r_ in ws.iter_rows(min_row=2, values_only=True) if r_[head.index("Outlet ID")] == "MX-1")
        return row[head.index("Total Count")], list(row[-3:])

    assert cells() == (5, [5, 0, 0])                                                 # the default is as before: unmarked shows 0
    assert cells(blank_as_zero="true") == (5, [5, 0, 0])
    assert cells(blank_as_zero="false") == (5, [5, 0, None])                         # a real 0 stays 0, an unmarked day is an empty cell
