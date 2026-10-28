"""Who may see Zepto capping, and city-wise vendor rates."""
import csv
import io
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.models.order_entry import OrderEntry
from app.models.store_cap import StoreCap
from app.services import capping_service, store_sync_service

from .conftest import API


def _csv(header, rows) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(header)
    w.writerows(rows)
    return out.getvalue().encode()


def _entry(db, store, day, n):
    db.add(OrderEntry(organization_id=store.organization_id, store_id=store.id, order_date=day, bottle_count=n, source="ADMIN"))
    db.commit()


# ---------------------------------------------------------------- who may use it

def test_every_active_zepto_employee_can_view_capping_whatever_their_state(client, make_user, make_store, db):
    code, pw, _ = make_user("employee", platform="zepto", region="SOUTH")
    for_other_state = make_store(platform="zepto", region=None, state="Uttar Pradesh", city="Noida")
    from .conftest import auth, login
    token = login(client, code, pw).json()["access_token"]
    r = client.get(f"{API}/capping", headers=auth(token))
    assert r.status_code == 200, r.text
    assert any(x["outlet_id"] == for_other_state.external_code for x in r.json()["rows"])   # a store in another state is listed


def test_blinkit_employees_and_the_unauthenticated_are_refused(client, emp):
    assert client.get(f"{API}/capping", headers=emp).status_code == 403
    assert client.get(f"{API}/capping").status_code in (401, 403)


def test_admin_can_view_and_a_deactivated_zepto_employee_cannot(client, admin, make_user):
    assert client.get(f"{API}/capping", headers=admin).status_code == 200
    code, pw, _ = make_user("employee", platform="zepto", region="SOUTH", active=False)
    from .conftest import login
    assert login(client, code, pw).status_code in (401, 403)


# ---------------------------------------------------------------- vendor rates

def _rates(client, headers, rows, header=("City", "Rate"), confirm=False, platform="blinkit"):
    return client.post(f"{API}/vendor-rates/upload", headers=headers, data={"platform": platform, "confirm": str(confirm).lower()},
                       files={"file": ("rates.csv", _csv(list(header), rows), "text/csv")})


def test_vendor_rates_need_the_rates_permission(client, emp, zepto_emp):
    assert client.get(f"{API}/vendor-rates", headers=emp).status_code == 403
    assert _rates(client, zepto_emp, [["Delhi", "5"]]).status_code == 403


def test_new_rates_apply_but_a_changed_rate_asks_first(client, admin, make_store):
    make_store(platform="blinkit", city="RateTown", code="RT-1")
    make_store(platform="blinkit", city="RateTown", code="RT-2")
    r = _rates(client, admin, [["RateTown", "9.50"]]).json()
    assert r["applied"] and r["new"] == 1 and not r["needs_confirmation"]
    # same rate again: nothing to do
    assert _rates(client, admin, [["RateTown", "9.50"]]).json()["unchanged"] == 1
    # a different rate: nothing is written until confirmed, and the answer lists the stores affected
    r = _rates(client, admin, [["ratetown", "11"]]).json()
    assert r["needs_confirmation"] and not r["applied"] and r["changed"] == 1 and r["stores_affected"] == 2
    assert r["changes"][0]["old_rate"] == "9.50" and r["changes"][0]["new_rate"] == "11.00"
    listed = {x["city"]: x["rate"] for x in client.get(f"{API}/vendor-rates", headers=admin).json()}
    assert listed["RateTown"] == "9.50"                      # still the old rate
    r = _rates(client, admin, [["RateTown", "11"]], confirm=True).json()
    assert r["applied"]
    listed = {x["city"]: x["rate"] for x in client.get(f"{API}/vendor-rates", headers=admin).json()}
    assert listed["RateTown"] == "11.00"


def test_a_store_rate_overrides_its_city_and_bad_rows_are_reported(client, admin, make_store):
    s = make_store(platform="blinkit", city="OverTown", code="OT-1")
    r = _rates(client, admin, [["OverTown", "8", ""], ["", "12", "OT-1"], ["OverTown", "abc", ""], ["", "5", "MISSING"]], header=("City", "Rate", "Outlet ID")).json()
    assert r["applied"] and r["new"] == 2 and any("skipped" in w for w in r["warnings"])
    rates = client.get(f"{API}/vendor-rates", headers=admin).json()
    assert any(x["outlet_id"] == "OT-1" and x["rate"] == "12.00" for x in rates)


def test_blinkit_report_with_vendor_rate_is_admin_only_and_multiplies_bottles(client, admin, emp, make_store, db):
    s = make_store(platform="blinkit", city="ReportTown", code="RP-1")
    day = date.today() - timedelta(days=1)
    _entry(db, s, day, 10)
    _rates(client, admin, [["ReportTown", "3.5"]])
    params = {"start": day.isoformat(), "end": day.isoformat(), "partner": "blinkit", "include_rate": "true"}
    assert client.get(f"{API}/orders/matrix/export", headers=emp, params=params).status_code in (403,)
    r = client.get(f"{API}/orders/matrix/export", headers=admin, params=params)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    head = [c.value for c in ws[1]]
    assert "Vendor Rate" in head and "Amount" in head
    row = next(row for row in ws.iter_rows(min_row=2, values_only=True) if row[head.index("Outlet ID")] == "RP-1")
    assert row[head.index("Vendor Rate")] == 3.5 and row[head.index("Amount")] == 35.0
    # the normal report is unchanged
    plain = load_workbook(io.BytesIO(client.get(f"{API}/orders/matrix/export", headers=admin, params={**params, "include_rate": "false"}).content)).active
    assert "Vendor Rate" not in [c.value for c in plain[1]]


def test_a_blinkit_person_never_sees_zepto_capping_even_holding_correct_any_entry(client, make_user, make_store):
    from .conftest import auth, login
    code, pw, _ = make_user("employee", platform="blinkit", region="NORTH", grants=("orders.view", "orders.correct", "capping.alerts"))
    h = auth(login(client, code, pw).json()["access_token"])
    assert client.get(f"{API}/capping", headers=h).status_code == 403
