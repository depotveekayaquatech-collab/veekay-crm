"""Template sheets: one per spreadsheet upload, in exactly the format that upload reads."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.services import cash_sync_service, order_import_service
from app.services import template_service as ts

from .conftest import API


def _book(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(r.content))


@pytest.mark.parametrize("kind", list(ts.KINDS))
def test_every_template_downloads_as_a_workbook_with_a_how_to_fill_tab(client, emp, kind):
    wb = _book(client.get(f"{API}/templates/{kind}", headers=emp))
    assert wb.sheetnames[-1] == "How to fill" and len(wb.sheetnames) >= 2
    assert wb[wb.sheetnames[0]].max_row >= 2                                   # a header and sample rows


def test_templates_need_a_login_and_a_known_kind(client, emp):
    assert client.get(f"{API}/templates/orders").status_code == 401
    assert client.get(f"{API}/templates/nonsense", headers=emp).status_code == 404
    assert client.get(f"{API}/templates/orders", params={"month": "2026-13x"}, headers=emp).status_code == 422


def test_the_order_template_is_read_by_the_order_import_with_the_tabs_year(client, emp):
    r = client.get(f"{API}/templates/orders", params={"month": "2026-10"}, headers=emp)
    wb = _book(r)
    tab = wb.sheetnames[0]
    assert tab == "NORTH OCT-2026"                                             # the tab carries the year, the headings do not
    rows = [tuple(c for c in row) for row in wb[tab].iter_rows(values_only=True)]
    assert rows[0][14] == "1-Oct" and rows[0][-1] == "31-Oct" and rows[0][6] == "Outlet Name" and rows[0][5] == "Outlet ID"
    found = order_import_service._locate(rows, tab)
    assert found is not None
    _, id_col, name_col, date_cols = found
    assert id_col == 5 and name_col == 6 and len(date_cols) == 31
    assert min(date_cols.values()) == date(2026, 10, 1) and max(date_cols.values()) == date(2026, 10, 31)


def test_the_other_templates_use_the_columns_their_uploads_look_for(client, emp):
    cash = _book(client.get(f"{API}/templates/cash", headers=emp))
    head = {str(c.value).strip().lower() for c in cash[cash.sheetnames[0]][1]}
    cols = cash_sync_service._UPLOAD_COLUMNS
    assert head & set(cols["date"]) and head & set(cols["amount"]) and head & set(cols["code"]) and head & set(cols["name"])

    att = _book(client.get(f"{API}/templates/attendance", params={"month": "2026-02"}, headers=emp))
    row = [c.value for c in att[att.sheetnames[0]][1]]
    assert row[:2] == ["Employee ID", "Name"] and row[2:] == list(range(1, 29))                # February 2026 has 28 days

    emps = _book(client.get(f"{API}/templates/employees", headers=emp))
    assert [c.value for c in emps["EMPLOYEES"][1]][:2] == ["Employee ID", "Employee Name"] and "State_Assignments" in emps.sheetnames
