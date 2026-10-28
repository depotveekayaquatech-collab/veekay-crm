"""
Template sheets: for every spreadsheet the CRM can read, a ready-made Excel file in exactly the format that upload understands (a sample sheet plus a "How to fill"
tab), so a person can shape their own sheet to match BEFORE it goes to the upload. They hold only invented example rows, never real data.

kinds: orders, employees, attendance, cash, cardsheet, stores, vendor-rates
"""
from __future__ import annotations

import calendar
import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

from fastapi import HTTPException, status

IST = ZoneInfo("Asia/Kolkata")
KINDS = ("orders", "employees", "attendance", "cash", "cardsheet", "stores", "vendor-rates")

STORE_SHEET_HEADER = [
    "S.NO", "Entity", "ZONE", "State", "City", "Outlet ID", "Outlet Name", "STATUS", "Rate - BCPL",
    "DATE RECEIVED", "START DATE BY BLINKIT", "start Date", "Address", "POC Name", "POC Contact no",
    "MAIL ID POC", "vendor", "Contact Details", "location",
]


def _month(value: str | None) -> date:
    if not value:
        n = datetime.now(IST)
        return date(n.year, n.month, 1)
    try:
        y, m = (int(x) for x in value.split("-"))
        return date(y, m, 1)
    except (ValueError, TypeError) as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Month must look like 2026-10.") from exc


def _day_headings(month: date) -> list[str]:
    last = calendar.monthrange(month.year, month.month)[1]
    return [f"{d}-{month:%b}" for d in range(1, last + 1)]          # 1-Oct, 2-Oct ... (no leading zero, the way the sheets are kept)


def _build(sheets: list[tuple[str, list[list], list[int]]], notes: list[str]) -> bytes:
    """sheets: [(tab name, rows (first row = header), column widths)], then the "How to fill" tab."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    side = Side(style="thin")
    for title, rows, widths in sheets:
        ws = wb.create_sheet(title[:31])
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(name="Calibri", size=11, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            c.border = Border(left=side, right=side, top=side, bottom=side)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    how = wb.create_sheet("How to fill")
    how.column_dimensions["A"].width = 110
    how.append(["How to fill this sheet"])
    how["A1"].font = Font(name="Calibri", size=13, bold=True)
    for line in notes:
        how.append([line])
        how.cell(row=how.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _orders(month: date) -> tuple[bytes, str]:
    days = _day_headings(month)
    head = ["Region", "Channel", "Entity", "STATE", "City", "Outlet ID", "Outlet Name", "POC Name", "Contact No", "PRIMARY VENDOR NAME", "PRIMARY LICENSE NUMBER",
            "SECONDARY VENDOR NAME", "SECONDARY LICENSE NUMBER", "Total Count", *days]
    pad = lambda row, vals: row + vals + [None] * (len(days) - len(vals))  # noqa: E731
    sample = [
        pad(["NORTH", "BLINKIT", "BCPL", "HARYANA", "Gurgaon", "BLK-1001", "Example Store One", "Example POC", "9000000001", "Example Vendor Pvt Ltd", "10000000000001", None, None, "=SUM(O2:%s2)" % _col(14 + len(days))], [7, 11, 11, 10]),
        pad(["NORTH", "BLINKIT", "BCPL", "HARYANA", "Gurgaon", "BLK-1002", "Example Store Two", "Example POC", "9000000002", "Example Vendor Pvt Ltd", "10000000000001", None, None, "=SUM(O3:%s3)" % _col(14 + len(days))], [None, 14, None, 9]),
    ]
    tab = f"NORTH {month:%b}-{month.year}".upper()
    notes = [
        "Only two things are needed: the store (Outlet Name, or Outlet ID) and a column for each date. Every other column (Region, Channel, POC, vendor...) is ignored, keep them if your sheet has them.",
        "Date headings can be 1-Oct, 01-Oct-2026, 2026-10-01 or a real Excel date. A heading with no year uses the year written in this tab's name (here %s)." % tab,
        "The store is found by its Outlet Name (capitals and extra spaces do not matter). If the name matches no store, the Outlet ID is used instead. Stores must already exist in the CRM.",
        "A BLANK cell marks nothing: that day is left for the employee to enter.",
        "A 0 in the sheet is only used where the CRM already has a 0 for that store and day. Anywhere else a 0 is ignored (it never replaces a number and never creates an entry), because a sheet made for the whole month has 0 in every day nobody has filled in yet. To record a real zero day, enter it on the store itself.",
        "A number from 1 to 200 is the bottles for that store and day. Existing entries are kept unless 'Overwrite existing entries' is ticked on the upload.",
        "Days in the future are ignored. In an Excel file every tab with these columns is read; other tabs (like this one) are ignored.",
        "The two example rows are invented. Delete them and paste your own stores.",
    ]
    return _build([(tab, [head, *sample], [10, 10, 9, 11, 13, 11, 28, 14, 13, 24, 18, 18, 18, 11, *[6] * len(days)])], notes), f"order-sheet-template-{month:%Y-%m}.xlsx"


def _col(n: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def _employees() -> tuple[bytes, str]:
    head = ["Employee ID", "Employee Name", "Password", "Region", "Status", "IsAdmin", "Platform", "IsAccountant", "Contact No"]
    rows = [head,
            ["EMP-1001", "Example Employee", "", "NORTH", "Active", "No", "Blinkit", "No", "9000000003"],
            ["EMP-1002", "Another Example", "", "SOUTH", "Active", "No", "Zepto", "No", "9000000004"]]
    assign = [["State", "Assigned Employee ID"], ["Karnataka", "EMP-1002"], ["Tamil Nadu", "EMP-1002"]]
    notes = [
        "Tab EMPLOYEES: one row per person. Employee ID and Employee Name are needed; the rest are optional. Headings are not case sensitive.",
        "Region and Platform must be ones that exist in the CRM (the Platform is Blinkit or Zepto). Status is Active or anything else for inactive. IsAdmin / IsAccountant: Yes or No.",
        "Contact No: a heading containing 'contact', 'phone' or 'mobile'.",
        "Password: leave it blank. Every new person gets a temporary password shown once after the upload (existing people's passwords are never touched).",
        "Tab State_Assignments (optional, for Zepto): which employee covers which state. Delete the tab if you do not need it. Any other tab is ignored.",
        "The example rows are invented. Delete them and paste your own people.",
    ]
    return _build([("EMPLOYEES", rows, [14, 22, 12, 11, 10, 10, 12, 13, 14]), ("State_Assignments", assign, [18, 22])], notes), "employees-template.xlsx"


def _attendance(month: date) -> tuple[bytes, str]:
    last = calendar.monthrange(month.year, month.month)[1]
    rows = [["Employee ID", "Name", *range(1, last + 1)],
            ["EMP-1001", "Example Employee", *(["P"] * 5 + [None] * (last - 5))],
            ["EMP-1002", "Another Example", *(["P", "P", None, "P"] + [None] * (last - 4))]]
    notes = [
        "One row per employee and one column per day of the month (1, 2, 3 ...). The month you are filling is the one you choose when you upload.",
        "Employee ID is how a row is matched. The Name column is only used when the ID cell is empty, and only if exactly one person has that name.",
        "Put P for a day the person was present. Only P writes anything: a blank, A or any other cell leaves that day exactly as it is.",
        "A day that already has a record is overwritten with the sheet's P (what it held before is kept in the activity log). Only days before today are written.",
        "The example rows are invented. Delete them and paste your own people.",
    ]
    return _build([(f"ATTENDANCE {month:%b-%Y}".upper(), rows, [14, 22, *[4] * last])], notes), f"attendance-template-{month:%Y-%m}.xlsx"


def _cash() -> tuple[bytes, str]:
    rows = [["Date", "Store Code", "Store / Place", "Type", "Reason", "Price per Item", "Amount (INR)"],
            [date.today().isoformat(), "BLK-1001", "Example Store One", "Store", "Cap", "5", "500"],
            [date.today().isoformat(), "", "Office", "Office", "Stationery", "", "250"]]
    notes = [
        "This is the same layout as the downloaded Cash Purchase Sheet, so that sheet can be uploaded back as it is.",
        "Needed: a Date, an Amount, and a Store Code (or a Store / Place name). Headings are not case sensitive; 'Outlet ID' also works for Store Code.",
        "Type: Store or Office. A row of type Office (or place 'Office') is not a store purchase and is counted separately.",
        "A purchase already in the CRM (same store, date and amount), or one uploaded before, is skipped, so uploading the same file twice never adds anything twice.",
        "Rows with a future date, no amount, or no store are skipped and counted with the reason. The example rows are invented: delete them.",
    ]
    return _build([("Cash purchases", rows, [12, 14, 26, 9, 18, 14, 14])], notes), "cash-purchase-template.xlsx"


def _cardsheet() -> tuple[bytes, str]:
    rows = [["Outlet", "Total"], ["Example Store One", 120], ["BLK-1002", 85]]
    notes = [
        "One row per outlet: the outlet (its name or its Outlet ID) and a total. A heading row is optional.",
        "On the page you pick the dates; each outlet's total is shared out at random across them, one standard card per outlet and month.",
        "Nothing is saved or added to any order or store. The example rows are invented: delete them.",
    ]
    return _build([("Outlets", rows, [30, 12])], notes), "random-card-sheet-template.xlsx"


def _stores() -> tuple[bytes, str]:
    rows = [STORE_SHEET_HEADER,
            [1, "Blinkit", "North Zone", "Delhi", "Delhi", "BLK-1001", "Blinkit - Example", "Live", None, None, None, "2026-03-01", "Sector 1, Rohini", "Store Manager", "9999999999", None, "Local Vendor", "8888888888", "Rohini"]]
    notes = [
        "The same columns as the platform's outlet sheet. Outlet ID and Outlet Name are needed; the rest are optional.",
        "STATUS is Live, Pending or Close. start Date is the day supply began (yyyy-mm-dd).",
        "Login and password columns are never read, so they are left out. The example row is invented: delete it.",
    ]
    return _build([("Stores", rows, [7, 10, 14, 12, 12, 12, 26, 9, 11, 14, 20, 12, 22, 16, 15, 20, 18, 16, 12])], notes), "stores-template.xlsx"


def _vendor_rates() -> tuple[bytes, str]:
    rows = [["City", "Rate", "Outlet ID", "Store Name"], ["Delhi", "4.50", None, None], [None, "5.00", "BLK-1001", None]]
    notes = [
        "A rate is per bottle. Give a City to set that city's rate, or an Outlet ID / Store Name to set one store's own rate (it overrides its city).",
        "A rate that differs from the current one asks for confirmation before it is applied. The example rows are invented: delete them.",
    ]
    return _build([("Rates", rows, [16, 10, 14, 26])], notes), "vendor-rates-template.xlsx"


def make(kind: str, month: str | None = None) -> tuple[bytes, str]:
    """(xlsx bytes, file name) of the template for `kind`."""
    if kind not in KINDS:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "There is no template for that.")
    m = _month(month)
    return {
        "orders": lambda: _orders(m), "employees": _employees, "attendance": lambda: _attendance(m), "cash": _cash,
        "cardsheet": _cardsheet, "stores": _stores, "vendor-rates": _vendor_rates,
    }[kind]()
