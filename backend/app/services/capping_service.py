"""
Zepto capping: how many bottles each store may take, and how close it is.

Who may use it (always re-checked on the server, whatever the app shows):
  * every ACTIVE employee whose platform is Zepto, whatever their state or region, and
  * administrators (they hold orders.correct and belong to no platform), so the Capping page works for them too.
  Blinkit employees (whatever else they hold) and partner logins are refused. Viewing and editing a store's cap are the same right.

The capping sheet gives ONE DAILY number per store (store_caps.daily_cap). The monthly limit a store must not exceed is that
number x 30 (always 30, whatever the month's length). No row = no cap; 0 is a real cap. Usage is the sum of the month's
order entries for the store, whoever entered them, so every employee sees the same figures. Alert levels:
  OVER   usage > monthly limit (daily x 30)
  AHEAD  usage is ahead of the pace: more than daily x day of the month (day 31 counts as 30) - current month only, not yet over
  NEAR   usage >= 80% of the monthly limit (and not over / ahead)        OK  below that        NO_CAP  no cap set
Example: sheet says 10 -> limit 300; on the 15th up to 150 is on pace, 160 marked gives an AHEAD alert.
Nothing limits a single day's entry: the only limit is the sheet's cap (or the one set by hand).

Who loads caps in bulk: developers and administrators (an uploaded sheet). Employees only change one store at a time.
Who is alerted: people holding capping.alerts - see alerts() - for the stores they cover.
"""
from __future__ import annotations

import calendar
import csv
import io
import re
import uuid
from datetime import date, datetime, timezone

from fastapi import HTTPException, status
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.order_entry import OrderEntry
from app.models.organization import Organization, OrganizationKind
from app.models.store import Store, StoreStatus
from app.models.store_cap import StoreCap
from app.models.user import User
from app.services import activity_service, assignment_service

PLATFORM = "zepto"
WARN_PERCENT = 80
CAP_DAYS = 30            # the monthly limit is the daily cap x 30, in every month
MAX_CAP = 1_000_000      # the largest daily cap accepted
_ORDER = {"OVER": 0, "AHEAD": 1, "NEAR": 2, "OK": 3, "NO_CAP": 4}
_MONTHS = ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def _today() -> date:
    return datetime.now(timezone.utc).date()


def is_eligible(user: User, perms: set[str]) -> bool:
    """Active Zepto employee, or an administrator. Never a Blinkit employee or a partner login."""
    if not user.is_active:
        return False
    slug = ((user.platform_organization.slug if user.platform_organization is not None else None) or "").strip().lower()
    if slug == PLATFORM:
        return True
    # Administrators (they hold orders.correct and belong to no platform). Someone who works on another platform (a Blinkit Team Lead, say) never sees Zepto
    # capping, even if "correct any entry" was ticked for them.
    return "orders.correct" in perms and not slug


def require_eligible(user: User, perms: set[str]) -> None:
    if not is_eligible(user, perms):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Capping is available to Zepto employees.")


def parse_month(value: str | None) -> date:
    """First day of the month. Default: this month. Future months are refused (there is nothing to report yet) and so is
    anything before the data starts."""
    today = _today()
    if not value:
        return date(today.year, today.month, 1)
    try:
        y, m = (int(x) for x in value.split("-"))
        d = date(y, m, 1)
    except (ValueError, TypeError):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Month must look like 2026-10.") from None
    if d > date(today.year, today.month, 1):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "That month has not started yet.")
    if d < date(2024, 1, 1):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "That month is too far back.")
    return d


def available_months(today: date | None = None) -> list[str]:
    today = today or _today()
    out, y, m = [], today.year, today.month
    for _ in range(12):
        out.append(f"{y}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return out


def monthly_limit(daily: int | None) -> int | None:
    """The most a store may take in a month: its daily cap x 30."""
    return None if daily is None else daily * CAP_DAYS


def alert_for(used: int, daily: int | None, day: int | None = None) -> str:
    """OVER / AHEAD / NEAR / OK / NO_CAP from the bottles used and the store's DAILY cap. Cap 0 is valid: zero used is OK,
    any use is OVER (no division). Pass `day` (the day of the month in progress) to also get AHEAD (used more than daily x day)."""
    if daily is None:
        return "NO_CAP"
    limit = daily * CAP_DAYS
    if used > limit:
        return "OVER"
    if limit == 0:
        return "OK"
    if day is not None and used > daily * min(day, CAP_DAYS):
        return "AHEAD"
    return "NEAR" if used * 100 >= WARN_PERCENT * limit else "OK"


def pace_for(daily: int | None, month: date, today: date) -> tuple[int | None, int | None]:
    """(day, bottles allowed so far) for the month in progress; (None, None) for a finished month or a store with no cap."""
    if daily is None or (month.year, month.month) != (today.year, today.month):
        return None, None
    return today.day, daily * min(today.day, CAP_DAYS)


def _partner(db: Session, org_id: uuid.UUID) -> Organization:
    p = db.execute(select(Organization).where(Organization.slug == PLATFORM, Organization.kind == OrganizationKind.PARTNER.value)).scalar_one_or_none()
    if p is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Zepto is not set up in the CRM.")
    return p


def _owners(db: Session, org_id: uuid.UUID, partner: Organization) -> dict[uuid.UUID, tuple[str, str]]:
    """{store_id: (employee code, employee name)}: the first Zepto employee (by name) whose states include the store.
    Same bulk lookup the dashboard uses: a few queries however many employees there are."""
    from app.repositories.employee_repository import EmployeeRepository

    employees = EmployeeRepository(db).all_for_platform(org_id, partner.id)
    scopes, _ = assignment_service.employee_store_scopes(db, partner, employees)
    out: dict[uuid.UUID, tuple[str, str]] = {}
    for emp in sorted(employees, key=lambda e: e.full_name.lower()):
        for sid in scopes.get(emp.id, []):
            out.setdefault(sid, (emp.employee_code, emp.full_name))
    return out


def _rows(db: Session, user: User, partner: Organization, month: date, store_ids: set[uuid.UUID] | None = None) -> list[dict]:
    """One row per live Zepto store (or only `store_ids`) with the month's usage and alert, worst first."""
    today = _today()
    end = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    q = (
        select(Store.id, Store.external_code, Store.name, Store.state, Store.city, StoreCap.daily_cap)
        .select_from(Store)
        .outerjoin(StoreCap, StoreCap.store_id == Store.id)
        .where(Store.organization_id == user.organization_id, Store.partner_organization_id == partner.id, Store.status == StoreStatus.LIVE.value)
    )
    if store_ids is not None:
        q = q.where(Store.id.in_(store_ids or {uuid.uuid4()}))
    stores = db.execute(q).all()
    usage = {
        r[0]: (int(r[1]), int(r[2]))
        for r in db.execute(
            select(OrderEntry.store_id, func.coalesce(func.sum(OrderEntry.bottle_count), 0), func.count(OrderEntry.id))
            .where(OrderEntry.order_date >= month, OrderEntry.order_date <= end, OrderEntry.store_id.in_([s[0] for s in stores] or [uuid.uuid4()]))
            .group_by(OrderEntry.store_id)
        )
    }
    owners = _owners(db, user.organization_id, partner)

    rows = []
    for sid, code, name, state, city, daily in stores:
        used, entries = usage.get(sid, (0, 0))
        day, allowed = pace_for(daily, month, today)
        level = alert_for(used, daily, day)
        cap = monthly_limit(daily)          # everything below is measured against the monthly limit (daily x 30)
        owner = owners.get(sid)
        rows.append({
            "store_id": sid, "outlet_id": code, "outlet_name": name, "state": state, "city": city,
            "owner_code": owner[0] if owner else None, "owner_name": owner[1] if owner else None,
            "daily_cap": daily, "monthly_cap": cap, "used": used, "entries": entries,
            "remaining": None if cap is None else max(cap - used, 0),
            "excess": 0 if cap is None else max(used - cap, 0),
            "percent": None if not cap else round(used * 100 / cap, 1),   # cap 0 or none: no percentage
            "allowed_to_date": allowed,                                     # None for a finished month / no cap
            "ahead_by": None if allowed is None else max(used - allowed, 0),
            "alert": level,
        })
    rows.sort(key=lambda r: (_ORDER[r["alert"]], -(r["percent"] if r["percent"] is not None else (10_000 if r["alert"] == "OVER" else -1)), r["outlet_name"].lower()))
    return rows


def _summary(rows: list[dict]) -> dict:
    counts = {k: sum(1 for r in rows if r["alert"] == k) for k in _ORDER}
    return {"over": counts["OVER"], "ahead": counts["AHEAD"], "near": counts["NEAR"], "ok": counts["OK"], "no_cap": counts["NO_CAP"], "total": len(rows)}


def report(db: Session, user: User, perms: set[str], month_str: str | None) -> dict:
    require_eligible(user, perms)
    month = parse_month(month_str)
    partner = _partner(db, user.organization_id)
    rows = _rows(db, user, partner, month)
    return {
        "success": True,
        "month": f"{month:%Y-%m}", "month_label": f"{_MONTHS[month.month]} {month.year}", "months": available_months(),
        "warn_percent": WARN_PERCENT,
        "summary": _summary(rows),
        "rows": rows,
    }


def alerts(db: Session, user: User, perms: set[str]) -> dict:
    """This month's stores that need the employee's attention (OVER or AHEAD), for people holding capping.alerts.
    An administrator sees every Zepto store; a Zepto employee only the stores they cover (their assigned states)."""
    if "capping.alerts" not in perms:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Capping alerts are not switched on for you.")
    require_eligible(user, perms)
    partner = _partner(db, user.organization_id)
    today = _today()
    month = date(today.year, today.month, 1)
    scope = None if "orders.correct" in perms else assignment_service.visible_store_ids(db, user)
    rows = [r for r in _rows(db, user, partner, month, scope) if r["alert"] in ("OVER", "AHEAD")]
    return {
        "success": True, "month": f"{month:%Y-%m}", "month_label": f"{_MONTHS[month.month]} {month.year}",
        "day": today.day, "days_in_month": CAP_DAYS,
        "over": sum(1 for r in rows if r["alert"] == "OVER"), "ahead": sum(1 for r in rows if r["alert"] == "AHEAD"),
        "rows": rows[:50],
    }


def set_cap(db: Session, user: User, perms: set[str], store_id: uuid.UUID, cap: int | None) -> dict:
    """Set (or clear, with None) one store's DAILY cap (its monthly limit is that x 30). Logged with the old and new value."""
    require_eligible(user, perms)
    partner = _partner(db, user.organization_id)
    store = db.get(Store, store_id)
    if store is None or store.organization_id != user.organization_id or store.partner_organization_id != partner.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Store not found.")
    if cap is not None and not (0 <= cap <= MAX_CAP):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"The daily cap must be a whole number from 0 to {MAX_CAP:,}.")
    row = db.execute(select(StoreCap).where(StoreCap.store_id == store_id)).scalar_one_or_none()
    before = row.daily_cap if row else None
    if before == cap:
        return {"store_id": store_id, "daily_cap": cap, "monthly_cap": monthly_limit(cap), "changed": False}
    if cap is None and row is not None:
        db.delete(row)
    elif row is None:
        db.add(StoreCap(store_id=store_id, daily_cap=cap, source="manual", updated_by_user_id=user.id))
    else:
        row.daily_cap, row.source, row.updated_by_user_id = cap, "manual", user.id
    if cap is not None:
        from app.services import capping_history_service    # the history is the record of when each value took effect

        capping_history_service.record_current_many(db, {store_id: cap}, user.id)
    activity_service.record(
        db, actor=user, action="store.cap_changed", entity_type="store", entity_id=store_id,
        metadata={"outlet": store.external_code, "name": store.name, "from": before, "to": cap},
    )
    db.commit()
    return {"store_id": store_id, "daily_cap": cap, "monthly_cap": monthly_limit(cap), "changed": True}


# ---------------------------------------------------------------- loading caps from the sheet / a file (developer)

def _cap_value(raw: str) -> tuple[int | None, bool]:
    """(cap, is_invalid). Blank -> (None, False): no cap in that row. Whole numbers only; 0 is valid."""
    text = (raw or "").strip().replace(",", "")
    if not text:
        return None, False
    try:
        f = float(text)
    except ValueError:
        return None, True
    if f != int(f) or not (0 <= f <= MAX_CAP):
        return None, True
    return int(f), False


_ID_HEADERS = {"outlet id", "outletid", "outlet code", "store id", "storeid", "store code", "outlet_id", "code", "id"}
_NAME_HEADERS = {"outlet name", "outletname", "store name", "storename", "outlet", "store", "name", "outlet_name"}
# The sheet's number is a DAILY cap whatever the column is called (older sheets are headed "Monthly Capping"): x 30 = the monthly limit.
_CAP_HEADERS = {"daily capping", "daily cap", "capping per day", "cap per day", "capping day", "capping", "cap", "monthly capping", "monthly cap", "limit", "daily limit"}
_HEADER_SCAN_ROWS = 25      # the header may sit under a title or a few blank rows
_MAX_SHEETS = 8


def _norm_name(text: str | None) -> str:
    """A store name for comparing: lower case, punctuation and repeated spaces ignored ("Zepto - Koramangala" = "zepto koramangala")."""
    return " ".join(re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).split())


def _norm_code(text: str | None) -> str:
    return re.sub(r"\s+", "", (text or "")).upper()


_ID_SET = {_norm_name(x) for x in _ID_HEADERS}
_NAME_SET = {_norm_name(x) for x in _NAME_HEADERS}
_CAP_SET = {_norm_name(x) for x in _CAP_HEADERS}


def _is_number(text: str) -> bool:
    value, invalid = _cap_value(text)
    return value is not None and not invalid


# ---------------------------------------------------------------- reading the file into plain rows of text

def rows_to_grid(rows: list[dict[str, str]]) -> list[list[str]]:
    """The dict rows the Google-sheet reader gives, as a grid again (header row first)."""
    if not rows:
        return []
    keys = list(rows[0].keys())
    return [list(keys)] + [[(r.get(k) or "") for k in keys] for r in rows]


def read_grids(filename: str, content: bytes) -> list[list[list[str]]]:
    """Every sheet of an upload as a grid of text cells (a .csv is one sheet). No assumption about where the header is."""
    from app.services import store_sync_service as sync

    if len(content) > sync.MAX_UPLOAD_BYTES:
        raise sync.SyncError("That file is too large (limit is 5 MB).")
    name = (filename or "").lower()

    def clean(cells) -> list[str]:
        out = [sync._cell_text(c).strip()[: sync.MAX_CELL_CHARS] for c in list(cells)[: sync.MAX_COLUMNS]]  # noqa: SLF001
        while out and not out[-1]:
            out.pop()
        return out

    def collect(rows) -> list[list[str]]:
        grid: list[list[str]] = []
        blanks = 0
        for cells in rows:
            r = clean(cells)
            blanks = 0 if any(r) else blanks + 1
            if blanks >= sync.BLANK_ROWS_STOP or len(grid) > sync.MAX_ROWS:
                break
            grid.append(r)
        while grid and not any(grid[-1]):
            grid.pop()
        return grid

    if name.endswith(".csv") or name.endswith(".txt"):
        return [collect(csv.reader(io.StringIO(content.decode("utf-8-sig", "replace"))))]
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
            raise sync.SyncError("Excel files need the 'openpyxl' package on the server.") from exc
        sync._check_unpacked_size(content)  # noqa: SLF001
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any parse failure is a bad file
            raise sync.SyncError("That file could not be read as an Excel workbook.") from exc
        try:
            return [collect(ws.iter_rows(values_only=True)) for ws in wb.worksheets[:_MAX_SHEETS]]
        finally:
            wb.close()
    raise sync.SyncError("Upload a .csv or .xlsx file.")


# ---------------------------------------------------------------- working out which column is which

def _resolve(code: str, name: str, by_code: dict[str, Store], by_name: dict[str, list[Store]]) -> tuple[Store | None, bool]:
    """(store, ambiguous). The Outlet Name decides (it is the store's identity). The Outlet ID only breaks a tie between stores that
    share a name, or stands in when the row has no name or the name is unknown."""
    if name:
        matches = by_name.get(_norm_name(name), [])
        if len(matches) == 1:
            return matches[0], False
        if len(matches) > 1:
            pick = [s for s in matches if code and _norm_code(s.external_code) == _norm_code(code)]
            return (pick[0], False) if len(pick) == 1 else (None, True)
    store = by_code.get(_norm_code(code)) if code else None
    return store, False


def _layouts(grid: list[list[str]], by_code: dict[str, Store], by_name: dict[str, list[Store]]) -> list[dict]:
    """Candidate readings of one sheet: which row is the header (if any) and which columns hold the store id, the store
    name and the daily cap. One by the header words, one by what the columns contain, so odd headers still work."""
    width = max((len(r) for r in grid), default=0)
    cell = lambda r, j: (grid[r][j] if j < len(grid[r]) else "").strip()  # noqa: E731

    # 1) by the header words
    header_row, id_col, name_col, cap_col = None, None, None, None
    for r in range(min(len(grid), _HEADER_SCAN_ROWS)):
        heads = [_norm_name(cell(r, j)) for j in range(width)]
        idc = next((j for j, h in enumerate(heads) if h in _ID_SET), None)
        nmc = next((j for j, h in enumerate(heads) if h in _NAME_SET), None)
        cpc = next((j for j, h in enumerate(heads) if h in _CAP_SET), None)
        if cpc is None:
            cpc = next((j for j, h in enumerate(heads) if re.search(r"\bcap(ping)?\b", h)), None)
        if (idc is not None or nmc is not None) and cpc is not None:
            header_row, id_col, name_col, cap_col = r, idc, nmc, cpc
            break
    layouts = []
    if header_row is not None:
        layouts.append({"header_row": header_row, "id": id_col, "name": name_col, "cap": cap_col})

    # 2) by what the columns contain (the header row is wherever the first line with text in the cap column is not a number)
    data_start = (header_row + 1) if header_row is not None else 0
    sample = [r for r in range(data_start, min(len(grid), data_start + 3000))]
    best_id, best_name = (None, 0.0), (None, 0.0)
    numeric: list[tuple[int, float]] = []
    for j in range(width):
        vals = [cell(r, j) for r in sample if cell(r, j)]
        if not vals:
            continue
        id_hits = sum(1 for v in vals if _norm_code(v) in by_code) / len(vals)
        name_hits = sum(1 for v in vals if _norm_name(v) in by_name) / len(vals)
        num_hits = sum(1 for v in vals if _is_number(v)) / len(vals)
        if id_hits > best_id[1]:
            best_id = (j, id_hits)
        if name_hits > best_name[1]:
            best_name = (j, name_hits)
        numeric.append((j, num_hits))
    id_j = best_id[0] if best_id[1] > 0 else None
    name_j = best_name[0] if best_name[1] > 0 and best_name[0] != id_j else None
    cap_candidates = [(h, j) for j, h in numeric if j not in (id_j, name_j) and h >= 0.5]
    cap_j = max(cap_candidates)[1] if cap_candidates else None      # the most numeric column (the right-most on a tie)
    if (id_j is not None or name_j is not None) and cap_j is not None:
        layouts.append({"header_row": header_row, "id": id_j, "name": name_j, "cap": cap_j, "guessed": True})
    return layouts


def _read_layout(grid: list[list[str]], lay: dict, by_code: dict[str, Store], by_name: dict[str, list[Store]]) -> dict:
    """The sheet's lines under this layout: (row number, id, name, cap text), and how many of them match a store."""
    cell = lambda r, j: (grid[r][j] if j is not None and j < len(grid[r]) else "").strip()  # noqa: E731
    start = (lay["header_row"] + 1) if lay["header_row"] is not None else 0
    lines, matched = [], 0
    for r in range(start, len(grid)):
        code, name, cap = cell(r, lay["id"]), cell(r, lay["name"]), cell(r, lay["cap"])
        if not (code or name):
            continue
        store, _ = _resolve(code, name, by_code, by_name)
        counts = store is not None and _is_number(cap)
        if (lay["header_row"] is None or lay.get("guessed")) and not matched and not counts:
            continue                                      # headings or a title above the first real store line
        lines.append((r + 1, code, name, cap))
        matched += 1 if counts else 0
    return {"lines": lines, "matched": matched, "layout": lay}


def import_caps(db: Session, actor: User, rows: list[dict[str, str]], *, overwrite: bool, source: str) -> dict:
    """Load caps from the dict rows of the Google capping tab."""
    return import_grids(db, actor, [rows_to_grid(rows)], overwrite=overwrite, source=source)


def import_grids(db: Session, actor: User, grids: list[list[list[str]]], *, overwrite: bool, source: str) -> dict:
    """Load caps from a sheet, whatever it looks like: the header can be on any of the first rows, the columns can be called
    anything (they are recognised by their words or by what they hold), a store is named by Outlet ID or Outlet Name
    (matched ignoring case, punctuation and spacing; a name two stores share is skipped and reported, never guessed), and an
    Excel file may have several sheets (the one that matches the most stores is used). The number is the DAILY cap; the
    monthly limit is that x 30. Existing caps are kept unless `overwrite` is on."""
    partner = _partner(db, actor.organization_id)
    res = {"rows_read": 0, "added": 0, "updated": 0, "kept": 0, "unknown_stores": 0, "ambiguous": 0, "invalid": 0, "warnings": [], "read_as": None}
    all_stores = list(db.execute(select(Store).where(Store.organization_id == actor.organization_id, Store.partner_organization_id == partner.id)).scalars())
    by_code = {_norm_code(s.external_code): s for s in all_stores}
    by_name: dict[str, list[Store]] = {}
    for st in all_stores:
        by_name.setdefault(_norm_name(st.name), []).append(st)

    best = None
    for grid in grids:
        for lay in _layouts(grid, by_code, by_name):
            got = _read_layout(grid, lay, by_code, by_name)
            if best is None or got["matched"] > best["matched"]:
                best = got
    if best is None or not best["lines"]:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "No stores with a capping number were found in this file. It needs a column of store names (or Outlet IDs) and a column with the daily cap.",
        )
    lay = best["layout"]
    by = "Outlet ID and name" if lay["id"] is not None and lay["name"] is not None else "Outlet ID" if lay["id"] is not None else "Outlet Name"
    res["read_as"] = f"stores matched by {by}; columns found from {'what they contain' if lay.get('guessed') else 'their headings'}"

    existing = {c.store_id: c for c in db.execute(select(StoreCap).where(StoreCap.store_id.in_([s.id for s in all_stores] or [uuid.uuid4()]))).scalars()}
    changed: dict[uuid.UUID, int] = {}          # caps this upload added or replaced, for the capping history
    seen: set[tuple[str, str]] = set()
    bad_rows: list[int] = []
    unknown_first: str | None = None
    ambiguous_names: list[str] = []
    for rownum, code, name, raw in best["lines"]:
        key = ("id", _norm_code(code)) if code else ("name", _norm_name(name))
        if not key[1] or key in seen:
            continue
        seen.add(key)
        res["rows_read"] += 1
        cap, invalid = _cap_value(raw)
        if invalid:
            res["invalid"] += 1
            bad_rows.append(rownum)
            continue
        if cap is None:
            continue
        store, ambiguous = _resolve(code, name, by_code, by_name)
        if ambiguous:
            res["ambiguous"] += 1
            if len(ambiguous_names) < 5:
                ambiguous_names.append(name)
            continue
        if store is None:
            res["unknown_stores"] += 1
            unknown_first = unknown_first or (code or name)
            continue
        cur = existing.get(store.id)
        if cur is None:
            row_new = StoreCap(store_id=store.id, daily_cap=cap, source=source, updated_by_user_id=actor.id)
            db.add(row_new)
            existing[store.id] = row_new
            changed[store.id] = cap
            res["added"] += 1
        elif cur.daily_cap == cap:
            res["kept"] += 1
        elif overwrite:
            cur.daily_cap, cur.source, cur.updated_by_user_id = cap, source, actor.id
            changed[store.id] = cap
            res["updated"] += 1
        else:
            res["kept"] += 1
    if bad_rows:
        res["warnings"].append(f"{res['invalid']} row(s) have a cap that is not a whole number from 0 to {MAX_CAP:,} and were skipped (first: row {bad_rows[0]}).")
    if res["unknown_stores"]:
        res["warnings"].append(f"{res['unknown_stores']} store(s) in the sheet are not Zepto stores in the CRM and were skipped (first: {unknown_first}).")
    if res["ambiguous"]:
        res["warnings"].append(f"{res['ambiguous']} store name(s) match more than one store and were skipped, add an Outlet ID column to settle them: {', '.join(ambiguous_names)}.")
    if res["kept"] and not overwrite:
        res["warnings"].append("Caps already set in the CRM were kept (turn on Overwrite to replace them with the sheet's values).")
    from app.services import capping_history_service

    capping_history_service.record_current_many(db, changed, actor.id)
    activity_service.record(
        db, actor=actor, action="store.caps_imported", entity_type="store", entity_id=partner.id,
        metadata={"source": source, "overwrite": overwrite, "added": res["added"], "updated": res["updated"], "kept": res["kept"]},
    )
    db.commit()
    return res
