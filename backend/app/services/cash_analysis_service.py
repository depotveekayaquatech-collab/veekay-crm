"""
Vendor analysis of store cash purchases (developer only): which vendors' stores needed the most cash purchases.

A store cash purchase means the vendor didn't deliver in time (or the store was new / short), so a vendor whose stores
keep showing up here is the one to chase. Vendors are ranked by number of cash purchases and flagged:

    HIGH   at least 3 purchases and at least 1.5x the average vendor
    WATCH  at least 2 purchases and above the average vendor

The average is taken over vendors that had any cash purchase in the period.
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from datetime import date
from decimal import Decimal

from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.cash_purchase import STORE_REASONS, CashPurchase, PurchaseKind
from app.models.store import Store, StoreStatus
from app.models.user import User

NO_VENDOR = "(No vendor on record)"
HIGH_MIN, HIGH_FACTOR, WATCH_MIN = 3, 1.5, 2


class StoreLine(BaseModel):
    store_id: str
    outlet_code: str
    outlet_name: str
    platform: str | None
    state: str | None
    city: str | None
    purchases: int
    amount: Decimal
    last_date: date
    top_reason: str | None


class VendorLine(BaseModel):
    rank: int
    flag: str                     # "HIGH" | "WATCH" | ""
    vendor: str
    vendor_number: str | None
    stores_with_purchases: int
    total_stores: int
    purchases: int
    amount: Decimal
    average_amount: Decimal
    last_date: date
    reasons: dict[str, int]
    stores: list[StoreLine]


class Analysis(BaseModel):
    start: date | None
    end: date | None
    total_purchases: int
    total_amount: Decimal
    vendors_count: int
    average_purchases: float
    high: int
    watch: int
    vendors: list[VendorLine]


def _key(name: str | None) -> str:
    return " ".join((name or "").split()).casefold()


def analyse(db: Session, user: User, start: date | None, end: date | None) -> Analysis:
    stmt = select(CashPurchase).where(
        CashPurchase.organization_id == user.organization_id,
        CashPurchase.kind == PurchaseKind.STORE.value,
        CashPurchase.store_id.is_not(None),
    )
    if start:
        stmt = stmt.where(CashPurchase.purchase_date >= start)
    if end:
        stmt = stmt.where(CashPurchase.purchase_date <= end)
    purchases = db.execute(stmt).scalars().unique().all()

    live = db.execute(
        select(Store.vendor_name, func.count()).where(
            Store.organization_id == user.organization_id, Store.status == StoreStatus.LIVE.value
        ).group_by(Store.vendor_name)
    ).all()
    stores_per_vendor: Counter = Counter()
    for name, n in live:
        stores_per_vendor[_key(name)] += n

    by_vendor: dict[str, list[CashPurchase]] = defaultdict(list)
    for p in purchases:
        if p.store is not None:
            by_vendor[_key(p.store.vendor_name)].append(p)

    counts = [len(v) for v in by_vendor.values()]
    avg = sum(counts) / len(counts) if counts else 0.0

    lines: list[VendorLine] = []
    for k, items in by_vendor.items():
        n = len(items)
        total = sum((p.amount for p in items), Decimal(0))
        flag = "HIGH" if n >= HIGH_MIN and n >= HIGH_FACTOR * avg else ("WATCH" if n >= WATCH_MIN and n > avg else "")
        by_store: dict = defaultdict(list)
        for p in items:
            by_store[p.store_id].append(p)
        stores = []
        for sid, ps in by_store.items():
            s = ps[0].store
            top = Counter(p.category for p in ps).most_common(1)[0][0]
            stores.append(StoreLine(
                store_id=str(sid), outlet_code=s.external_code, outlet_name=s.name,
                platform=s.partner_organization.name if s.partner_organization else None, state=s.state, city=s.city,
                purchases=len(ps), amount=sum((p.amount for p in ps), Decimal(0)),
                last_date=max(p.purchase_date for p in ps), top_reason=STORE_REASONS.get(top, top.title()),
            ))
        stores.sort(key=lambda x: (-x.purchases, x.outlet_name.lower()))
        first = items[0].store
        reasons = Counter(STORE_REASONS.get(p.category, p.category.title()) for p in items)
        lines.append(VendorLine(
            rank=0, flag=flag, vendor=" ".join((first.vendor_name or "").split()) or NO_VENDOR, vendor_number=first.vendor_number,
            stores_with_purchases=len(stores), total_stores=stores_per_vendor.get(k, 0), purchases=n, amount=total,
            average_amount=(total / n).quantize(Decimal("0.01")), last_date=max(p.purchase_date for p in items),
            reasons=dict(reasons.most_common()), stores=stores,
        ))
    lines.sort(key=lambda v: (-v.purchases, -v.amount, v.vendor.lower()))
    for i, v in enumerate(lines, 1):
        v.rank = i
    return Analysis(
        start=start, end=end, total_purchases=len(purchases), total_amount=sum((p.amount for p in purchases), Decimal(0)),
        vendors_count=len(lines), average_purchases=round(avg, 2), high=sum(v.flag == "HIGH" for v in lines),
        watch=sum(v.flag == "WATCH" for v in lines), vendors=lines,
    )


def analysis_xlsx(db: Session, user: User, start: date | None, end: date | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    a = analyse(db, user, start, end)
    wb = Workbook()
    side = Side(style="thin")
    fills = {"HIGH": PatternFill("solid", fgColor="F8CBAD"), "WATCH": PatternFill("solid", fgColor="FFE699")}

    def head(ws):
        for c in ws[1]:
            c.font = Font(name="Calibri", size=11, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            c.border = Border(left=side, right=side, top=side, bottom=side)
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Vendor summary"
    ws.append(["Rank", "Flag", "Vendor", "Vendor number", "Cash purchases", "Total amount (INR)", "Avg per purchase (INR)",
               "Stores with cash purchases", "Total live stores", "Last purchase", "Main reasons", "Store names"])
    head(ws)
    for v in a.vendors:
        ws.append([
            v.rank, v.flag or None, v.vendor, v.vendor_number, v.purchases, float(v.amount), float(v.average_amount),
            v.stores_with_purchases, v.total_stores, v.last_date,
            ", ".join(f"{k} ({n})" for k, n in v.reasons.items()),
            ", ".join(f"{s.outlet_name} ({s.purchases})" for s in v.stores),
        ])
        row = ws.max_row
        if v.flag:
            for c in ws[row][:3]:
                c.fill = fills[v.flag]
        ws.cell(row, 10).number_format = "dd-mmm-yyyy"
        ws.cell(row, 6).number_format = ws.cell(row, 7).number_format = "#,##0.00"
        ws.cell(row, 11).alignment = ws.cell(row, 12).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKL", (6, 8, 30, 15, 11, 15, 14, 13, 11, 13, 40, 70)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Stores by vendor")
    ws2.append(["Vendor", "Flag", "Store ID", "Store name", "Platform", "State", "City", "Cash purchases", "Amount (INR)", "Last purchase", "Main reason"])
    head(ws2)
    for v in a.vendors:
        for s in v.stores:
            ws2.append([v.vendor, v.flag or None, s.outlet_code, s.outlet_name, s.platform, s.state, s.city, s.purchases,
                        float(s.amount), s.last_date, s.top_reason])
            row = ws2.max_row
            ws2.cell(row, 9).number_format = "#,##0.00"
            ws2.cell(row, 10).number_format = "dd-mmm-yyyy"
            if v.flag:
                ws2.cell(row, 1).fill = ws2.cell(row, 2).fill = fills[v.flag]
    for col, w in zip("ABCDEFGHIJK", (30, 8, 14, 32, 12, 16, 16, 11, 14, 13, 30)):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Notes")
    period = f"{start:%d-%b-%Y} to {end:%d-%b-%Y}" if start and end else ("from " + f"{start:%d-%b-%Y}" if start else ("up to " + f"{end:%d-%b-%Y}" if end else "all time"))
    for line in (
        "Vendor analysis of store cash purchases",
        f"Period: {period}",
        f"Cash purchases counted: {a.total_purchases} for ₹{a.total_amount:,.2f} across {a.vendors_count} vendors",
        f"Average cash purchases per vendor (vendors with any): {a.average_purchases}",
        f"HIGH = at least {HIGH_MIN} purchases and at least {HIGH_FACTOR}x the average. WATCH = at least {WATCH_MIN} purchases and above the average.",
        "Only store purchases are counted (office purchases are not). The vendor is the vendor recorded on the store.",
    ):
        ws3.append([line])
    ws3["A1"].font = Font(bold=True, size=12)
    ws3.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
