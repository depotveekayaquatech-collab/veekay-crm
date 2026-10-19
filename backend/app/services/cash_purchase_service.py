"""
Cash purchases: who may record / see what.

  cash.add   record a purchase; sees only the purchases they recorded themselves
  cash.view  (admin, accounts) sees every purchase and downloads the sheet

Store purchases may only name a store the person can see (their own stores; any live store for cash.view).
"""
from __future__ import annotations

import hashlib
import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

from fastapi import HTTPException, status
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.core import storage
from app.models.cash_purchase import (
    OFFICE_REASONS,
    OTHER,
    REASONS_BY_KIND,
    STORE_ITEM,
    STORE_REASONS,
    CashPurchase,
    PurchaseKind,
)
from app.models.store import Store, StoreStatus
from app.models.user import User
from app.schemas.cash_purchase import (
    CashPurchaseCreate,
    CashPurchaseList,
    CashPurchaseOptions,
    CashPurchaseOut,
    Reason,
    StoreChoice,
)
from app.services import activity_service, assignment_service
from app.services.compliance_service import _CONTENT_TYPES, _prepare

class DuplicatePurchase(Exception):
    """The purchase looks like one already recorded. Answered as 409 with `detail` (readable text, so an older app can show it)
    plus `code` and `matches` for an app that can ask the person (see main.py)."""

    def __init__(self, message: str, matches: list[dict]):
        super().__init__(message)
        self.message = message
        self.matches = matches


MAX_AMOUNT = Decimal("1000000")
MAX_EXPORT_ROWS = 20000
# A late entry is fine, but not one from last year.
BACKDATE_DAYS = 62
_IST = timezone(timedelta(hours=5, minutes=30))


def _today() -> date:
    return datetime.now(_IST).date()


def _can_view_all(perms: set[str]) -> bool:
    return "cash.view" in perms


def _stores_for(db: Session, user: User, perms: set[str]) -> list[Store]:
    if _can_view_all(perms):
        rows = db.execute(
            select(Store).where(Store.organization_id == user.organization_id, Store.status == StoreStatus.LIVE.value)
        ).scalars().all()
    else:
        rows = assignment_service.visible_stores(db, user)
    return sorted(rows, key=lambda s: (s.name or "").lower())


def options(db: Session, user: User, perms: set[str]) -> CashPurchaseOptions:
    return CashPurchaseOptions(
        office_reasons=[Reason(code=c, label=l) for c, l in OFFICE_REASONS.items()],
        store_reasons=[Reason(code=c, label=l) for c, l in STORE_REASONS.items()],
        stores=[
            StoreChoice(
                id=s.id, name=s.name, code=s.external_code, state=s.state,
                platform=s.partner_organization.name if s.partner_organization else None,
            )
            for s in _stores_for(db, user, perms)
        ],
        can_view_all=_can_view_all(perms),
        max_amount=MAX_AMOUNT,
    )


def _label(kind: str, category: str) -> str:
    return REASONS_BY_KIND.get(kind, {}).get(category, category.replace("_", " ").title())


def _out(p: CashPurchase) -> CashPurchaseOut:
    return CashPurchaseOut(
        id=p.id, kind=p.kind, purchase_date=p.purchase_date,
        store_id=p.store_id, store_name=p.store.name if p.store else None,
        store_code=p.store.external_code if p.store else None,
        category=p.category, category_label=_label(p.kind, p.category), other_reason=p.other_reason,
        amount=p.amount, notes=p.notes, has_proof=bool(p.proof_key), proof_name=p.proof_name,
        added_by=p.created_by.full_name if p.created_by else None, created_at=p.created_at,
        possible_duplicate=p.duplicate_of_id is not None,
    )


def _match_out(p: CashPurchase, user: User, perms: set[str], same_details: bool, same_file: bool) -> dict:
    """What the person is shown about the earlier purchase. Another employee's name goes only to someone who may see everything."""
    who = "You" if p.created_by_user_id == user.id else (p.created_by.full_name if p.created_by and _can_view_all(perms) else "Another employee")
    return {
        "id": str(p.id), "purchase_date": p.purchase_date.isoformat(), "amount": str(p.amount), "kind": p.kind,
        "store_name": p.store.name if p.store else None, "category_label": _label(p.kind, p.category),
        "added_by": who, "created_at": p.created_at.isoformat(), "same_details": same_details, "same_file": same_file,
    }


def _clean(v: str | None) -> str | None:
    v = (v or "").strip()
    return v or None


def _find_duplicates(
    db: Session, org_id: uuid.UUID, *, kind: str, store_id: uuid.UUID | None, purchase_date: date, category: str,
    other_reason: str | None, amount: Decimal, proof_sha256: str,
) -> list[tuple[CashPurchase, bool, bool]]:
    """Earlier purchases that look like this one: (purchase, same details, same bill file). Everyone's purchases count, not
    only the caller's, because two people can record the same bill."""
    same = and_(
        CashPurchase.kind == kind, CashPurchase.purchase_date == purchase_date, CashPurchase.category == category,
        CashPurchase.amount == amount,
        CashPurchase.store_id == store_id if store_id else CashPurchase.store_id.is_(None),
    )
    rows = db.execute(
        select(CashPurchase).where(CashPurchase.organization_id == org_id, or_(same, CashPurchase.proof_sha256 == proof_sha256))
        .order_by(CashPurchase.created_at.desc()).limit(10)
    ).scalars().unique().all()
    out = []
    for p in rows:
        details = (
            p.kind == kind and p.purchase_date == purchase_date and p.category == category and p.amount == amount
            and p.store_id == store_id
            and (category != OTHER or (p.other_reason or "").strip().lower() == (other_reason or "").strip().lower())
        )
        file_same = p.proof_sha256 == proof_sha256
        if details or file_same:
            out.append((p, details, file_same))
    return out[:5]


def _duplicate_message(found: list[tuple[CashPurchase, bool, bool]]) -> str:
    _, details, file_same = found[0]
    what = (
        "This looks exactly like a purchase already recorded (same details and the same bill photo)." if details and file_same
        else "This bill photo was already used for another purchase." if file_same
        else "A purchase with the same date, amount and reason is already recorded."
    )
    return f"{what} Please check it. Is it the same purchase, or a different one?"


def create(
    db: Session, user: User, perms: set[str], body: CashPurchaseCreate, proof_name: str, proof: bytes,
    confirm_duplicate: bool = False,
) -> CashPurchaseOut:
    today = _today()
    if body.purchase_date > today:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "The purchase date can't be in the future.")
    if body.purchase_date < today - timedelta(days=BACKDATE_DAYS):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"Purchases older than {BACKDATE_DAYS} days can't be added.")
    if body.amount > MAX_AMOUNT:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"The amount can't be more than {MAX_AMOUNT:,.0f}.")

    if body.category not in REASONS_BY_KIND[body.kind]:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Pick a reason from the list.")
    other = _clean(body.other_reason)
    if body.category == OTHER:
        if not other or len(other) < 3:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Describe what the purchase was for.")
    else:
        other = None

    # Payment proof: one photo (JPG / PNG / WEBP) or a PDF, checked by content, not by file name.
    proof_bytes, ext = _prepare([(proof_name or "payment proof", proof)])

    store_id = None
    if body.kind == PurchaseKind.STORE.value:
        if body.store_id is None:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Pick the store this purchase was for.")
        if body.store_id not in {s.id for s in _stores_for(db, user, perms)}:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "You can't record purchases for that store.")
        store_id = body.store_id

    proof_sha256 = hashlib.sha256(proof).hexdigest()
    found = _find_duplicates(
        db, user.organization_id, kind=body.kind, store_id=store_id, purchase_date=body.purchase_date, category=body.category,
        other_reason=other, amount=body.amount, proof_sha256=proof_sha256,
    )
    if found and not confirm_duplicate:
        # keep a record that the warning was shown, then ask: nothing is saved yet
        activity_service.record(
            db, actor=user, action="cash_purchase.duplicate_flagged", entity_type="cash_purchase", entity_id=found[0][0].id,
            metadata={"amount": str(body.amount), "date": body.purchase_date.isoformat(), "matches": [str(p.id) for p, _, _ in found]},
        )
        db.commit()
        raise DuplicatePurchase(_duplicate_message(found), [_match_out(p, user, perms, d, f) for p, d, f in found])

    row = CashPurchase(
        organization_id=user.organization_id, kind=body.kind, store_id=store_id,
        purchase_date=body.purchase_date, category=body.category, other_reason=other,
        amount=body.amount, notes=_clean(body.notes), created_by_user_id=user.id,
        proof_sha256=proof_sha256, duplicate_of_id=found[0][0].id if found else None,
    )
    db.add(row)
    db.flush()
    row.proof_key = f"Cash purchases/{body.purchase_date:%Y-%m}/{row.id}.{ext}"
    row.proof_name = f"cash-purchase-{body.purchase_date.isoformat()}-{str(row.id)[:8]}.{ext}"
    row.proof_content_type = _CONTENT_TYPES[ext]
    row.proof_size_bytes = len(proof_bytes)
    storage.save_verified(row.proof_key, proof_bytes)
    storage.log.info("file stored", extra={"category": "cash.proof", "file_key": row.proof_key, "size_out": len(proof_bytes)})
    activity_service.record(
        db, actor=user, action="cash_purchase.created", entity_type="cash_purchase", entity_id=row.id,
        metadata={"kind": row.kind, "category": row.category, "amount": str(row.amount),
                  "store_id": str(store_id) if store_id else None,
                  **({"confirmed_different_from": [str(p.id) for p, _, _ in found]} if found else {})},
    )
    db.commit()
    db.refresh(row)
    return _out(row)


def _filtered(user: User, perms: set[str], *, kind, start, end, category, store_id, q):
    stmt = select(CashPurchase).where(CashPurchase.organization_id == user.organization_id)
    if not _can_view_all(perms):
        stmt = stmt.where(CashPurchase.created_by_user_id == user.id)
    if kind:
        stmt = stmt.where(CashPurchase.kind == kind)
    if start:
        stmt = stmt.where(CashPurchase.purchase_date >= start)
    if end:
        stmt = stmt.where(CashPurchase.purchase_date <= end)
    if category:
        stmt = stmt.where(CashPurchase.category == category)
    if store_id:
        stmt = stmt.where(CashPurchase.store_id == store_id)
    if q and q.strip():
        raw = q.strip()
        for ch in ("\\", "%", "_"):
            raw = raw.replace(ch, "\\" + ch)
        like = f"%{raw}%"
        stmt = stmt.where(or_(
            CashPurchase.notes.ilike(like, escape="\\"),
            CashPurchase.other_reason.ilike(like, escape="\\"),
        ))
    return stmt


def listing(
    db: Session, user: User, perms: set[str], *, page: int, page_size: int,
    kind: str | None, start: date | None, end: date | None, category: str | None,
    store_id: uuid.UUID | None, q: str | None,
) -> CashPurchaseList:
    base = _filtered(user, perms, kind=kind, start=start, end=end, category=category, store_id=store_id, q=q)
    sub = base.order_by(None).with_only_columns(CashPurchase.kind, CashPurchase.amount).subquery()
    total, amount_total, office_total = db.execute(
        select(
            func.count(), func.coalesce(func.sum(sub.c.amount), 0),
            func.coalesce(func.sum(sub.c.amount).filter(sub.c.kind == PurchaseKind.OFFICE.value), 0),
        )
    ).one()
    rows = db.execute(
        base.order_by(CashPurchase.purchase_date.desc(), CashPurchase.created_at.desc())
        .offset((page - 1) * page_size).limit(page_size)
    ).scalars().unique().all()
    return CashPurchaseList(
        items=[_out(r) for r in rows], total=total, page=page, page_size=page_size,
        amount_total=amount_total, office_total=office_total, store_total=amount_total - office_total,
    )


def export_xlsx(
    db: Session, user: User, perms: set[str], *,
    kind: str | None, start: date | None, end: date | None, category: str | None,
    store_id: uuid.UUID | None, q: str | None,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    stmt = _filtered(user, perms, kind=kind, start=start, end=end, category=category, store_id=store_id, q=q)
    rows = db.execute(
        stmt.order_by(CashPurchase.purchase_date, CashPurchase.created_at).limit(MAX_EXPORT_ROWS)
    ).scalars().unique().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Cash purchases"
    ws.append(["Date", "Type", "Store / Place", "Store code", "State", "Item", "Reason", "Details",
               "Proof attached", "Amount (INR)", "Added by", "Added on", "Possible duplicate"])
    side = Side(style="thin")
    for c in ws[1]:
        c.font = Font(name="Calibri", size=11, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="top")
        c.border = Border(left=side, right=side, top=side, bottom=side)

    total = Decimal(0)
    for p in rows:
        total += p.amount
        details = " — ".join(x for x in (p.other_reason, p.notes) if x)
        ws.append([
            p.purchase_date, "Office" if p.kind == PurchaseKind.OFFICE.value else "Store",
            p.store.name if p.store else "Office", p.store.external_code if p.store else None,
            p.store.state if p.store else None,
            STORE_ITEM if p.store_id else _label(p.kind, p.category),
            _label(p.kind, p.category) if p.store_id else None, details or None,
            "Yes" if p.proof_key else "No", float(p.amount), p.created_by.full_name if p.created_by else None,
            p.created_at.astimezone(_IST).strftime("%Y-%m-%d %H:%M"),
            "Yes - recorded as different" if p.duplicate_of_id else None,
        ])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = "dd-mmm-yyyy"
        ws.cell(r, 10).number_format = "#,##0.00"
    ws.append([])
    ws.append(["TOTAL", None, None, None, None, None, None, None, None, float(total)])
    last = ws.max_row
    bold = Font(name="Calibri", size=11, bold=True)
    ws.cell(last, 1).font = bold
    ws.cell(last, 10).font = bold
    ws.cell(last, 10).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKLM", (13, 9, 30, 12, 16, 22, 30, 36, 22, 14, 22, 17, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_proof(db: Session, user: User, perms: set[str], purchase_id: uuid.UUID) -> tuple[bytes, str, str]:
    """(bytes, content_type, file_name) of the payment proof. cash.view: any; otherwise only your own purchase."""
    row = db.get(CashPurchase, purchase_id)
    if row is None or row.organization_id != user.organization_id or (
        not _can_view_all(perms) and row.created_by_user_id != user.id
    ):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Purchase not found.")
    if not row.proof_key or not storage.exists(row.proof_key):
        if row.proof_key:
            storage.log_missing(row.proof_key, "cash_purchases", row.id)
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No payment proof is stored for this purchase.")
    return storage.read(row.proof_key), row.proof_content_type or "application/octet-stream", row.proof_name or "proof"
