"""
Import / refresh stores from the partner outlet-master Google Sheets.

Each sheet is read through its public CSV export endpoint
(https://docs.google.com/spreadsheets/d/<id>/export?format=csv&gid=<gid>),
so the sheet must be shared as "Anyone with the link - Viewer" or Published
to web. No Google credentials are used or stored.

Match key: (partner_organization, external_code) == (platform, Outlet ID).
Rows are upserted; a store already in the CRM but absent from the sheet is
left untouched (the sheet only adds / updates, never deletes).
Credential columns in the sheet (logins / passwords) are deliberately ignored.
"""
from __future__ import annotations

import csv
import io
import re
import urllib.error
import urllib.request
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.organization import Organization, OrganizationKind
from app.models.region import Region
from app.models.store import Store, StoreStatus
from app.models.user import User
from app.repositories.store_repository import StoreRepository
from app.services import activity_service, assignment_service

_CSV_URL = "https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"

# Sheet header -> Store field. Header lookup is case-insensitive and
# whitespace-trimmed; the first alias that is present wins.
_COLUMN_ALIASES: dict[str, list[str]] = {
    "external_code": ["outlet id", "outlet_id", "store id", "store code", "outlet code", "id"],
    "name": ["outlet name", "store name", "outlet", "name"],
    "region": ["zone", "region"],
    "entity": ["entity"],
    # Veekay's own start date wins; fall back to the partner's go-live date column.
    "start_date": ["start date", "start_date", "start date by blinkit", "go-live date", "go live date", "live date"],
    "state": ["state"],
    "city": ["city"],
    "address": ["address", "location"],
    "status": ["status"],
    "poc_name": ["poc name", "poc", "store manager name", "store maneger name", "store manager", "manager name"],
    "poc_number": ["poc contact no", "poc contact", "poc number", "poc contact number", "poc no",
                   "store manager number", "manager number", "number", "phone", "mobile"],
    "vendor_name": ["vendor", "vendor name"],
    "vendor_number": ["contact details", "vendor contact", "vendor number", "vendor contact no", "contact number"],
}

# Order matters: "not live yet" / "inactive" must be caught before the bare
# words "live" / "active" would claim them. Whole-word matches only, so
# e.g. "deliver" never counts as "live".
_PENDING_FIRST = re.compile(
    r"\b(yet to|not (yet )?(live|started|active)|to go live|to be live|awaiting|await|pending|"
    r"on[- ]?hold|hold|in[- ]?progress|in process|onboarding|upcoming)\b"
)
_CLOSE_RE = re.compile(
    r"\b(closed?|inactive|stopped|stop|churn(ed)?|terminated|discontinued|dead|lost|shut( down)?)\b"
)
_LIVE_RE = re.compile(r"\b(live|active|running|working|operational|go[- ]?live)\b")


def normalize_status(raw: str | None) -> str:
    """Free-text sheet status -> LIVE / PENDING / CLOSE. Unknown -> PENDING so
    it surfaces for review rather than silently going live."""
    if not raw:
        return StoreStatus.PENDING.value
    text = raw.strip().lower()
    if _PENDING_FIRST.search(text):
        return StoreStatus.PENDING.value
    if _CLOSE_RE.search(text):
        return StoreStatus.CLOSE.value
    if _LIVE_RE.search(text):
        return StoreStatus.LIVE.value
    return StoreStatus.PENDING.value


@dataclass
class SyncResult:
    platform: str
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    rows_read: int = 0
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "platform": self.platform,
            "created": self.created,
            "updated": self.updated,
            "unchanged": self.unchanged,
            "rows_read": self.rows_read,
            "warnings": self.warnings[:50],
        }


# --------------------------------------------------------------------------


def _fetch_rows(sheet_id: str, gid: str) -> list[dict[str, str]]:
    url = _CSV_URL.format(sheet_id=sheet_id, gid=gid)
    req = urllib.request.Request(url, headers={"User-Agent": "veekay-crm-store-sync/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:  # noqa: S310 - fixed docs.google.com host
            body = resp.read().decode("utf-8-sig", "replace")
    except urllib.error.HTTPError as exc:
        if exc.code in (401, 403):
            raise SyncError(
                "The sheet is not publicly readable. Share it as "
                "'Anyone with the link - Viewer' (or File - Share - Publish to web)."
            ) from exc
        raise SyncError(f"Google Sheets returned HTTP {exc.code} for this sheet.") from exc
    except urllib.error.URLError as exc:
        raise SyncError(f"Could not reach Google Sheets: {exc.reason}") from exc

    return _parse_csv_text(body)


def _parse_csv_text(body: str) -> list[dict[str, str]]:
    reader = csv.reader(io.StringIO(body))
    try:
        header = next(reader)
    except StopIteration:
        return []
    keys = [h.strip() for h in header]
    return [dict(zip(keys, row)) for row in reader if any(c.strip() for c in row)]


MAX_UPLOAD_BYTES = 5 * 1024 * 1024

# Column widths from the Store model. Real sheets have long free-text cells
# (multi-line addresses, several phone numbers in one cell); trim instead of
# letting one cell fail the whole import with a database error.
_FIELD_LIMITS = {
    "name": 128, "entity": 64, "state": 64, "city": 64, "address": 255,
    "poc_name": 128, "poc_number": 32, "vendor_name": 128, "vendor_number": 32,
}
_CODE_LIMIT = 64


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # phone numbers / ids Excel stored as floats
    return str(value)


def _rows_from_xlsx(content: bytes) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
        raise SyncError("Excel files need the 'openpyxl' package on the server.") from exc
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is a bad file
        raise SyncError("That file could not be read as an Excel workbook.") from exc
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        return []
    keys = [_cell_text(h).strip() for h in header]
    return [
        dict(zip(keys, [_cell_text(c) for c in row]))
        for row in it
        if any(_cell_text(c).strip() for c in row)
    ]


def rows_from_upload(filename: str, content: bytes) -> list[dict[str, str]]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise SyncError("That file is too large (limit is 5 MB).")
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return _parse_csv_text(content.decode("utf-8-sig", "replace"))
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _rows_from_xlsx(content)
    raise SyncError("Upload a .csv or .xlsx file.")


def _resolve_columns(sample: dict[str, str]) -> dict[str, str]:
    """{store_field: actual_header} for headers present in the sheet."""
    lower_to_actual = {k.strip().lower(): k for k in sample}
    resolved: dict[str, str] = {}
    for field_name, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in lower_to_actual:
                resolved[field_name] = lower_to_actual[alias]
                break
    return resolved


_DATE_FORMATS = (
    "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y",
    "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
)


def _parse_date(raw: str | None) -> date | None:
    """Sheet cell -> date. Day-first for ambiguous forms (Indian sheets); Excel serials accepted."""
    if not raw:
        return None
    text = raw.strip()
    if not text:
        return None
    parsed: date | None = None
    if text.isdigit() and 30000 <= int(text) <= 70000:  # Excel serial day number
        parsed = date(1899, 12, 30) + timedelta(days=int(text))
    else:
        for candidate in (text, text[:10]):  # "2026-09-01 00:00:00" -> first 10 chars
            for fmt in _DATE_FORMATS:
                try:
                    parsed = datetime.strptime(candidate, fmt).date()
                    break
                except ValueError:
                    continue
            if parsed:
                break
    if parsed is None or not (2015 <= parsed.year <= date.today().year + 2):
        return None
    return parsed


def _clean(value: str | None) -> str | None:
    if value is None:
        return None
    v = value.strip()
    return v or None


# --------------------------------------------------------------------------


class SyncError(Exception):
    """Raised for a whole-sheet failure (bad URL, private sheet, missing key column)."""


def _create_region(
    db: Session, org_id: uuid.UUID, name: str, cache: dict[str, Region]
) -> Region:
    name = name.strip()[:128]
    base_code = "".join(c for c in name.upper() if c.isalnum())[:24] or "REGION"
    code = base_code
    n = 1
    existing_codes = {r.code for r in cache.values()}
    while code in existing_codes:
        n += 1
        code = f"{base_code[:22]}{n}"
    region = Region(organization_id=org_id, name=name, code=code, is_active=True)
    db.add(region)
    db.flush()
    cache[name.lower()] = region
    return region


def sync_platform(db: Session, actor: User, partner_slug: str) -> SyncResult:
    partner_slug = partner_slug.strip().lower()
    sources = settings.sheet_sources()
    if partner_slug not in sources:
        raise SyncError(f"No sheet is configured for '{partner_slug}'. Set STORE_SYNC_SHEETS.")
    sheet_id, gid = sources[partner_slug]

    partner = db.execute(
        select(Organization).where(
            Organization.slug == partner_slug,
            Organization.kind == OrganizationKind.PARTNER.value,
        )
    ).scalar_one_or_none()
    if partner is None:
        raise SyncError(f"'{partner_slug}' is not a configured platform in the CRM.")

    return _import_rows(db, actor, partner, partner_slug, _fetch_rows(sheet_id, gid), source="sheet")


def import_file(db: Session, actor: User, partner_slug: str, filename: str, content: bytes) -> SyncResult:
    """Import stores for one platform from an uploaded CSV / Excel file."""
    partner_slug = partner_slug.strip().lower()
    partner = db.execute(
        select(Organization).where(
            Organization.slug == partner_slug,
            Organization.kind == OrganizationKind.PARTNER.value,
        )
    ).scalar_one_or_none()
    if partner is None:
        raise SyncError(f"'{partner_slug}' is not a configured platform in the CRM.")
    return _import_rows(db, actor, partner, partner_slug, rows_from_upload(filename, content), source="file")


def _import_rows(
    db: Session,
    actor: User,
    partner: Organization,
    partner_slug: str,
    rows: list[dict[str, str]],
    *,
    source: str,
) -> SyncResult:
    result = SyncResult(platform=partner_slug, rows_read=len(rows))
    if not rows:
        result.warnings.append("The sheet has no data rows.")
        return result

    cols = _resolve_columns(rows[0])
    if "external_code" not in cols or "name" not in cols:
        raise SyncError(
            "Could not find the 'Outlet ID' and 'Outlet Name' columns in the sheet header."
        )

    regions_by_name = {
        r.name.strip().lower(): r
        for r in db.execute(
            select(Region).where(Region.organization_id == actor.organization_id)
        ).scalars()
    }
    repo = StoreRepository(db)
    seen_codes: set[str] = set()
    trimmed: dict[str, list[int]] = {}  # field -> sheet rows that were cut to fit
    bad_dates: list[int] = []            # rows whose start date couldn't be read

    for i, row in enumerate(rows, start=2):  # sheet row number (1 = header)
        code = _clean(row.get(cols["external_code"]))
        name = _clean(row.get(cols["name"]))
        if not code or not name:
            continue
        if len(code) > _CODE_LIMIT:
            result.warnings.append(f"Row {i}: Outlet ID is longer than {_CODE_LIMIT} characters — skipped.")
            continue
        if code in seen_codes:
            result.warnings.append(f"Row {i}: duplicate Outlet ID '{code}' — skipped.")
            continue
        seen_codes.add(code)

        status = normalize_status(row.get(cols["status"]) if "status" in cols else None)
        region_id: uuid.UUID | None = None
        if "region" in cols:
            raw_region = _clean(row.get(cols["region"]))
            if raw_region:
                key = raw_region.lower().replace(" zone", "").strip()
                base = re.sub(r"\s*\d+$", "", key).strip()  # "South 2" -> "south"
                region = regions_by_name.get(key) or regions_by_name.get(raw_region.lower()) or regions_by_name.get(base)
                if region is None and assignment_service.is_employee_model(partner_slug):
                    pass  # employee-model platforms (Zepto) are scoped by state, not region — never invent regions
                elif region is None:
                    region = _create_region(db, actor.organization_id, raw_region, regions_by_name)
                    result.warnings.append(f"Row {i}: created new region '{region.name}' from zone '{raw_region}'.")
                region_id = region.id if region is not None else None

        fields = {
            "name": name,
            "entity": _clean(row.get(cols["entity"])) if "entity" in cols else None,
            "start_date": _parse_date(row.get(cols["start_date"])) if "start_date" in cols else None,
            "state": _clean(row.get(cols["state"])) if "state" in cols else None,
            "city": _clean(row.get(cols["city"])) if "city" in cols else None,
            "address": _clean(row.get(cols["address"])) if "address" in cols else None,
            "poc_name": _clean(row.get(cols["poc_name"])) if "poc_name" in cols else None,
            "poc_number": _clean(row.get(cols["poc_number"])) if "poc_number" in cols else None,
            "vendor_name": _clean(row.get(cols["vendor_name"])) if "vendor_name" in cols else None,
            "vendor_number": _clean(row.get(cols["vendor_number"])) if "vendor_number" in cols else None,
            "status": status,
            "region_id": region_id,
        }

        if "start_date" in cols and _clean(row.get(cols["start_date"])) and fields["start_date"] is None:
            bad_dates.append(i)

        for fname, limit in _FIELD_LIMITS.items():
            val = fields.get(fname)
            if isinstance(val, str) and len(val) > limit:
                fields[fname] = val[:limit].rstrip()
                trimmed.setdefault(fname, []).append(i)

        existing = repo.get_by_code(partner.id, code)
        if existing is None:
            repo.add(Store(
                organization_id=actor.organization_id,
                partner_organization_id=partner.id,
                external_code=code,
                **fields,
            ))
            result.created += 1
        else:
            changed = False
            for k, v in fields.items():
                # don't wipe a manually-filled field when the sheet cell is blank
                if v is None and getattr(existing, k) is not None and k != "status":
                    continue
                if getattr(existing, k) != v:
                    setattr(existing, k, v)
                    changed = True
            if changed:
                result.updated += 1
            else:
                result.unchanged += 1

    if bad_dates:
        result.warnings.append(
            f"{len(bad_dates)} row{'s' if len(bad_dates) > 1 else ''} had a start date that could not be read "
            f"and was left blank (first: row {bad_dates[0]})."
        )

    for fname, rows_hit in trimmed.items():
        result.warnings.append(
            f"{len(rows_hit)} row{'s' if len(rows_hit) > 1 else ''} had a '{fname}' value longer than "
            f"{_FIELD_LIMITS[fname]} characters and was trimmed (first: row {rows_hit[0]})."
        )

    activity_service.record(
        db, actor=actor, action="store.synced", entity_type="store", entity_id=partner.id,
        metadata={
            "platform": partner_slug,
            "source": source,
            "created": result.created,
            "updated": result.updated,
            "rows_read": result.rows_read,
        },
    )
    db.commit()
    return result


def sync_all(db: Session, actor: User) -> list[SyncResult]:
    out: list[SyncResult] = []
    for slug in settings.sheet_sources():
        try:
            out.append(sync_platform(db, actor, slug))
        except SyncError as exc:
            r = SyncResult(platform=slug)
            r.warnings.append(str(exc))
            out.append(r)
    return out
