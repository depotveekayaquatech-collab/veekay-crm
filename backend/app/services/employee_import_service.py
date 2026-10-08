"""
Import employees (and Zepto state assignments) from the old EMPLOYEES sheet.

EMPLOYEES columns (header match is case-insensitive):
  Employee ID, Employee Name, Password, Region, Status, IsAdmin, Platform,
  IsAccountant, and an optional phone column (a header containing
  contact / phone / mobile, otherwise column I).

State_Assignments (optional tab, Zepto): State, Assigned Employee ID.

Passwords: the sheet stores them in plain text. They are NEVER stored as
text here. By default every *new* user gets a generated temporary password,
returned once in the result so an admin can hand it out. With
`keep_sheet_passwords` the sheet password is hashed on import instead, so
people keep signing in with what they know. Existing users' passwords are
never touched.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.security import generate_temp_password, hash_password
from app.models.organization import Organization, OrganizationKind
from app.models.region import Region
from app.models.role import Role
from app.models.state_assignment import StateAssignment
from app.models.user import User, UserRole, UserStatus
from app.repositories.user_repository import UserRepository
from app.services import activity_service
from app.services.store_sync_service import MAX_UPLOAD_BYTES, SyncError, _create_region, _parse_csv_text

DEFAULT_EMPLOYEE_GRANTS = ["orders.view", "orders.mark", "compliance.upload", "cash.add"]
_TRUE = {"true", "yes", "y", "1", "admin"}


@dataclass
class EmployeeImportResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped: int = 0
    assignments: int = 0
    credentials: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {**self.__dict__, "warnings": self.warnings[:50]}


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------

def _truthy(v: object) -> bool:
    return str(v).strip().lower() in _TRUE if v is not None else False


def _text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_tabs(filename: str, content: bytes) -> dict[str, list[list[object]]]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise SyncError("That file is too large (limit is 5 MB).")
    name = (filename or "").lower()
    if name.endswith((".csv", ".txt")):
        import csv
        text = content.decode("utf-8-sig", "replace")
        return {"employees": [list(r) for r in csv.reader(io.StringIO(text))]}
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise SyncError("That file could not be read as an Excel workbook.") from exc
        wanted = {"employees", "state_assignments"}
        return {
            ws.title.strip().lower(): [list(r) for r in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets
            if ws.title.strip().lower() in wanted  # other tabs are never read
        }
    raise SyncError("Upload a .csv or .xlsx file.")


def _col(header: list[str], *names: str) -> int | None:
    for n in names:
        if n in header:
            return header.index(n)
    return None


# --------------------------------------------------------------------------
# import
# --------------------------------------------------------------------------

def import_employees(
    db: Session, actor: User, filename: str, content: bytes, *, keep_sheet_passwords: bool = False
) -> EmployeeImportResult:
    tabs = _read_tabs(filename, content)
    rows = tabs.get("employees")
    if not rows or len(rows) < 2:
        raise SyncError("No EMPLOYEES sheet with data found in this file.")

    header = [_text(h).lower() for h in rows[0]]
    i_code = _col(header, "employee id", "emp id", "employee code", "id")
    i_name = _col(header, "employee name", "name", "full name")
    if i_code is None or i_name is None:
        raise SyncError("Could not find the 'Employee ID' and 'Employee Name' columns.")
    i_pw = _col(header, "password")
    i_region = _col(header, "region", "zone")
    i_status = _col(header, "status")
    i_admin = _col(header, "isadmin", "is admin", "admin")
    i_plat = _col(header, "platform")
    i_acct = _col(header, "isaccountant", "is accountant", "accountant")
    i_phone = next((i for i, h in enumerate(header) if any(k in h for k in ("contact", "phone", "mobile"))), None)
    if i_phone is None and len(header) > 8:
        i_phone = 8  # spec: phone is column I when there's no labelled column

    org_id = actor.organization_id
    result = EmployeeImportResult()
    urepo = UserRepository(db)

    partners = {
        key: p
        for p in db.execute(
            select(Organization).where(Organization.kind == OrganizationKind.PARTNER.value)
        ).scalars()
        for key in (p.slug.lower(), p.name.lower())
    }
    regions = {
        r.name.strip().lower(): r
        for r in db.execute(select(Region).where(Region.organization_id == org_id)).scalars()
    }
    roles = {r.code: r for r in db.execute(select(Role)).scalars()}
    if "employee" not in roles:
        raise SyncError("Employee role is not configured on the server.")

    def cell(row: list[object], idx: int | None) -> str:
        return _text(row[idx]) if idx is not None and idx < len(row) else ""

    seen: set[str] = set()
    for n, row in enumerate(rows[1:], start=2):
        code, name = cell(row, i_code), cell(row, i_name)
        if not code and not name:
            continue
        if not code or not name:
            result.skipped += 1
            result.warnings.append(f"Row {n}: needs both an Employee ID and a name — skipped.")
            continue
        if len(code) > 32 or code.lower() in seen:
            result.skipped += 1
            result.warnings.append(f"Row {n}: duplicate or over-long Employee ID '{code[:32]}' — skipped.")
            continue
        seen.add(code.lower())

        plat_raw = cell(row, i_plat)
        platform = partners.get(plat_raw.lower()) if plat_raw else None
        if plat_raw and platform is None:
            result.warnings.append(f"Row {n} ({code}): platform '{plat_raw}' is not configured — left blank.")

        region = None
        region_raw = cell(row, i_region)
        if region_raw and platform is not None and platform.slug.lower() not in ("zepto",):
            region = regions.get(region_raw.lower())
            if region is None:
                region = _create_region(db, org_id, region_raw, regions)
                result.warnings.append(f"Row {n} ({code}): created new region '{region.name}'.")

        active = (cell(row, i_status) or "ACTIVE").strip().upper() == "ACTIVE"
        is_admin = _truthy(cell(row, i_admin))
        is_acct = _truthy(cell(row, i_acct))
        phone = cell(row, i_phone)[:32] or None

        wanted_roles = []
        if is_admin:
            wanted_roles.append("admin")
        if is_acct:
            wanted_roles.append("accountant")
        if not wanted_roles:
            wanted_roles.append("employee")
        for rc in wanted_roles:
            if rc not in roles:
                result.warnings.append(f"Row {n} ({code}): role '{rc}' doesn't exist on the server yet — not assigned.")
        wanted_roles = [rc for rc in wanted_roles if rc in roles]

        user = db.execute(
            select(User).where(User.organization_id == org_id, func.lower(User.employee_code) == code.lower())
        ).scalar_one_or_none()

        if user is None:
            sheet_pw = cell(row, i_pw)
            if keep_sheet_passwords and sheet_pw:
                password, temp = sheet_pw, False
                if len(sheet_pw) < 8:
                    result.warnings.append(f"Row {n} ({code}): the sheet password is shorter than 8 characters — consider resetting it.")
            else:
                password, temp = generate_temp_password(), True
            user = User(
                organization_id=org_id,
                platform_organization_id=platform.id if platform else None,
                region_id=region.id if region else None,
                employee_code=code, full_name=name[:128], phone=phone,
                password_hash=hash_password(password),
                must_change_password=True,  # sheet / generated passwords are never the long-term one
                status=UserStatus.ACTIVE.value if active else UserStatus.DEACTIVATED.value,
                is_active=active,
            )
            db.add(user)
            db.flush()
            for rc in wanted_roles:
                db.add(UserRole(user_id=user.id, role_id=roles[rc].id))
            if wanted_roles == ["employee"]:
                urepo.set_direct_permissions(user.id, DEFAULT_EMPLOYEE_GRANTS)
            if temp:
                result.credentials.append({"employee_code": code, "full_name": name, "temp_password": password})
            result.created += 1
            continue

        # existing user: refresh profile + add roles; never touch password or direct permissions
        changed = False
        for attr, val in (
            ("full_name", name[:128]),
            ("phone", phone or user.phone),
            ("platform_organization_id", platform.id if platform else user.platform_organization_id),
            ("region_id", region.id if region else user.region_id),
            ("is_active", active),
            ("status", UserStatus.ACTIVE.value if active else UserStatus.DEACTIVATED.value),
        ):
            if getattr(user, attr) != val:
                setattr(user, attr, val)
                changed = True
        have = set(urepo.get_role_codes(user.id))
        for rc in wanted_roles:
            if rc != "employee" and rc not in have:
                db.add(UserRole(user_id=user.id, role_id=roles[rc].id))
                changed = True
        result.updated += 1 if changed else 0
        result.unchanged += 0 if changed else 1

    db.flush()

    # ---- Zepto state assignments ------------------------------------------
    sa = tabs.get("state_assignments")
    if sa and len(sa) > 1:
        sh = [_text(h).lower() for h in sa[0]]
        i_state = _col(sh, "state")
        i_emp = _col(sh, "assigned employee id", "employee id", "assigned to")
        if i_state is not None and i_emp is not None:
            by_code = {
                u.employee_code.lower(): u
                for u in db.execute(select(User).where(User.organization_id == org_id)).scalars()
            }
            for n, row in enumerate(sa[1:], start=2):
                state, emp = cell(row, i_state), cell(row, i_emp)
                if not state or not emp:
                    continue
                u = by_code.get(emp.lower())
                if u is None or u.platform_organization_id is None:
                    result.warnings.append(f"State_Assignments row {n}: employee '{emp}' not found or has no platform — skipped.")
                    continue
                existing = db.execute(
                    select(StateAssignment).where(
                        StateAssignment.partner_organization_id == u.platform_organization_id,
                        func.lower(StateAssignment.state) == state.lower(),
                    )
                ).scalar_one_or_none()
                if existing is None:
                    db.add(StateAssignment(
                        organization_id=org_id, partner_organization_id=u.platform_organization_id,
                        state=state, assigned_user_id=u.id,
                    ))
                    result.assignments += 1
                elif existing.assigned_user_id != u.id:
                    existing.assigned_user_id = u.id
                    result.assignments += 1
        else:
            result.warnings.append("State_Assignments tab found but it needs 'State' and 'Assigned Employee ID' columns.")

    activity_service.record(
        db, actor=actor, action="employee.imported", entity_type="employee", entity_id=actor.id,
        metadata={  # never includes passwords
            "created": result.created, "updated": result.updated, "assignments": result.assignments,
            "kept_sheet_passwords": keep_sheet_passwords,
        },
    )
    db.commit()
    return result
