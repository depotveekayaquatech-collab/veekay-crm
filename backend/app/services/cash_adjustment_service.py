"""
Developer cash-purchase adjustments: add bottles bought with cash on top of the store's real order entry.

    quantity = ceil(amount / price)          (never rounded down; the developer may override it)
    new order qty = existing entry + quantity   (the employee's entry is never replaced)

Every applied change is stored twice: on the entry (`cash_adjustment`, so employee entry = bottle_count -
cash_adjustment) and as an append-only row in `cash_adjustments` (the developer-only history).
The API layer requires the reserved `cash.adjust` permission for everything in here.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_CEILING, Decimal

from fastapi import HTTPException, status
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.cash_adjustment import CashAdjustment
from app.models.order_entry import EntrySource, OrderEntry
from app.models.store import Store, StoreStatus
from app.models.user import User
from app.schemas.cash_adjustment import (
    AdjustmentList,
    AdjustmentOut,
    AdjustRequest,
    ApplyResult,
    FailedStore,
    Preview,
    PreviewRow,
    StoreChoice,
)
from app.services import activity_service
from app.services.order_service import _check_order_date

_IST = timezone(timedelta(hours=5, minutes=30))
MAX_EXPORT_ROWS = 50000


def calc_quantity(amount: Decimal, price: Decimal) -> int:
    """Whole items the purchase covers, always rounded UP: 100 / 30 = 3.33 -> 4."""
    return int((amount / price).to_integral_value(rounding=ROUND_CEILING))


def _invalid(message: str) -> HTTPException:
    return HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, message)


def _quantities(body: AdjustRequest) -> tuple[int, int]:
    auto = calc_quantity(body.amount, body.price_per_item)
    final = auto if body.quantity is None else body.quantity
    if final < 1:
        raise _invalid("The final quantity must be at least 1.")
    if final > settings.MAX_BOTTLE_COUNT:
        raise _invalid(f"The final quantity can't be more than {settings.MAX_BOTTLE_COUNT} per store.")
    return auto, final


def _live_stores(db: Session, user: User, ids: list[uuid.UUID]) -> dict[uuid.UUID, Store]:
    rows = db.execute(
        select(Store).where(Store.organization_id == user.organization_id, Store.id.in_(set(ids)))
    ).scalars().unique().all()
    return {s.id: s for s in rows}


def _rows(db: Session, user: User, body: AdjustRequest, final: int, *, lock: bool = False) -> list[PreviewRow]:
    stores = _live_stores(db, user, body.store_ids)
    ordered = list(dict.fromkeys(body.store_ids))
    entries = {}
    if ordered:
        stmt = select(OrderEntry).where(OrderEntry.store_id.in_(ordered), OrderEntry.order_date == body.purchase_date)
        entries = {e.store_id: e for e in db.execute(stmt.with_for_update() if lock else stmt).scalars()}
    out: list[PreviewRow] = []
    for sid in ordered:
        s = stores.get(sid)
        if s is None:
            out.append(PreviewRow(store_id=sid, outlet_code="?", outlet_name="Unknown store", has_entry=False,
                                  employee_entry=None, previous=0, added=final, new_total=final, error="Store not found."))
            continue
        e = entries.get(sid)
        previous = e.bottle_count if e else 0
        real = (e.bottle_count - e.cash_adjustment) if e else None
        error = None
        if s.status != StoreStatus.LIVE.value:
            error = "This store isn't live, so it can't take orders."
        elif previous + final > settings.MAX_BOTTLE_COUNT:
            error = f"The new total ({previous + final}) is more than {settings.MAX_BOTTLE_COUNT}."
        out.append(PreviewRow(
            store_id=sid, outlet_code=s.external_code, outlet_name=s.name, has_entry=e is not None,
            employee_entry=real if real else None,
            previous=previous, added=final, new_total=previous + final, error=error,
        ))
    return out


def preview(db: Session, user: User, body: AdjustRequest) -> Preview:
    _check_order_date(body.purchase_date, allow_any_month=True)
    auto, final = _quantities(body)
    rows = _rows(db, user, body, final)
    return Preview(
        purchase_date=body.purchase_date, amount=body.amount, price_per_item=body.price_per_item,
        auto_qty=auto, final_qty=final, overridden=final != auto, total_stores=len(rows), rows=rows,
        can_apply=bool(rows) and all(r.error is None for r in rows),
    )


def apply_to_entry(db: Session, user: User, store: Store, order_date: date, final: int) -> tuple[int, int, int | None]:
    """Add `final` to one store's entry under a row lock. Returns (previous, new_total, employee_entry)."""
    def locked() -> OrderEntry | None:
        return db.execute(
            select(OrderEntry).where(OrderEntry.store_id == store.id, OrderEntry.order_date == order_date).with_for_update()
        ).scalar_one_or_none()

    entry = locked()
    if entry is None:
        try:
            with db.begin_nested():
                db.add(OrderEntry(
                    organization_id=user.organization_id, store_id=store.id, order_date=order_date,
                    bottle_count=final, cash_adjustment=final, source=EntrySource.ADMIN.value, marked_by_user_id=user.id,
                ))
                db.flush()
            return 0, final, None
        except IntegrityError:
            # An employee saved this day a moment ago: add to their entry instead.
            entry = locked()
            if entry is None:
                raise
    previous = entry.bottle_count
    if previous + final > settings.MAX_BOTTLE_COUNT:
        raise _invalid(f"The new total ({previous + final}) is more than {settings.MAX_BOTTLE_COUNT}.")
    entry.bottle_count = previous + final
    entry.cash_adjustment = entry.cash_adjustment + final
    # The employee's real entry (and who marked it) is kept; only the adjustment on top of it grows.
    return previous, entry.bottle_count, entry.bottle_count - entry.cash_adjustment


def apply(db: Session, user: User, body: AdjustRequest) -> ApplyResult:
    _check_order_date(body.purchase_date, allow_any_month=True)
    auto, final = _quantities(body)
    # Validate everything first: if anything is wrong, nothing is changed.
    problems = [r for r in _rows(db, user, body, final) if r.error]
    if problems:
        raise _invalid("; ".join(f"{p.outlet_name} ({p.outlet_code}): {p.error}" for p in problems[:5]))

    stores = _live_stores(db, user, body.store_ids)
    batch = uuid.uuid4()
    applied: list[PreviewRow] = []
    failed: list[FailedStore] = []
    for sid in dict.fromkeys(body.store_ids):
        s = stores[sid]
        try:
            with db.begin_nested():
                previous, total, real = apply_to_entry(db, user, s, body.purchase_date, final)
                db.add(CashAdjustment(
                    organization_id=user.organization_id, batch_id=batch, store_id=s.id,
                    outlet_code=s.external_code, outlet_name=s.name, purchase_date=body.purchase_date,
                    amount=body.amount, price_per_item=body.price_per_item, auto_qty=auto, final_qty=final,
                    previous_qty=previous, final_order_qty=total, created_by_user_id=user.id, created_by_name=user.full_name,
                ))
                db.flush()
            applied.append(PreviewRow(
                store_id=s.id, outlet_code=s.external_code, outlet_name=s.name, has_entry=previous > 0 or real is not None,
                employee_entry=real, previous=previous, added=final, new_total=total,
            ))
        except Exception as exc:  # noqa: BLE001 — report per store, keep going with the rest
            failed.append(FailedStore(
                store_id=s.id, outlet_code=s.external_code, outlet_name=s.name,
                error=exc.detail if isinstance(exc, HTTPException) else "The database rejected this change.",
            ))
    if applied:
        activity_service.record(
            db, actor=user, action="cash_adjustment.applied", entity_type="cash_adjustment", entity_id=batch,
            metadata={"date": body.purchase_date.isoformat(), "amount": str(body.amount), "price": str(body.price_per_item),
                      "auto_qty": auto, "final_qty": final, "stores": len(applied), "failed": len(failed)},
        )
    db.commit()
    return ApplyResult(batch_id=batch, auto_qty=auto, final_qty=final, applied=applied, failed=failed)


# --------------------------------------------------------------------------
# stores to pick from + the developer-only history
# --------------------------------------------------------------------------

def search_stores(db: Session, user: User, q: str | None, limit: int) -> list[StoreChoice]:
    stmt = select(Store).where(Store.organization_id == user.organization_id, Store.status == StoreStatus.LIVE.value)
    if q and q.strip():
        raw = q.strip()
        for ch in ("\\", "%", "_"):
            raw = raw.replace(ch, "\\" + ch)
        like = f"%{raw}%"
        stmt = stmt.where(or_(
            Store.name.ilike(like, escape="\\"), Store.external_code.ilike(like, escape="\\"),
            Store.city.ilike(like, escape="\\"), Store.state.ilike(like, escape="\\"),
        ))
    rows = db.execute(stmt.order_by(Store.name).limit(limit)).scalars().unique().all()
    return [
        StoreChoice(id=s.id, name=s.name, code=s.external_code, state=s.state,
                    platform=s.partner_organization.name if s.partner_organization else None)
        for s in rows
    ]


def _history_stmt(user: User, start: date | None, end: date | None, q: str | None):
    stmt = select(CashAdjustment).where(CashAdjustment.organization_id == user.organization_id)
    if start:
        stmt = stmt.where(CashAdjustment.purchase_date >= start)
    if end:
        stmt = stmt.where(CashAdjustment.purchase_date <= end)
    if q and q.strip():
        raw = q.strip()
        for ch in ("\\", "%", "_"):
            raw = raw.replace(ch, "\\" + ch)
        like = f"%{raw}%"
        stmt = stmt.where(or_(CashAdjustment.outlet_name.ilike(like, escape="\\"), CashAdjustment.outlet_code.ilike(like, escape="\\")))
    return stmt


def _out(a: CashAdjustment) -> AdjustmentOut:
    return AdjustmentOut(
        id=a.id, batch_id=a.batch_id, created_at=a.created_at, developer=a.created_by_name,
        outlet_code=a.outlet_code, outlet_name=a.outlet_name, purchase_date=a.purchase_date, amount=a.amount,
        price_per_item=a.price_per_item, auto_qty=a.auto_qty, final_qty=a.final_qty,
        previous_qty=a.previous_qty, final_order_qty=a.final_order_qty,
        sync_code=f"SYNC-{a.sync_id.hex[:8].upper()}" if a.sync_id else None,
        purchase_code=f"CP-{a.purchase_id.hex[:8].upper()}" if a.purchase_id else None,
        source=a.source, action=a.action,
    )


def history(db: Session, user: User, *, page: int, page_size: int, start: date | None, end: date | None, q: str | None) -> AdjustmentList:
    base = _history_stmt(user, start, end, q)
    total = db.execute(select(func.count()).select_from(base.subquery())).scalar_one()
    rows = db.execute(
        base.order_by(CashAdjustment.created_at.desc(), CashAdjustment.outlet_code).offset((page - 1) * page_size).limit(page_size)
    ).scalars().all()
    return AdjustmentList(items=[_out(r) for r in rows], total=total, page=page, page_size=page_size)


def history_xlsx(db: Session, user: User, *, start: date | None, end: date | None, q: str | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    rows = db.execute(
        _history_stmt(user, start, end, q).order_by(CashAdjustment.created_at, CashAdjustment.outlet_code).limit(MAX_EXPORT_ROWS)
    ).scalars().all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Cash purchase history"
    ws.append(["Sync ID", "Purchase ID", "Timestamp", "Developer", "Outlet ID", "Outlet Name", "Cash Purchase Date",
               "Cash Purchase Amount", "Price Per Item", "Auto Calculated Qty", "Final Qty Applied", "Previous Order Qty",
               "Final Order Qty", "Source", "Action"])
    side = Side(style="thin")
    for c in ws[1]:
        c.font = Font(name="Calibri", size=11, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        c.border = Border(left=side, right=side, top=side, bottom=side)
    for a in rows:
        ws.append([
            f"SYNC-{a.sync_id.hex[:8].upper()}" if a.sync_id else None,
            f"CP-{a.purchase_id.hex[:8].upper()}" if a.purchase_id else None,
            a.created_at.astimezone(_IST).strftime("%d-%b-%Y %H:%M"), a.created_by_name, a.outlet_code, a.outlet_name,
            a.purchase_date, float(a.amount), float(a.price_per_item), a.auto_qty, a.final_qty, a.previous_qty, a.final_order_qty,
            a.source, a.action,
        ])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7).number_format = "dd-mmm-yyyy"
        ws.cell(r, 8).number_format = ws.cell(r, 9).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKLMNO", (16, 14, 19, 22, 14, 32, 17, 16, 13, 14, 14, 14, 14, 22, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
